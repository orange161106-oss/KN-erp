import io
from decimal import Decimal
import openpyxl
import pytest

from app.models.inventory_masters import Consumable, Unit
from app.schemas.stock_statements import RowIssue
from app.services.stock_statement_service import (
    detect_month_in_text,
    export_error_report_excel,
    export_stock_statement_excel,
    get_statement_end_date,
    normalize_header,
    parse_and_preview_stock_statement,
    parse_raw_decimal,
)


def test_normalize_header():
    assert normalize_header("Item Id") == "item id"
    assert normalize_header("Item Description") == "item description"
    assert normalize_header("Cust. Part No.") == "cust part no"
    assert normalize_header("Opening Qty.") == "opening qty"
    assert normalize_header("Closing Value") == "closing value"


def test_parse_raw_decimal():
    issues: list[RowIssue] = []

    # Clean positive
    assert parse_raw_decimal(1500, 1, "Opening Qty.", issues) == Decimal("1500.0000")
    assert parse_raw_decimal("1500.50", 1, "Opening Qty.", issues) == Decimal("1500.5000")

    # Indian digit grouping
    assert parse_raw_decimal("1,71,219.00", 2, "Opening Value", issues) == Decimal("171219.0000")
    assert parse_raw_decimal("10,50,000", 3, "Closing Value", issues) == Decimal("1050000.0000")

    # Parenthesized negatives
    assert parse_raw_decimal("(500.00)", 4, "Diff", issues) == Decimal("-500.0000")

    # Hyphen or dash as zero
    assert parse_raw_decimal("-", 5, "Receipt Qty.", issues) == Decimal("0.0000")

    # Real zero
    assert parse_raw_decimal(0, 6, "Issue Qty.", issues) == Decimal("0.0000")
    assert parse_raw_decimal("0.00", 7, "Issue Value", issues) == Decimal("0.0000")

    # Blank / whitespace as None (never fabricated as zero)
    assert parse_raw_decimal("", 8, "Closing Qty.", issues) is None
    assert parse_raw_decimal(None, 9, "Closing Qty.", issues) is None
    assert parse_raw_decimal("   ", 10, "Closing Qty.", issues) is None


def test_detect_month_in_text():
    assert detect_month_in_text("Stock Statement for October 2026") == (10, 2026)
    assert detect_month_in_text("KNL - Oct'26 Statement") == (10, 2026)
    assert detect_month_in_text("Jan 2025 Inventory") == (1, 2025)
    assert detect_month_in_text("No date here") is None


def test_statement_end_date():
    assert str(get_statement_end_date(2026, 10)) == "2026-10-31"
    assert str(get_statement_end_date(2024, 2)) == "2024-02-29"  # Leap year
    assert str(get_statement_end_date(2025, 2)) == "2025-02-28"


def create_statement_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    # Row 1: Title
    ws.cell(row=1, column=1, value="KNL Consumables Stock Statement - October 2026")

    # Row 2: Headers (12 exact columns)
    headers = [
        "Item Id", "Item Description", "Cust. Part No.", "Unit",
        "Opening Qty.", "Opening Value", "Receipt Qty.", "Receipt Value",
        "Issue Qty.", "Issue Value", "Closing Qty.", "Closing Value",
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=2, column=col_idx, value=h)

    # Row 3: Item 1 (Balanced)
    row3 = [
        "TEST-CONS-01", "Welding Wire ER70S-6", "CPN-01", "KGS",
        "100.00", "15000.00", "50.00", "7500.00", "30.00", "4500.00", "120.00", "18000.00"
    ]
    for col_idx, val in enumerate(row3, start=1):
        ws.cell(row=3, column=col_idx, value=val)

    # Row 4: Item 2 (Arithmetic mismatch warning: Closing Qty 100 instead of 90)
    row4 = [
        "TEST-CONS-02", "Grinding Wheel 4 inch", "CPN-02", "NOS",
        "100.00", "5000.00", "10.00", "500.00", "20.00", "1000.00", "100.00", "4500.00"
    ]
    for col_idx, val in enumerate(row4, start=1):
        ws.cell(row=4, column=col_idx, value=val)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_error_report():
    issues = [
        RowIssue(row=3, column="Item Id", description="Item 'UNKNOWN-01' not in master", severity="ERROR"),
        RowIssue(row=4, column="Closing Qty.", description="Closing Qty mismatch", severity="WARNING"),
    ]
    excel_bytes, filename = export_error_report_excel(issues, "test_upload.xlsx")
    assert len(excel_bytes) > 0
    assert filename.endswith(".xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Excel Row"
    assert ws.cell(row=2, column=3).value == "ERROR"
    assert ws.cell(row=2, column=4).value == "Item 'UNKNOWN-01' not in master"
