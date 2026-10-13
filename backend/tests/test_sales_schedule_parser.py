import io
from decimal import Decimal
import openpyxl
import pytest
from sqlalchemy import select

from app.models.masters import Customer, CustomerAlias
from app.services.sales_schedule_parser import (
    normalize_alias,
    parse_sales_schedule_data,
    seed_default_customers_and_aliases,
)


def create_knl_sample_workbook() -> bytes:
    """Create a synthetic workbook mimicking KNL's 'Pro.Order.pdf' layout."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prd. Order"

    # Row 1: Month label
    ws.cell(row=1, column=1, value="")
    ws.cell(row=1, column=2, value="")
    ws.cell(row=1, column=3, value="")
    ws.cell(row=1, column=4, value="")
    ws.cell(row=1, column=5, value="Oct'26")  # Qty column R0
    ws.cell(row=1, column=6, value="Oct'26")  # Qty column R1
    ws.cell(row=1, column=7, value="Oct'26 R0 Value (in Lakhs)")  # Value column to ignore

    # Row 2: Schedule given on
    ws.cell(row=2, column=5, value="22.09.2026")
    ws.cell(row=2, column=6, value="28.09.2026")
    ws.cell(row=2, column=7, value="")

    # Row 3: Month date
    ws.cell(row=3, column=5, value="01.10.2026")
    ws.cell(row=3, column=6, value="01.10.2026")
    ws.cell(row=3, column=7, value="")

    # Row 4: Revision label / Columns
    ws.cell(row=4, column=1, value="S. No")
    ws.cell(row=4, column=2, value="Part No.")
    ws.cell(row=4, column=3, value="Item ID")
    ws.cell(row=4, column=4, value="Part Description")
    ws.cell(row=4, column=5, value="R0")
    ws.cell(row=4, column=6, value="R1")
    ws.cell(row=4, column=7, value="Value")

    # Block 1: ASHOK LEYLAND (No heading banner! Starts at row 5)
    # Part 1
    ws.cell(row=5, column=1, value=1)
    ws.cell(row=5, column=2, value="AL-PART-01")
    ws.cell(row=5, column=3, value="ITM-AL-01")
    ws.cell(row=5, column=4, value="BRACKET ASSY AL")
    ws.cell(row=5, column=5, value=100)
    ws.cell(row=5, column=6, value=120)
    ws.cell(row=5, column=7, value=5.4)

    # Part 2
    ws.cell(row=6, column=1, value=2)
    ws.cell(row=6, column=2, value="AL-PART-02")
    ws.cell(row=6, column=3, value="ITM-AL-02")
    ws.cell(row=6, column=4, value="FLANGE SHAFT AL")
    ws.cell(row=6, column=5, value=250)
    ws.cell(row=6, column=6, value=250)
    ws.cell(row=6, column=7, value=12.1)

    # Closing Row for Block 1
    ws.cell(row=7, column=4, value="ASHOK LEYLAND - TOTAL")
    ws.cell(row=7, column=5, value=350)  # Sum: 100 + 250 = 350
    ws.cell(row=7, column=6, value=370)  # Sum: 120 + 250 = 370
    ws.cell(row=7, column=7, value=17.5)

    # Block 2: TAFE FED Component (Has heading banner)
    ws.cell(row=8, column=4, value="TAFE FED Component")  # Heading
    # Part 1
    ws.cell(row=9, column=1, value=3)
    ws.cell(row=9, column=2, value="TF-COMP-01")
    ws.cell(row=9, column=3, value="ITM-TF-01")
    ws.cell(row=9, column=4, value="GEAR HOUSING TAFE")
    ws.cell(row=9, column=5, value=50)
    ws.cell(row=9, column=6, value=60)
    # Closing Row for Block 2
    ws.cell(row=10, column=4, value="TAFE FED Component - TOTAL")
    ws.cell(row=10, column=5, value=50)
    ws.cell(row=10, column=6, value=60)

    # Block 3: JOB WORK
    ws.cell(row=11, column=4, value="JOB WORK")  # Heading
    # Part 1
    ws.cell(row=12, column=1, value=4)
    ws.cell(row=12, column=2, value="JW-OP-01")
    ws.cell(row=12, column=3, value="ITM-JW-01")
    ws.cell(row=12, column=4, value="MACHINING OF CRANK CASE")
    ws.cell(row=12, column=5, value=80)
    ws.cell(row=12, column=6, value=80)
    # Closing Row for Block 3
    ws.cell(row=13, column=4, value="JOB WORK - TOTAL")
    ws.cell(row=13, column=5, value=80)
    ws.cell(row=13, column=6, value=80)

    # Grand Total
    ws.cell(row=14, column=4, value="Grand Total")
    ws.cell(row=14, column=5, value=480)  # 350 + 50 + 80 = 480
    ws.cell(row=14, column=6, value=510)  # 370 + 60 + 80 = 510

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_alias_normalizer():
    assert normalize_alias("ASHOK LEYLAND - TOTAL") == "ashok leyland"
    assert normalize_alias("TAFE FED Component - TOTAL") == "tafe fed component"
    assert normalize_alias("TVS CHILD PARTS - TOTAL") == "tvs child parts"
    assert normalize_alias("JOB WORK - TOTAL") == "job work"
    assert normalize_alias("GREAVES MOBILITY TOTAL") == "greaves mobility"
    assert normalize_alias("  MINDA   ") == "minda"


def test_sales_schedule_two_pass_parsing(db_session):
    seed_default_customers_and_aliases(db_session)
    excel_bytes = create_knl_sample_workbook()

    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))

    res = parse_sales_schedule_data(rows, db_session)

    # 1. Available Revisions check
    revs = [r["revision_label"] for r in res["available_revisions"]]
    assert "R0" in revs
    assert "R1" in revs
    # Value column must NOT be in revisions
    assert not any("VALUE" in r["revision_label"] for r in res["available_revisions"])

    # 2. Blocks check
    blocks = res["blocks"]
    assert len(blocks) == 3

    # Block 1: First block Ashok Leyland has no heading row, named from closing row
    b1 = blocks[0]
    assert b1["customer_name"] == "Ashok Leyland"
    assert b1["customer_code"] == "CUST-AL"
    assert b1["part_count"] == 2
    assert b1["reconciliation_ok"] is True
    assert b1["quantities"]["2026-10_R0"] == "350.0000"
    assert b1["quantities"]["2026-10_R1"] == "370.0000"

    # Block 2: TAFE FED Component
    b2 = blocks[1]
    assert b2["customer_name"] == "TAFE FED"
    assert b2["group_label"] == "Component"
    assert b2["part_count"] == 1
    assert b2["reconciliation_ok"] is True

    # Block 3: JOB WORK
    b3 = blocks[2]
    assert b3["customer_name"] == "Job Work"
    assert b3["section_type"] == "JOB_WORK"
    assert b3["part_count"] == 1
    assert b3["reconciliation_ok"] is True

    # 3. Grand Total check
    gt = res["grand_total"]
    assert gt["reconciliation_ok"] is True
    assert gt["parts_sum"]["2026-10_R0"] == "480.0000"
    assert gt["sheet_grand_total"]["2026-10_R0"] == "480.0000"

    # 4. Flat rows check
    flat_rows = res["_all_flat_rows"]
    # 4 parts * 2 revisions = 8 flat rows
    assert len(flat_rows) == 8
    # Every row has Customer, Customer ID
    for row in flat_rows:
        assert row["customer"] in ("Ashok Leyland", "TAFE FED", "Job Work")
        assert row["customer_code"] in ("CUST-AL", "CUST-TAFE-FED", "CUST-JOBWORK")
        assert row["revision"] in ("R0", "R1")

    assert res["can_confirm"] is True
    assert len(res["unmapped_aliases"]) == 0


def test_sales_schedule_mismatch_and_unknown_heading(db_session):
    seed_default_customers_and_aliases(db_session)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Prd. Order"

    # Header rows
    ws.cell(row=1, column=4, value="Part Description")
    ws.cell(row=1, column=5, value="Oct'26")
    ws.cell(row=2, column=5, value="22.09.2026")
    ws.cell(row=3, column=5, value="01.10.2026")
    ws.cell(row=4, column=1, value="S. No")
    ws.cell(row=4, column=2, value="Part No.")
    ws.cell(row=4, column=3, value="Item ID")
    ws.cell(row=4, column=4, value="Part Description")
    ws.cell(row=4, column=5, value="R0")

    # Unknown Heading block
    ws.cell(row=5, column=4, value="MYSTERY AUTOMOTIVE")  # Unknown heading!
    ws.cell(row=6, column=1, value=1)
    ws.cell(row=6, column=2, value="MYS-01")
    ws.cell(row=6, column=3, value="ITM-MYS-01")
    ws.cell(row=6, column=4, value="ROD")
    ws.cell(row=6, column=5, value=100)
    # Closing row with mismatched total (says 150 but parts sum is 100)
    ws.cell(row=7, column=4, value="MYSTERY AUTOMOTIVE - TOTAL")
    ws.cell(row=7, column=5, value=150)

    # Grand Total (says 200 but sum is 100)
    ws.cell(row=8, column=4, value="Grand Total")
    ws.cell(row=8, column=5, value=200)

    rows = list(ws.iter_rows(values_only=True))
    res = parse_sales_schedule_data(rows, db_session)

    # Unknown heading blocked
    assert "MYSTERY AUTOMOTIVE" in res["unmapped_aliases"]
    assert res["can_confirm"] is False
    assert any("MYSTERY AUTOMOTIVE" in err for err in res["errors"])

    # Reconciliation mismatch detected
    b = res["blocks"][0]
    assert b["reconciliation_ok"] is False
    assert any("differs from sheet total" in w for w in res["warnings"])
    assert res["grand_total"]["reconciliation_ok"] is False
    assert any("Grand total check failed" in w for w in res["warnings"])
