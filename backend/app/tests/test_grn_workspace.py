import io
from uuid import uuid4
import openpyxl
from openpyxl.chartsheet.chartsheet import Chartsheet
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import Settings
from app.db.base import Base
from app.db.session import get_db
from app.main import create_app
from app.models.auth import Permission, Role, User
from app.models.grn import GoodsReceiptRecord
from app.modules.po_grn.router import order_database
from app.security.passwords import PasswordService
from app.security.tokens import create_access_token

BASE = '/api/v1/grns/workspace'
ALL_WORKSPACE_PERMISSIONS = [
    'purchase.grns.read',
    'purchase.grns.create',
    'purchase.grns.update',
    'purchase.grns.delete',
    'purchase.grns.import',
    'purchase.grns.export',
]


@pytest.fixture
def workspace_client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def configure(connection, record):
        connection.create_function("btrim", 1, lambda value: value.strip())
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    settings = Settings(
        _env_file=None, app_env="test",
        database_url="postgresql+psycopg://test_user@127.0.0.1/kn_unit_test",
        auth_secret_key="test-signing-key-for-workspace-testing",
    )
    passwords = PasswordService()

    with Session(engine, expire_on_commit=False) as session:
        grants = [Permission(code=p, description='Workspace test grant') for p in ALL_WORKSPACE_PERMISSIONS]
        role = Role(code="WORKSPACE_TESTER", name="Workspace tester", permissions=grants)
        user = User(username="workspace_operator", password_hash=passwords.hash("password"), roles=[role])
        session.add(user)
        session.commit()

        app = create_app(settings)

        def database():
            with Session(engine, expire_on_commit=False, autoflush=False) as req_session:
                yield req_session

        app.dependency_overrides[get_db] = database
        app.dependency_overrides[order_database] = database
        headers = {"Authorization": "Bearer " + create_access_token(user.id, settings)}

        with TestClient(app) as client:
            yield client, session, user, headers
    engine.dispose()


def create_sample_workbook(sheets_data: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    first = True
    for sheet_name, rows in sheets_data.items():
        if first:
            ws = wb.active
            ws.title = sheet_name
            first = False
        else:
            ws = wb.create_sheet(title=sheet_name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_workspace_records_permission_denied(workspace_client):
    client, session, user, headers = workspace_client
    # Clear permissions
    user.roles[0].permissions = []
    session.commit()
    res = client.get(BASE + '/records', headers=headers)
    assert res.status_code == 403, res.text


def test_workspace_save_and_list_records(workspace_client):
    client, session, user, headers = workspace_client

    # 1. Save new draft rows
    payload = {
        'records': [
            {
                'id': str(uuid4()),
                'row_index': 1,
                'part_number': '1222A00201',
                'item_id': 'ITEM-001',
                'description': 'Seat Spacer',
                'quantity': '10.0000',
                'unit': 'Nos',
                'po_number': 'PO-9001',
                'supplier_name': 'Alpha Supplies',
                'status': 'SAVED',
            },
            {
                'id': str(uuid4()),
                'row_index': 2,
                'part_number': '1311A01702',
                'item_id': 'ITEM-002',
                'description': 'Pivot Brake',
                'quantity': '5.0000',
                'unit': 'Nos',
                'po_number': 'PO-9002',
                'supplier_name': 'Beta Metals',
                'status': 'SAVED',
            },
        ],
        'deleted_ids': [],
        'reason': 'Initial test batch',
    }
    save_res = client.post(BASE + '/save', headers=headers, json=payload)
    assert save_res.status_code == 200, save_res.text
    data = save_res.json()
    assert data['saved_count'] == 2
    assert len(data['records']) == 2

    # 2. List records
    list_res = client.get(BASE + '/records', headers=headers)
    assert list_res.status_code == 200
    records = list_res.json()
    assert len(records) == 2
    assert records[0]['part_number'] == '1222A00201'
    assert records[0]['description'] == 'Seat Spacer'
    assert records[1]['part_number'] == '1311A01702'

    # 3. Search records
    search_res = client.get(BASE + '/records?search=Pivot', headers=headers)
    assert search_res.status_code == 200
    search_records = search_res.json()
    assert len(search_records) == 1
    assert search_records[0]['part_number'] == '1311A01702'


def test_workspace_delete_record(workspace_client):
    client, session, user, headers = workspace_client

    rec_id = str(uuid4())
    client.post(BASE + '/save', headers=headers, json={
        'records': [{
            'id': rec_id,
            'part_number': 'DEL-001',
            'item_id': 'DEL-ITEM',
            'description': 'To be deleted',
            'quantity': '1.0000',
            'unit': 'Nos',
            'status': 'SAVED',
        }],
        'deleted_ids': [],
        'reason': 'Setup row for deletion',
    })

    # Delete row
    del_res = client.delete(BASE + f'/records/{rec_id}?reason=Confirmed+deletion', headers=headers)
    assert del_res.status_code == 204

    # Verify not returned in active list
    list_res = client.get(BASE + '/records', headers=headers)
    assert not any(r['id'] == rec_id for r in list_res.json())


def test_workspace_excel_inspect_and_import(workspace_client):
    client, session, user, headers = workspace_client

    sheets = {
        'Consumable Plan': [
            ['Part Number', 'Item ID', 'Description', 'Quantity', 'Unit', 'PO Number', 'Supplier'],
            ['1222A00201', 'ITEM-1', 'Seat Spacer', 10, 'Nos', 'PO-101', 'Alpha Corp'],
            ['1321A02401', 'ITEM-2', 'Radiator RH', 2, 'Nos', 'PO-102', 'Beta Corp'],
        ],
        'Production': [
            ['Part Number', 'Item ID', 'Description', 'Quantity'],
            ['PROD-999', 'P-1', 'Production Item', 50],
        ],
    }
    excel_bytes = create_sample_workbook(sheets)

    # 1. Inspect Excel
    files = {'file': ('consumable_plan.xlsx', excel_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    inspect_res = client.post(BASE + '/inspect', headers=headers, files=files)
    assert inspect_res.status_code == 200, inspect_res.text
    inspect_data = inspect_res.json()
    assert inspect_data['filename'] == 'consumable_plan.xlsx'
    assert len(inspect_data['sheets']) == 2
    sheet_names = [s['name'] for s in inspect_data['sheets']]
    assert 'Consumable Plan' in sheet_names
    assert 'Production' in sheet_names

    plan_sheet = next(s for s in inspect_data['sheets'] if s['name'] == 'Consumable Plan')
    assert plan_sheet['row_count'] == 2
    assert 'Part Number' in plan_sheet['headers']

    # 2. Import selected sheet
    files = {'file': ('consumable_plan.xlsx', excel_bytes, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    import_res = client.post(
        BASE + '/import-sheet',
        headers=headers,
        files=files,
        data={'sheet_name': 'Consumable Plan', 'mode': 'APPEND', 'reason': 'Testing import'},
    )
    assert import_res.status_code == 200, import_res.text
    import_data = import_res.json()
    assert import_data['imported_count'] == 2
    assert len(import_data['records']) == 2
    assert import_data['records'][0]['part_number'] == '1222A00201'


def test_workspace_excel_export(workspace_client):
    client, session, user, headers = workspace_client

    client.post(BASE + '/save', headers=headers, json={
        'records': [{
            'id': str(uuid4()),
            'part_number': 'EXP-101',
            'item_id': 'ITEM-EXP',
            'description': 'Exportable Item',
            'quantity': '42.5000',
            'unit': 'Nos',
            'status': 'SAVED',
        }],
        'deleted_ids': [],
        'reason': 'Setup for export',
    })

    export_res = client.get(BASE + '/export', headers=headers)
    assert export_res.status_code == 200
    assert 'spreadsheetml.sheet' in export_res.headers['content-type']
    content = export_res.content
    assert len(content) > 100

    # Parse exported workbook to verify integrity
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert 'Goods Receipts' in wb.sheetnames
    ws = wb['Goods Receipts']
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][1] == 'Part Number'
    assert any(r[1] == 'EXP-101' for r in rows[1:])


def test_workspace_chartsheet_handling(workspace_client, monkeypatch):
    client, session, user, headers = workspace_client
    from unittest.mock import MagicMock

    mock_wb = MagicMock()
    mock_wb.sheetnames = ['DataSheet', 'ChartSheet1']

    mock_data_sheet = MagicMock()
    mock_data_sheet.iter_rows.return_value = [
        ['Part Number', 'Item ID', 'Description', 'Quantity'],
        ['CHART-01', 'ITM-01', 'Chart Item', 10],
    ]

    mock_chart_sheet = MagicMock(spec=[])  # Has no iter_rows attribute, simulating openpyxl Chartsheet

    def get_sheet(name):
        return mock_data_sheet if name == 'DataSheet' else mock_chart_sheet

    mock_wb.__getitem__.side_effect = get_sheet
    monkeypatch.setattr('openpyxl.load_workbook', lambda *args, **kwargs: mock_wb)

    files = {'file': ('with_chart.xlsx', b'dummy_excel_bytes', 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
    inspect_res = client.post(BASE + '/inspect', headers=headers, files=files)
    assert inspect_res.status_code == 200, inspect_res.text
    data = inspect_res.json()
    sheet_names = [s['name'] for s in data['sheets']]
    assert 'DataSheet' in sheet_names
    assert 'ChartSheet1' not in sheet_names


def test_workspace_csv_inspect_and_import(workspace_client):
    client, session, user, headers = workspace_client

    csv_content = b"Part Number,Item ID,Description,Quantity,Unit\nCSV-100,CSV-ITEM,CSV Cable,25,Nos\n"

    # 1. Inspect CSV
    files = {'file': ('consumables.csv', csv_content, 'text/csv')}
    inspect_res = client.post(BASE + '/inspect', headers=headers, files=files)
    assert inspect_res.status_code == 200, inspect_res.text
    inspect_data = inspect_res.json()
    assert len(inspect_data['sheets']) == 1
    assert inspect_data['sheets'][0]['row_count'] == 1

    # 2. Import CSV
    files = {'file': ('consumables.csv', csv_content, 'text/csv')}
    import_res = client.post(
        BASE + '/import-sheet',
        headers=headers,
        files=files,
        data={'sheet_name': 'CSV Data', 'mode': 'APPEND', 'reason': 'Testing CSV import'},
    )
    assert import_res.status_code == 200, import_res.text
    import_data = import_res.json()
    assert import_data['imported_count'] == 1
    assert import_data['records'][0]['part_number'] == 'CSV-100'

