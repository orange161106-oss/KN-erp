"""Tests for Tool Life (Output per Unit) requirement and purchase calculation (Step 1 acceptance tests T1-T5, T7)."""
import math
from decimal import Decimal
from io import BytesIO
import openpyxl
import pytest

from app.domain.purchase_engine.monthly import calculate_monthly_purchase
from app.models.inventory_masters import Consumable, Supplier, Unit
from app.models.prd import PlanningVersion
from app.models.requirements import CalculatedRequirement
from app.schemas.auth import CurrentUser
from app.services.purchase_workspace import list_purchase_workspace, _export_purchase_view
from app.services.requirements_workspace import export_requirements_excel
from app.schemas.requirements_workspace import RequirementWorkspaceRecord
from app.tests.test_permission_foundation import db, actor
from app.tests.test_central_basic_planning import catalog


def test_t1_aac100001qc_plug_tap_purchase_math():
    """T1: AAC100001QC Plug tap: Sch 5550, tool life 2000, stock 17, MOQ 10, rate 246 -> Req=3, Order=0, Value=Rs 0."""
    gross_sch = Decimal('5550')
    tool_life = Decimal('2000')
    effective_req = Decimal(str(math.ceil(gross_sch / tool_life)))
    assert effective_req == Decimal('3')

    # Formula: MSL (4) + Req (3) - Stock (17) = -10 <= 0 -> Order 0, Value 0
    res = calculate_monthly_purchase(
        required=effective_req,
        closing_stock=Decimal('17'),
        msl=Decimal('4'),
        moq=Decimal('10'),
        rate=Decimal('246')
    )
    assert res['raw_quantity'] == Decimal('-10')
    assert res['recommended_quantity'] == Decimal('0')
    assert res['purchase_value'] == Decimal('0')


def test_t2_aac100002qc_plug_tap_purchase_math():
    """T2: AAC100002QC: Sch 9432, tool life 2000 -> Req=5. Stock 0, MSL 4, MOQ 10, Rate 246 -> Order=10, Value=Rs 2460."""
    gross_sch = Decimal('9432')
    tool_life = Decimal('2000')
    effective_req = Decimal(str(math.ceil(gross_sch / tool_life)))
    assert effective_req == Decimal('5')

    # Formula: MSL (4) + Req (5) - Stock (0) = 9 -> Raised to MOQ (10) -> Order 10, Value 2460
    res = calculate_monthly_purchase(
        required=effective_req,
        closing_stock=Decimal('0'),
        msl=Decimal('4'),
        moq=Decimal('10'),
        rate=Decimal('246')
    )
    assert res['raw_quantity'] == Decimal('9')
    assert res['recommended_quantity'] == Decimal('10')
    assert res['purchase_value'] == Decimal('2460')


def test_t3_aac100003qc_plug_tap_purchase_math():
    """T3: AAC100003QC: Sch 2868, tool life 3759 -> Req=1. Stock 0, MSL 5, MOQ 10, Rate 270 -> Order=10, Value=Rs 2700."""
    gross_sch = Decimal('2868')
    tool_life = Decimal('3759')
    effective_req = Decimal(str(math.ceil(gross_sch / tool_life)))
    assert effective_req == Decimal('1')

    # Formula: MSL (5) + Req (1) - Stock (0) = 6 -> Raised to MOQ (10) -> Order 10, Value 2700
    res = calculate_monthly_purchase(
        required=effective_req,
        closing_stock=Decimal('0'),
        msl=Decimal('5'),
        moq=Decimal('10'),
        rate=Decimal('270')
    )
    assert res['raw_quantity'] == Decimal('6')
    assert res['recommended_quantity'] == Decimal('10')
    assert res['purchase_value'] == Decimal('2700')


def test_t4_nbq000031qc_production_rate():
    """T4: NBQ000031QC: Tool life 1 -> Required = Gross (4788). No division."""
    gross_sch = Decimal('4788')
    tool_life = Decimal('1')
    rule_type = 'PRODUCTION_RATE'
    if rule_type == 'TOOL_LIFE' and tool_life > 1:
        effective_req = Decimal(str(math.ceil(gross_sch / tool_life)))
    else:
        effective_req = gross_sch
    assert effective_req == Decimal('4788')


def test_t5_missing_output_per_unit_flagged(db, catalog):
    """T5: Output per Unit blank on a TOOL_LIFE consumable is flagged OUTPUT_PER_UNIT_MISSING, no purchase value."""
    catalog[5].rule_type = "TOOL_LIFE"
    catalog[5].output_per_unit = None
    db.commit()

    from app.tests.test_central_basic_planning import calculated
    calculated(db, catalog)

    user = actor(is_super_admin=True)
    res = list_purchase_workspace(db, month=10, year=2026, current_user=user)
    row = next((r for r in res.records if r.consumable_code == catalog[5].code), None)
    assert row is not None
    assert row.status == "OUTPUT_PER_UNIT_MISSING"
    assert row.remarks == "Output per Unit missing"
    assert row.purchase_value is None
    assert row.system_recommended_qty is None


def test_t7_normal_view_output_unit_is_numeric():
    """T7: Output / Unit column in Normal View is numeric ('n') for tool-life items."""
    from app.schemas.purchase_workspace import PurchaseWorkspaceRecord
    record = PurchaseWorkspaceRecord(
        id="1:1:1",
        planning_period="2026-10",
        planning_month=10,
        planning_year=2026,
        source_row_count=1,
        plan_id="1",
        revision="R0",
        consumable_code="AAC100001QC",
        consumable_name="Plug tap",
        unit="Nos",
        gross_sch=Decimal('5550'),
        output_per_unit=Decimal('2000'),
        output_unit="2000",
        required_qty=Decimal('3'),
        minimum_order_qty=Decimal('10'),
        minimum_stock_level=Decimal('4'),
        stock_qty=Decimal('17'),
        system_recommended_qty=Decimal('0'),
        purchase_value=Decimal('0'),
        unit_rate=Decimal('246'),
        supplier_code="SUVSGENHO01",
        supplier_name="V.s.g enterprises",
        category="Tool",
        status="NO_PURCHASE_REQUIRED"
    )

    excel_bytes = _export_purchase_view([record], month=10, year=2026, view="normal")
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    # Find the 'Output / Unit' and 'Sch. Oct-26' columns
    headers = [ws.cell(3, col).value for col in range(1, ws.max_column + 1)]
    assert "Output / Unit" in headers
    output_col = headers.index("Output / Unit") + 1
    sch_col = headers.index("Sch. Oct-26") + 1

    cell = ws.cell(4, output_col)
    assert cell.data_type == "n", f"Expected numeric type 'n', got {cell.data_type} with value {cell.value}"
    assert cell.value == 2000

    sch_cell = ws.cell(4, sch_col)
    assert sch_cell.value == Decimal('5550') or sch_cell.value == 5550


def test_requirements_export_has_output_per_unit_and_required_qty(db, catalog):
    """Verify that export_requirements_excel includes Output per Unit and Required Qty."""
    catalog[5].rule_type = "TOOL_LIFE"
    catalog[5].output_per_unit = Decimal('2000')
    catalog[5].purchasing_unit = "Nos"
    db.commit()

    from app.tests.test_central_basic_planning import calculated
    calculated(db, catalog)

    excel_bytes = export_requirements_excel(db, month=10, year=2026)
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    ws = wb.active

    col_headers = [ws.cell(5, c).value for c in range(1, ws.max_column + 1)]
    assert "Output per Unit" in col_headers
    assert "Required Qty" in col_headers

    out_col = col_headers.index("Output per Unit") + 1
    req_col = col_headers.index("Required Qty") + 1

    # Row 6 is the first data row
    out_cell = ws.cell(6, out_col)
    assert out_cell.value == 2000
    assert out_cell.data_type == "n"



def test_t6_supplier_and_purchasing_unit(db, catalog):
    """T6: Supplier code, name and purchasing unit ('Nos') are populated; no 'Not selected' or hardcoded 'PCS'."""
    supplier = catalog[6]
    material = catalog[5]
    material.rule_type = "TOOL_LIFE"
    material.output_per_unit = Decimal('2000')
    material.purchasing_unit = "Nos"
    material.supplier_id = supplier.id
    db.commit()

    from app.tests.test_central_basic_planning import calculated
    calculated(db, catalog)

    user = actor(is_super_admin=True)
    res = list_purchase_workspace(db, month=10, year=2026, current_user=user)
    row = next((r for r in res.records if r.consumable_code == material.code), None)
    assert row is not None
    assert row.supplier_code == supplier.code
    assert row.supplier_name == supplier.name
    assert row.unit == "Nos"
    assert row.supplier_code != "Not selected"
    assert row.supplier_name != "Not selected"
    assert row.unit != "PCS"

