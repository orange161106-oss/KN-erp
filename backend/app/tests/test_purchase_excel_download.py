from io import BytesIO

from openpyxl import load_workbook

from app.tests.test_purchase_workspace import purchase_client


def test_purchase_excel_download_requires_auth_and_preserves_selection_and_filters(purchase_client):
    client, _, _, headers = purchase_client
    base = '/api/v1/purchasing/workspace'
    export = f'{base}/export?month=10&year=2026'
    assert client.get(export).status_code == 401
    row = client.get(f'{base}?month=10&year=2026', headers=headers).json()['records'][0]
    response = client.get(export, params={'month':10, 'year':2026, 'record_ids': [row['id']]}, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    assert 'purchase_recommendations_2026_10.xlsx' in response.headers['content-disposition']
    workbook = load_workbook(BytesIO(response.content))
    assert workbook.active.cell(5, 1).value == 'WLD-01'
    assert workbook.active.cell(5, 6).value == 80
    workbook.close()
    for params in ({'record_ids': ['not-selected']}, {'search': 'no-matching-material'}, {'status': 'SUBMITTED'}):
        filtered = client.get(export, params={'month':10, 'year':2026, **params}, headers=headers)
        assert filtered.status_code == 200
        workbook = load_workbook(BytesIO(filtered.content))
        assert workbook.active.max_row == 4
        workbook.close()
