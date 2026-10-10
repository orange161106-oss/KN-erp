import io
from uuid import uuid4
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import Settings
from app.db.base import Base
from app.db.session import get_db
from app.main import create_app
from app.models.auth import Permission, Role, User
from app.models.prd import PRDRecord
from app.security.passwords import PasswordService
from app.security.tokens import create_access_token

PRD_BASE = '/api/v1/prd/workspace'
REQ_BASE = '/api/v1/requirements/workspace'
ALL_PERMISSIONS = [
    'prd.plan.read',
    'prd.plan.create',
    'prd.plan.update',
    'prd.plan.delete',
    'prd.plan.import',
    'prd.plan.export',
    'inventory.stock.read',
]


@pytest.fixture
def workspace_client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def configure(connection, record):
        connection.create_function("btrim", 1, lambda value: value.strip())
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    settings = Settings(
        _env_file=None,
        app_env="test",
        database_url="postgresql+psycopg://test_user@127.0.0.1/kn_unit_test",
        auth_secret_key="test-signing-key-for-workspace-testing",
    )
    passwords = PasswordService()

    with Session(engine, expire_on_commit=False) as session:
        grants = [Permission(code=p, description='Workspace test grant') for p in ALL_PERMISSIONS]
        role = Role(code="ADMIN", name="Administrator", permissions=grants)
        user = User(username="admin_operator", password_hash=passwords.hash("password"), roles=[role], can_view_planning=True, can_run_calculations=True)
        session.add(user)
        session.commit()

        app = create_app(settings)

        def database():
            with Session(engine, expire_on_commit=False, autoflush=False) as req_session:
                yield req_session

        app.dependency_overrides[get_db] = database
        token = create_access_token(user.id, settings)
        with TestClient(app) as client:
            client.headers.update({"Authorization": f"Bearer {token}"})
            yield client, session, user
    engine.dispose()


def test_prd_save_and_list(workspace_client):
    client, session, user = workspace_client

    payload = {
        "records": [
            {
                "row_index": 1,
                "plant": "Plant 1",
                "customer": "Toyota",
                "product_code": "PRD-T01",
                "description": "Chassis Mount Bracket",
                "planned_quantity": "500.0000",
                "uom": "Nos",
                "target_period": "2026-10",
                "planning_version": "V1",
                "status": "SAVED",
                "remarks": "Batch A",
            },
            {
                "row_index": 2,
                "plant": "Plant 2",
                "customer": "Honda",
                "product_code": "PRD-H02",
                "description": "Radiator Side Plate",
                "planned_quantity": "300.0000",
                "uom": "Nos",
                "target_period": "2026-10",
                "planning_version": "V1",
                "status": "CONFIRMED",
                "remarks": "Batch B",
            },
        ],
        "deleted_ids": [],
        "reason": "Initial PRD upload",
    }

    res = client.post(f"{PRD_BASE}/save", json=payload)
    assert res.status_code == 200
    data = res.json()
    assert data["saved_count"] == 2
    assert len(data["records"]) == 2

    # Query records
    res_get = client.get(f"{PRD_BASE}/records")
    assert res_get.status_code == 200
    records = res_get.json()
    assert len(records) == 2
    assert records[0]["product_code"] == "PRD-T01"


def test_prd_bulk_delete(workspace_client):
    client, session, user = workspace_client

    # Create two records directly
    rec1 = PRDRecord(
        plant="Plant 1",
        product_code="DEL-01",
        description="Delete target 1",
        planned_quantity=10,
        uom="Nos",
        target_period="2026-10",
        created_by=user.id,
    )
    rec2 = PRDRecord(
        plant="Plant 1",
        product_code="DEL-02",
        description="Delete target 2",
        planned_quantity=20,
        uom="Nos",
        target_period="2026-10",
        created_by=user.id,
    )
    session.add_all([rec1, rec2])
    session.commit()

    # Bulk delete rec1
    del_res = client.post(
        f"{PRD_BASE}/bulk-delete",
        json={"ids": [str(rec1.id)], "reason": "Test bulk delete"},
    )
    assert del_res.status_code == 200
    assert del_res.json()["deleted_count"] == 1

    # Verify rec1 is gone and rec2 remains
    res_get = client.get(f"{PRD_BASE}/records")
    records = res_get.json()
    assert len(records) == 1
    assert records[0]["product_code"] == "DEL-02"


def test_prd_export_excel(workspace_client):
    client, session, user = workspace_client

    rec = PRDRecord(
        plant="Plant 1",
        product_code="EXP-01",
        description="Export target",
        planned_quantity=100,
        uom="Nos",
        target_period="2026-10",
        created_by=user.id,
    )
    session.add(rec)
    session.commit()

    export_res = client.get(f"{PRD_BASE}/export")
    assert export_res.status_code == 200
    assert "spreadsheetml" in export_res.headers["content-type"]

    wb = openpyxl.load_workbook(io.BytesIO(export_res.content))
    ws = wb.active
    assert ws.cell(row=1, column=2).value == "Plant"
    assert ws.cell(row=1, column=3).value == "Customer"
    assert ws.cell(row=1, column=4).value == "Product/Part No."
    assert ws.cell(row=2, column=4).value == "EXP-01"


def test_requirements_workspace(workspace_client):
    client, session, user = workspace_client

    # Get records
    res = client.get(f"{REQ_BASE}/records")
    assert res.status_code == 200
    assert isinstance(res.json(), list)

    # Recalculate
    recalc_res = client.post(f"{REQ_BASE}/recalculate")
    assert recalc_res.status_code == 200
    recalc_data = recalc_res.json()
    assert "record_count" in recalc_data
    assert "critical_shortages" in recalc_data

    # Export
    exp_res = client.get(f"{REQ_BASE}/export")
    assert exp_res.status_code == 200
    assert "spreadsheetml" in exp_res.headers["content-type"]
