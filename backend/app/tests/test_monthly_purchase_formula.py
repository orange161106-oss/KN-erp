"""KNL workbook AD4/AE4 and read-only uploaded-source integration."""
from datetime import date
from decimal import Decimal as D
from io import BytesIO

import openpyxl
import pytest
from sqlalchemy import event, func, select

from app.domain.purchase_engine.monthly import RULE_VERSION, calculate_monthly_purchase
from app.models.audit import AuditLog
from app.models.purchase_approval import PurchaseApproval
from app.models.purchase_order import PurchaseDemandEvidence
from app.models.stock_statements import StockStatementBatch, StockStatementLine, StockParameterBatch, StockParameter
from app.services.purchase_workspace import list_purchase_workspace, export_purchase_workspace_excel
from app.services.monthly_purchase import MonthlyPurchaseInputs
from app.tests.test_permission_foundation import db
from app.tests.test_central_basic_planning import catalog, calculated, approved


@pytest.mark.parametrize('required,stock,msl,moq,expected', [
    ('2550', '235', '981', '675', '3296'),  # CO2 example in the supplied workbook.
    ('5', '4', '2', '10', '10'),           # Positive gap below MOQ.
    ('5', '7', '2', '10', '0'),           # MOQ cannot create demand at exact coverage.
    ('5', '100', '2', '10', '0'),
    ('5', '-3', '2', '0', '10'),          # Published negative closing stock is retained.
    ('1.0001', '.3001', '.2500', '.9', '.9500'),  # No whole-piece or pack rounding.
])
def test_company_monthly_formula(required, stock, msl, moq, expected):
    result = calculate_monthly_purchase(required=D(required), closing_stock=D(stock), msl=D(msl), moq=D(moq), rate=D('9'))
    assert result['recommended_quantity'] == D(expected)
    assert result['purchase_value'] == D(expected) * D('9')
    if expected == '3296':
        assert result['purchase_value'] == D('29664')


@pytest.mark.parametrize('value', [None, D('NaN'), D('-1'), D('1.00001'), D('1e14')])
def test_unknown_invalid_or_excess_precision_input_never_becomes_zero(value):
    with pytest.raises(ValueError):
        calculate_monthly_purchase(required=D('1'), closing_stock=D('0'), msl=value, moq=D('0'))


def sources(db, catalog, *, stocks=(('20', 'NOS'),), msl='10', moq='50', month=10, year=2026):
    previous = date(year - 1, 12, 1) if month == 1 else date(year, month - 1, 1)
    statement = StockStatementBatch(statement_month=previous.month, statement_year=previous.year,
        statement_date=previous, source_filename='previous-stock.xlsx', file_hash='s' * 64, status='active')
    parameters = StockParameterBatch(month=month, year=year, effective_date=date(year, month, 1),
        source_filename='selected-limits.xlsx', file_hash='p' * 64, status='active')
    db.add_all([statement, parameters]); db.flush()
    for index, (quantity, unit) in enumerate(stocks, 1):
        db.add(StockStatementLine(batch_id=statement.id, item_id=catalog[5].code,
            consumable_id=catalog[5].id, source_row=index, unit=unit,
            closing_qty=D(quantity) if quantity is not None else None))
    parameter = StockParameter(batch_id=parameters.id, item_id=catalog[5].code, consumable_id=catalog[5].id,
        min_stock_level=D(msl) if msl is not None else None,
        min_order_qty=D(moq) if moq is not None else None, source_row=1)
    db.add(parameter)
    catalog[5].rate = D('9')
    catalog[5].supplier_id = catalog[6].id
    db.commit()
    return statement, parameters, parameter


def show(db, **kwargs):
    return list_purchase_workspace(db, month=10, year=2026, **kwargs)


def test_shared_stock_applied_once_after_product_demand_and_export_uses_same_rule(db, catalog):
    calculated(db, catalog)
    statement, parameters, param = sources(db, catalog, stocks=(('20','NOS'), ('5','PCS')))
    audits = db.scalar(select(func.count(AuditLog.id)))
    queries = []
    def capture(*args): queries.append(args[2])
    event.listen(db.bind, 'before_cursor_execute', capture)
    try:
        result = show(db)
    finally:
        event.remove(db.bind, 'before_cursor_execute', capture)
    row = result.records[0]
    assert len(queries) <= 14  # Batched source/master reads, independent of product count.
    assert row.required_qty == 230 and row.source_row_count == 2
    assert row.stock_qty == 25 and row.shortage_qty == 215 and row.system_recommended_qty == 215
    assert row.minimum_stock_level == 10 and row.minimum_order_qty == 50 and row.purchase_value == 1935
    assert row.stock_period == '2026-09' and row.calculation_basis == RULE_VERSION
    assert row.calculation_sources['parameter_id'] == str(param.id)
    assert row.calculation_sources['stock_batch_id'] == str(statement.id)
    assert row.supplier_id == catalog[6].id and row.po_pending_qty is None
    assert row.status == 'REQUIREMENT_REVIEW_REQUIRED'  # Preview never approves company demand.
    sheet = openpyxl.load_workbook(BytesIO(export_purchase_workspace_excel(db, month=10, year=2026))).active
    assert [sheet.cell(5, c).value for c in (6,7,9,10,19,20,21,22)] == [230,25,215,215,10,50,9,1935]
    assert sheet.cell(5, 8).value == 'Not used by monthly formula'
    assert sheet.cell(5, 23).value == '2026-09'
    assert 'Pending POs are not deducted' in sheet.cell(5,25).value
    assert db.scalar(select(func.count(PurchaseApproval.id))) == 0
    assert db.scalar(select(func.count(PurchaseDemandEvidence.id))) == 0
    assert db.scalar(select(func.count(AuditLog.id))) == audits


@pytest.mark.parametrize('stocks,msl,moq,issue', [
    ((('20','KG'),), '10','50','does not match'),
    ((('20',None),), '10','50','does not match'),
    ((('20','249'),), '10','50','does not match'),
    (((None,'PCS'),), '10','50','Closing quantity'),
    ((), '10','50','missing from the 2026-09'),
    ((('20','PCS'),), None,'50','Both MSL and MOQ'),
    ((('20','PCS'),), '10',None,'Both MSL and MOQ'),
])
def test_missing_values_or_incompatible_units_are_explicit(db, catalog, stocks, msl, moq, issue):
    calculated(db, catalog); sources(db, catalog, stocks=stocks, msl=msl, moq=moq)
    row = show(db, treat_missing_as_zero=False).records[0]
    assert row.system_recommended_qty is None and any(issue in text for text in row.input_issues)


def test_treat_missing_as_zero_calculates_with_zero_stock(db, catalog):
    """When treat_missing_as_zero is True (default), missing or blank stock is treated as 0 and calculated."""
    calculated(db, catalog)
    sources(db, catalog, stocks=(), msl='10', moq='50')
    row = show(db, treat_missing_as_zero=True).records[0]
    assert row.stock_qty == D('0')
    assert row.system_recommended_qty == D('240')  # MSL 10 + Req 230 - Stock 0 = 240
    assert 'Stock: 0 (item not in stock statement)' in row.calculation_details



def test_ambiguous_limits_and_duplicate_active_batches_cannot_be_chosen_silently(db, catalog):
    approved(db, catalog)
    statement, parameters, _ = sources(db, catalog)
    db.add(StockParameter(batch_id=parameters.id, item_id=catalog[5].code, category='Another department',
        min_stock_level=D('90'), min_order_qty=D('50'), source_row=2)); db.commit()
    row = show(db).records[0]
    assert row.status == 'MONTHLY_INPUTS_REQUIRED' and row.system_recommended_qty is None
    assert any('Ambiguous' in issue for issue in row.input_issues)
    db.add(StockStatementBatch(statement_month=9, statement_year=2026, statement_date=date(2026,9,1),
        source_filename='duplicate.xlsx', file_hash='d'*64, status='active')); db.commit()
    assert any('Multiple active stock statement' in issue for issue in show(db).records[0].input_issues)


def test_wrong_month_is_not_fallback_stock_and_january_uses_previous_december(db, catalog):
    calculated(db, catalog)
    sources(db, catalog, month=1, year=2027)
    october = show(db).records[0]
    assert october.stock_qty is october.system_recommended_qty is None
    row = MonthlyPurchaseInputs(db, month=1, year=2027, material_ids={catalog[5].id}).calculate(
        material_id=catalog[5].id, unit_id=catalog[1].id, uom='PCS', required=D('230'))
    assert row['stock_period'] == '2026-12' and row['recommended_quantity'] == 220


def test_zero_inputs_no_rate_and_stock_coverage_preserve_real_zero(db, catalog):
    approved(db, catalog)
    sources(db, catalog, stocks=(('230','PCS'),), msl='0', moq='0')
    catalog[5].rate = None; db.commit()
    row = show(db).records[0]
    assert row.system_recommended_qty == 0 and row.status == 'NO_PURCHASE_REQUIRED'
    assert row.minimum_stock_level == row.minimum_order_qty == 0
    assert row.unit_rate is row.purchase_value is None and 'Rate is missing' in row.calculation_details


def test_saved_assessment_is_never_recomputed_under_new_monthly_rule(db, catalog):
    identity = approved(db, catalog); sources(db, catalog)
    approval = PurchaseApproval(planning_version_id=identity, consumable_id=catalog[5].id, supplier_id=catalog[6].id,
        raw_calculated_qty=D('101'), system_recommended_qty=D('125'), uom='PCS', requested_by=catalog[0].id)
    db.add(approval); db.flush()
    db.add(PurchaseDemandEvidence(approval_id=approval.id, submission_key='saved-original-evidence', payload_hash='h'*64,
        submitted_by=catalog[0].id, report={'current_stock':'100','confirmed_incoming_before_receipt':'50'})); db.commit()
    row = show(db).records[0]
    assert row.calculation_basis == 'M5.1_V1' and row.status == 'SUBMITTED'
    assert row.system_recommended_qty == 125 and row.stock_qty == 100 and row.po_pending_qty == 50
    assert row.shortage_qty == 101 and row.minimum_stock_level is None
    assert db.scalar(select(func.count(PurchaseApproval.id))) == 1
