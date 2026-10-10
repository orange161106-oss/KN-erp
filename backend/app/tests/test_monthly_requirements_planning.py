import io
from decimal import Decimal
from uuid import UUID, uuid4
import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine, event
from sqlalchemy.orm import Session
from sqlalchemy.pool import StaticPool

from app.core.config import Settings
from app.db.base import Base
from app.db.session import get_db
from app.main import create_app
from app.models.auth import Role, User
from app.models.inventory_masters import Consumable, Unit
from app.models.prd import PlanningVersion
from app.models.requirements import MonthlyRequirementRecord
from app.models.rules import ConsumptionNorm
from app.security.passwords import PasswordService
from app.security.tokens import create_access_token

BASE = "/api/v1/requirements/workspace"


@pytest.fixture
def monthly_client():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def configure(connection, record):
        connection.create_function("btrim", 1, lambda value: value.strip() if value else "")
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    settings = Settings(
        _env_file=None,
        app_env="test",
        database_url="postgresql+psycopg://test_user@127.0.0.1/kn_unit_test",
        auth_secret_key="test-secret-key-for-monthly-req-testing-32b",
    )
    passwords = PasswordService()

    with Session(engine, expire_on_commit=False) as session:
        # Create Super Admin User
        admin_user = User(
            username="admin_planner",
            full_name="Planner Admin",
            password_hash=passwords.hash("password"),
            is_super_admin=True,
            requirements_read=True,
            requirements_create=True,
            requirements_update=True,
            requirements_delete=True,
        )
        # Create View-Only User
        viewer_user = User(
            username="viewer_user",
            full_name="Read Only Viewer",
            password_hash=passwords.hash("password"),
            is_super_admin=False,
            requirements_read=True,
            requirements_create=False,
            requirements_update=False,
            requirements_delete=False,
        )
        # Create Master unit & consumable & rule
        unit = Unit(code="KGS", name="Kilograms")
        session.add_all([admin_user, viewer_user, unit])
        session.flush()

        consumable = Consumable(code="WLD-01", name="Welding Wire 1.2mm", unit_id=unit.id)
        session.add(consumable)
        session.flush()

        # Add Consumption Norm for WLD-01: rule_type PRODUCTION_RATE, rate=0.05
        norm = ConsumptionNorm(
            rule_type="PRODUCTION_RATE",
            consumable_id=consumable.id,
            unit_id=unit.id,
            version=1,
            parameters={"usage_rate": 0.05, "scrap_factor": 0.0},
            rounding_policy="ROUND_HALF_UP",
            rounding_precision=2,
            is_active=True,
        )
        session.add(norm)
        session.commit()

        app = create_app(settings)

        def database():
            with Session(engine, expire_on_commit=False, autoflush=False) as req_session:
                yield req_session

        app.dependency_overrides[get_db] = database
        admin_token = create_access_token(admin_user.id, settings)
        viewer_token = create_access_token(viewer_user.id, settings)

        with TestClient(app) as client:
            client.headers.update({"Authorization": f"Bearer {admin_token}"})
            yield client, session, admin_user, viewer_user, viewer_token
    engine.dispose()


def test_empty_monthly_plan_metadata(monthly_client):
    client, session, admin_user, _, _ = monthly_client
    resp = client.get(f"{BASE}/plan-metadata?month=10&year=2026")
    assert resp.status_code == 200
    data = resp.json()
    assert data["has_plan"] is False
    assert data["planning_period"] == "2026-10"
    assert data["record_count"] == 0


def test_import_valid_excel_records(monthly_client):
    client, session, admin_user, _, _ = monthly_client

    # 1. Create a mock Excel file with the 9 KNL columns
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Oct 2026 Plan"
    ws.append(["FOR COMPONENT", "", "FOR CONSUMABLES", "", "", "", "", "", ""])
    ws.append([
        "Used For - Part Name",
        "Used Part No. (Production Order)",
        "Consumable Item ID",
        "Consumable Name",
        "Process Name",
        "Part Thickness (mm)",
        "Number of Processes",
        "Production Order for Selected Month",
        "Scheduled Consumable Quantity for Selected Month",
    ])
    ws.append([
        "Chassis Bracket",
        "PRD-CH-100",
        "WLD-01",
        "Welding Wire 1.2mm",
        "MIG Welding",
        "2.5000",
        "2",
        "10000.0000",
        "500.0000",
    ])
    ws.append([
        "Battery Tray",
        "PRD-BT-200",
        "PNT-01",
        "Epoxy Primer",
        "Powder Coating",
        "1.6000",
        "1",
        "5000.0000",
        "120.0000",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    # 2. Inspect Excel
    inspect_resp = client.post(
        f"{BASE}/inspect-excel",
        files={"file": ("test_plan.xlsx", file_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"month": 10, "year": 2026},
    )
    assert inspect_resp.status_code == 200
    inspect_data = inspect_resp.json()
    assert inspect_data["valid_rows"] == 2
    assert inspect_data["error_count"] == 0
    assert inspect_data["plan_already_exists"] is False

    # 3. Import records
    import_resp = client.post(
        f"{BASE}/import",
        json={
            "planning_month": 10,
            "planning_year": 2026,
            "source_filename": "test_plan.xlsx",
            "records": inspect_data["preview_rows"],
            "overwrite": False,
        },
    )
    assert import_resp.status_code == 200
    import_data = import_resp.json()
    assert import_data["inserted_count"] == 2
    assert import_data["plan_metadata"]["has_plan"] is True
    assert import_data["plan_metadata"]["created_by_name"] == "Planner Admin"

    # 4. Retrieve records
    list_resp = client.get(f"{BASE}/records?month=10&year=2026")
    assert list_resp.status_code == 200
    rows = list_resp.json()
    assert len(rows) == 2
    assert rows[0]["part_name"] == "Chassis Bracket"
    assert rows[0]["part_number"] == "PRD-CH-100"
    assert rows[0]["consumable_code"] == "WLD-01"
    assert rows[0]["production_order_qty"] == "10000.0000"


def test_excel_validation_errors(monthly_client):
    client, session, _, _, _ = monthly_client

    # Create invalid Excel file (missing quantity, negative thickness, duplicate rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([
        "Used For - Part Name",
        "Used Part No. (Production Order)",
        "Consumable Item ID",
        "Consumable Name",
        "Process Name",
        "Part Thickness (mm)",
        "Number of Processes",
        "Production Order for Selected Month",
    ])
    # Row 1: negative quantity
    ws.append(["Frame Bar", "PRD-FB-01", "WLD-01", "Wire", "Welding", "2.0", "1", "-50"])
    # Row 2: negative thickness & zero processes
    ws.append(["Frame Bar", "PRD-FB-02", "WLD-01", "Wire", "Welding", "-1.5", "0", "100"])
    # Row 3: duplicate of Row 2
    ws.append(["Frame Bar", "PRD-FB-02", "WLD-01", "Wire", "Welding", "1.5", "1", "100"])

    buf = io.BytesIO()
    wb.save(buf)

    resp = client.post(
        f"{BASE}/inspect-excel",
        files={"file": ("invalid_plan.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"month": 10, "year": 2026},
    )
    assert resp.status_code == 200
    data = resp.json()
    assert data["error_count"] > 0
    errors = [e["error_description"] for e in data["errors"]]
    assert any("must be positive" in e for e in errors)
    assert any("cannot be negative" in e for e in errors)
    assert any("Duplicate entry" in e for e in errors)


def test_month_isolation_october_vs_november(monthly_client):
    client, session, admin_user, _, _ = monthly_client

    # Import October plan
    client.post(
        f"{BASE}/import",
        json={
            "planning_month": 10,
            "planning_year": 2026,
            "source_filename": "oct_plan.xlsx",
            "records": [
                {
                    "part_name": "Oct Part",
                    "part_number": "PO-1001",
                    "consumable_code": "CONS-001",
                    "consumable_name": "Consumable 1",
                    "process_name": "Proc 1",
                    "production_order_qty": "10000.0000",
                }
            ],
            "overwrite": False,
        },
    )

    # Import November plan
    client.post(
        f"{BASE}/import",
        json={
            "planning_month": 11,
            "planning_year": 2026,
            "source_filename": "nov_plan.xlsx",
            "records": [
                {
                    "part_name": "Nov Part",
                    "part_number": "PO-1002",
                    "consumable_code": "CONS-002",
                    "consumable_name": "Consumable 2",
                    "process_name": "Proc 2",
                    "production_order_qty": "15000.0000",
                }
            ],
            "overwrite": False,
        },
    )

    # Verify October has only October data
    oct_rows = client.get(f"{BASE}/records?month=10&year=2026").json()
    assert len(oct_rows) == 1
    assert oct_rows[0]["part_number"] == "PO-1001"
    assert oct_rows[0]["production_order_qty"] == "10000.0000"

    # Verify November has only November data
    nov_rows = client.get(f"{BASE}/records?month=11&year=2026").json()
    assert len(nov_rows) == 1
    assert nov_rows[0]["part_number"] == "PO-1002"
    assert nov_rows[0]["production_order_qty"] == "15000.0000"


def test_batch_editing_and_timestamp_updating(monthly_client):
    client, session, admin_user, _, _ = monthly_client

    # Import initial plan
    import_resp = client.post(
        f"{BASE}/import",
        json={
            "planning_month": 10,
            "planning_year": 2026,
            "source_filename": "plan.xlsx",
            "records": [
                {
                    "part_name": "Bracket",
                    "part_number": "PRD-01",
                    "consumable_code": "WLD-01",
                    "consumable_name": "Wire",
                    "process_name": "Weld",
                    "part_thickness": "1.0000",
                    "process_count": 1,
                    "production_order_qty": "100.0000",
                }
            ],
            "overwrite": False,
        },
    )
    meta_before = import_resp.json()["plan_metadata"]
    created_at_before = meta_before["created_at"]
    assert meta_before["updated_at"] is None

    # Load record ID
    records = client.get(f"{BASE}/records?month=10&year=2026").json()
    rec_id = records[0]["id"]

    # Edit permitted fields (production order qty, thickness, process count, remarks)
    edit_resp = client.put(
        f"{BASE}/records",
        json={
            "records": [
                {
                    "id": rec_id,
                    "production_order_qty": "250.0000",
                    "part_thickness": "3.2000",
                    "process_count": 4,
                    "remarks": "Modified for rush order",
                }
            ]
        },
    )
    assert edit_resp.status_code == 200
    edit_data = edit_resp.json()
    assert edit_data["updated_count"] == 1
    updated_rec = edit_data["records"][0]
    assert updated_rec["production_order_qty"] == "250.0000"
    assert updated_rec["part_thickness"] == "3.2000"
    assert updated_rec["process_count"] == 4
    assert updated_rec["remarks"] == "Modified for rush order"

    # Verify created_at was preserved and updated_at was updated
    meta_after = edit_data["plan_metadata"]
    assert meta_after["created_at"].rstrip("Z") == created_at_before.rstrip("Z")
    assert meta_after["updated_at"] is not None
    assert meta_after["updated_by_name"] == "Planner Admin"


def test_deterministic_recalculation(monthly_client):
    client, session, _, _, _ = monthly_client

    # Import plan with WLD-01 (configured with rule rate = 0.05) and unknown CONS-99 (unconfigured)
    client.post(
        f"{BASE}/import",
        json={
            "planning_month": 10,
            "planning_year": 2026,
            "source_filename": "plan.xlsx",
            "records": [
                {
                    "part_name": "Bracket",
                    "part_number": "PRD-01",
                    "consumable_code": "WLD-01",
                    "consumable_name": "Wire",
                    "process_name": "Weld",
                    "production_order_qty": "1000.0000",
                },
                {
                    "part_name": "Custom Arm",
                    "part_number": "PRD-02",
                    "consumable_code": "CONS-99",
                    "consumable_name": "Special Sealant",
                    "process_name": "Assembly",
                    "production_order_qty": "500.0000",
                },
            ],
            "overwrite": False,
        },
    )

    # Run recalculation for October 2026
    recalc_resp = client.post(f"{BASE}/recalculate?month=10&year=2026")
    assert recalc_resp.status_code == 200
    data = recalc_resp.json()
    assert data["record_count"] == 2
    assert data["plan_metadata"]["status"] == "CALCULATED"

    records = data["records"]
    # WLD-01: 1000 * 0.05 = 50.0000
    wld_rec = next(r for r in records if r["consumable_code"] == "WLD-01")
    assert wld_rec["status"] == "Critical shortage" or wld_rec["status"] == "Calculated"
    assert Decimal(wld_rec["scheduled_consumable_qty"]) == Decimal("50.0000")

    # CONS-99 has no configured norm -> flagged for review
    unconfigured = next(r for r in records if r["consumable_code"] == "CONS-99")
    assert unconfigured["status"] == "Configuration required"
    assert "No active Consumption Norm configured" in unconfigured["remarks"]


def test_permission_enforcement(monthly_client):
    client, session, _, viewer_user, viewer_token = monthly_client

    # Viewer has read-only access
    viewer_client = TestClient(client.app)
    viewer_client.headers.update({"Authorization": f"Bearer {viewer_token}"})

    # Read records -> allowed (200)
    read_resp = viewer_client.get(f"{BASE}/records?month=10&year=2026")
    assert read_resp.status_code == 200

    # Import -> forbidden (403)
    import_resp = viewer_client.post(
        f"{BASE}/import",
        json={"planning_month": 10, "planning_year": 2026, "source_filename": "x.xlsx", "records": []},
    )
    assert import_resp.status_code == 403

    # Edit -> forbidden (403)
    edit_resp = viewer_client.put(f"{BASE}/records", json={"records": []})
    assert edit_resp.status_code == 403

    # Recalculate -> forbidden (403)
    recalc_resp = viewer_client.post(f"{BASE}/recalculate?month=10&year=2026")
    assert recalc_resp.status_code == 403


def test_excel_export_with_period_and_grouped_headers(monthly_client):
    client, session, _, _, _ = monthly_client

    # Import records
    client.post(
        f"{BASE}/import",
        json={
            "planning_month": 10,
            "planning_year": 2026,
            "source_filename": "oct_plan.xlsx",
            "records": [
                {
                    "part_name": "Export Test Part",
                    "part_number": "PO-EXP-01",
                    "consumable_code": "WLD-01",
                    "consumable_name": "Wire",
                    "process_name": "Weld",
                    "part_thickness": "2.0000",
                    "process_count": 2,
                    "production_order_qty": "1200.0000",
                    "scheduled_consumable_qty": "60.0000",
                }
            ],
            "overwrite": False,
        },
    )

    # Export Excel
    export_resp = client.get(f"{BASE}/export?month=10&year=2026")
    assert export_resp.status_code == 200
    assert export_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    # Verify Excel workbook contents
    wb = openpyxl.load_workbook(io.BytesIO(export_resp.content))
    ws = wb.active
    # Title and Period
    assert "KNL CONSUMABLE REQUIREMENTS PLANNING" in str(ws.cell(row=1, column=1).value)
    assert "October 2026" in str(ws.cell(row=2, column=1).value)
    # Group headers in Row 4
    assert ws.cell(row=4, column=2).value == "FOR COMPONENT"
    assert ws.cell(row=4, column=4).value == "FOR CONSUMABLES"
    # Column headers in Row 5
    headers = [ws.cell(row=5, column=c).value for c in range(1, 16)]
    assert "Used For – Part Name" in headers
    assert "Used Part No. (Production Order)" in headers
    assert "Consumable Item ID" in headers
    assert "Production Order for Selected Month" in headers
    assert "Scheduled Consumable Quantity for Selected Month" in headers
    # Data row in Row 6
    assert ws.cell(row=6, column=2).value == "Export Test Part"
    assert ws.cell(row=6, column=3).value == "PO-EXP-01"
    assert ws.cell(row=6, column=9).value == 1200.0
