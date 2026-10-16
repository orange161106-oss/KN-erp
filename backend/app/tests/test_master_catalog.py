"""Targeted undated Product Master preview/save, count and deletion checks."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import UUID

from openpyxl import Workbook, load_workbook
import pytest
from sqlalchemy import func, select

from app.core.errors import ApplicationError
from app.models.audit import AuditLog
from app.models.master_uploads import MasterUploadBatch
from app.models.inventory_masters import Consumable, Unit
from app.schemas.master_uploads import ExistingMasterDelete, MasterCatalogSave
from app.services import master_catalog as service
from app.tests.test_master_uploads import workbook, upload
from app.tests.test_stock_statement_delete import db, identity, client_for


@pytest.mark.parametrize("category", ["consumables", "sales_parts", "assemble"])
def test_preview_does_not_save_and_explicit_save_export_delete_preserve_ids(db, category):
    content = workbook(category, rate="12")
    preview = service.preview(db, category, "synthetic.xlsx", content)
    model = service.shared.MODELS[category]
    assert preview.row_count == 1
    assert db.scalar(select(func.count()).select_from(model)) == 0
    assert db.scalar(select(func.count()).select_from(AuditLog)) == 0
    saved = service.save_upload(db, category, "synthetic.xlsx", content,
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    assert saved.row_count == 1 and saved.filename == "synthetic.xlsx"
    db.expire_all()
    assert service.get_catalog(db, category).review_token == saved.review_token
    assert db.scalar(select(func.count()).select_from(MasterUploadBatch)) == 0
    row_id = UUID(saved.records[0]["id"])
    exported = load_workbook(BytesIO(service.export_catalog(db, category)))
    assert exported.active.max_row == 2
    exported.close()
    empty = service.delete_catalog(db, category,
        ExistingMasterDelete(review_token=saved.review_token, reason="Replace workbook"), identity(db))
    assert empty.row_count == 0 and not db.get(model, row_id).is_active
    next_preview = service.preview(db, category, "replacement.xlsx", content)
    replacement = service.save_upload(db, category, "replacement.xlsx", content,
        MasterCatalogSave(review_token=next_preview.review_token), identity(db))
    assert replacement.records[0]["id"] == str(row_id)


def counted_workbook():
    book = Workbook()
    sheet = book.active
    sheet.append(["KNL workbook"])
    sheet.append(["Item ID", "Description", "Output/Unit", "Rate", "Category"])
    sheet.append(["MAT-1", "First", "PCS", "1.1234", "First category"])
    sheet.append([None, "Missing required identifier", None, None, None])
    sheet.append(["MAT-2", "Numeric output quantity", "2000", "2.0001", "Tools"])
    sheet.append(["MAT-1", "Duplicate", "PCS", "1.1234", "Last category"])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_numeric_source_units_can_be_saved_unassigned_then_completed_later(db):
    content = counted_workbook()
    preview = service.preview(db, "consumables", "synthetic.xlsx", content)
    assert preview.report == {"data_rows": 4, "identified_rows": 3, "unique_rows": 2,
        "missing_identifier_rows": [4], "duplicate_rows": [{"identifier": "MAT-1", "first_row": 3, "last_row": 6}],
        "deferred_stock_units": [{"identifier": "MAT-2", "part_no": None, "process": None, "source_value": "2000"}]}
    assert preview.validation_errors == []
    assert preview.records[0]["category"] == "Last category"
    assert db.scalar(select(func.count()).select_from(Unit)) == 0
    saved = service.save_upload(db, "consumables", "synthetic.xlsx", content,
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    material = db.get(Consumable, UUID(next(row["id"] for row in saved.records if row["item_id"] == "MAT-2")))
    assert material.unit_id is None and material.output_unit is None
    assert saved.report["deferred_stock_units"][0]["source_value"] == "2000"
    saved = service.save_changes(db, "consumables", MasterCatalogSave(review_token=saved.review_token,
        updates=[{"id": material.id, "changes": {"output_unit": "KG"}}]), identity(db))
    assert material.rate == Decimal("2.0001") and material.output_unit == "KG"
    assert db.get(Unit, material.unit_id).code == "KG"
    assert db.scalar(select(Unit).where(Unit.code == "2000")) is None
    assert not service.valid_unit("2E3")


def test_explicit_unit_column_wins_over_numeric_output_and_blank_is_not_guessed(db):
    book = Workbook()
    sheet = book.active
    sheet.append(["Item ID", "Description", "Output/Unit", "Unit", "Rate"])
    sheet.append(["MAT-1", "Explicit unit", 2000, "KG", 2])
    sheet.append(["MAT-2", "Missing reviewed unit", 1500, None, 3])
    content = BytesIO()
    book.save(content)
    book.close()
    preview = service.preview(db, "consumables", "synthetic.xlsx", content.getvalue())
    assert preview.records[0]["output_unit"] == "KG"
    assert preview.records[1]["output_unit"] is None
    assert preview.validation_errors == []


def test_existing_monthly_upload_is_visible_and_delete_rolls_back_with_history(db):
    previous = upload(db)
    catalog = service.get_catalog(db, "consumables")
    assert catalog.row_count == 1 and catalog.filename == "synthetic.xlsx"
    payload = ExistingMasterDelete(review_token=catalog.review_token, reason="Replace old upload")
    def fail():
        db.flush()
        raise RuntimeError("synthetic persistence failure")
    with patch.object(db, "commit", side_effect=fail), pytest.raises(RuntimeError):
        service.delete_catalog(db, "consumables", payload, identity(db))
    assert service.get_catalog(db, "consumables").row_count == 1
    assert db.get(MasterUploadBatch, previous.batch_id).status == "active"
    service.delete_catalog(db, "consumables", payload, identity(db))
    assert service.get_catalog(db, "consumables").row_count == 0
    assert db.get(MasterUploadBatch, previous.batch_id).status == "deleted"
    assert db.get(MasterUploadBatch, previous.batch_id).records


def test_saved_catalog_edits_use_permissions_tokens_and_protect_identifiers(db):
    previous = upload(db)
    catalog = service.get_catalog(db, "consumables")
    row_id = previous.records[0]["id"]
    payload = MasterCatalogSave(review_token=catalog.review_token,
        updates=[{"id": row_id, "changes": {"description": "Reviewed", "rate": "2.3456"}}])
    saved = service.save_changes(db, "consumables", payload, identity(db))
    assert saved.records[0]["description"] == "Reviewed"
    assert db.get(Consumable, UUID(row_id)).rate == Decimal("2.3456")
    with pytest.raises(ApplicationError) as stale:
        service.save_changes(db, "consumables", payload, identity(db))
    assert stale.value.status_code == 409
    protected = MasterCatalogSave(review_token=saved.review_token, updates=[{"id": row_id, "changes": {"item_id": "RENAME"}}])
    with pytest.raises(ApplicationError) as invalid:
        service.save_changes(db, "consumables", protected, identity(db))
    assert invalid.value.status_code == 422
    endpoint = "/api/v1/masters/product-master/consumables/workspace"
    with client_for(db, identity(db, masters_read=True)) as client:
        assert client.get(endpoint).json()["row_count"] == 1
        assert client.post(endpoint + "/preview", files={"file": ("synthetic.xlsx", workbook())}).status_code == 403
        assert client.patch(endpoint, json=payload.model_dump(mode="json")).status_code == 403
        assert client.request("DELETE", endpoint, json={"review_token": saved.review_token, "reason": "Not allowed"}).status_code == 403


def test_upload_http_preview_then_explicit_save_without_dates(db):
    endpoint = "/api/v1/masters/product-master/consumables/workspace"
    with client_for(db, identity(db, masters_create=True, masters_read=True)) as client:
        preview = client.post(endpoint + "/preview", files={"file": ("synthetic.xlsx", workbook())})
        assert preview.status_code == 200
        assert client.get(endpoint).json()["row_count"] == 0
        saved = client.post(endpoint, files={"file": ("synthetic.xlsx", workbook())},
            data={"review": MasterCatalogSave(review_token=preview.json()["review_token"]).model_dump_json()})
        assert saved.status_code == 200 and saved.json()["row_count"] == 1
        assert "month" not in saved.json() and "year" not in saved.json()


def assembly_workbook(*, stock_unit=None):
    book = Workbook()
    sheet = book.active
    sheet.append(["Consumable Description ", "Part No. of saleable item", "Used for- Name of Saleable part",
                  "Used Part no (prd order)", "Process name", "Material thickness in mm / Gsm",
                  "No of process per part", "Output/  Unit", " Rate"] + (["Stock Unit"] if stock_unit else []))
    sheet.append(["Synthetic tap", "PART-1", "Synthetic assembly", "PRD-1", "Retapping", 2.5, 2, 2000, 12.3456]
                 + ([stock_unit] if stock_unit else []))
    sheet.append(["General material", None, "All welded parts", None, "Welding"])
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_assembly_workbook_display_headers_preview_all_fields_without_writing(db):
    preview = service.preview(db, "assemble", "assembly.xlsx", assembly_workbook())
    assert preview.row_count == 1
    row = preview.records[0]
    assert {key: value for key, value in row.items() if key != "id"} == {
        "consumable_item": "SYNTHETIC TAP", "part_no_saleable_item": "PART-1",
        "used_for_name_of_saleable_part": "Synthetic assembly", "used_part_no_prd_order": "PRD-1",
        "process_name": "Retapping", "material_thickness_in_mm_gsm": "2.5",
        "no_of_process_per_part": "2", "output_unit": None, "rate": "12.3456"}
    assert preview.report["missing_identifier_rows"] == [3]
    assert preview.report["identified_rows"] == 1
    assert preview.validation_errors == []
    assert preview.report["deferred_stock_units"][0]["source_value"] == "2000"
    assert db.scalar(select(func.count()).select_from(service.shared.MODELS["assemble"])) == 0
    assert db.scalar(select(func.count()).select_from(AuditLog)) == 0
    assert db.scalar(select(func.count()).select_from(Unit)) == 0


def test_assembly_display_headers_http_review_and_save_with_explicit_stock_unit(db):
    endpoint = "/api/v1/masters/product-master/assemble/workspace"
    content = assembly_workbook(stock_unit="KG")
    with client_for(db, identity(db, masters_create=True, masters_read=True)) as client:
        preview = client.post(endpoint + "/preview", files={"file": ("assembly.xlsx", content)})
        assert preview.status_code == 200
        assert preview.json()["records"][0]["output_unit"] == "KG"
        assert preview.json()["validation_errors"] == []
        assert client.get(endpoint).json()["row_count"] == 0
        saved = client.post(endpoint, files={"file": ("assembly.xlsx", content)},
            data={"review": MasterCatalogSave(review_token=preview.json()["review_token"]).model_dump_json()})
        assert saved.status_code == 200
        row = saved.json()["records"][0]
        assert row["used_part_no_prd_order"] == "PRD-1" and row["process_name"] == "RETAPPING"
        assert Decimal(row["no_of_process_per_part"]) == 2
        assert Decimal(row["rate"]) == Decimal("12.3456")
        assert row["output_unit"] == "KG"
        assert db.scalar(select(Unit).where(Unit.code == "2000")) is None


def test_assembly_can_save_without_unit_then_add_it_with_update_permission(db):
    endpoint = "/api/v1/masters/product-master/assemble/workspace"
    content = assembly_workbook()
    with client_for(db, identity(db, masters_create=True, masters_read=True)) as client:
        preview = client.post(endpoint + "/preview", files={"file": ("assembly.xlsx", content)}).json()
        response = client.post(endpoint, files={"file": ("assembly.xlsx", content)},
            data={"review": MasterCatalogSave(review_token=preview["review_token"]).model_dump_json()})
        assert response.status_code == 200
        saved = response.json()
        row = saved["records"][0]
        assert row["output_unit"] is None
        mapping = db.get(service.shared.MODELS["assemble"], UUID(row["id"]))
        material = db.get(Consumable, mapping.consumable_id)
        assert material.unit_id is None
        payload = {"review_token": saved["review_token"], "updates": [{"id": row["id"], "changes": {"output_unit": "PCS"}}]}
        assert client.patch(endpoint, json=payload).status_code == 403
    with client_for(db, identity(db, masters_update=True, masters_read=True)) as client:
        result = client.patch(endpoint, json=payload)
        assert result.status_code == 200
        assert result.json()["records"][0]["output_unit"] == "PCS"
        assert db.get(Unit, material.unit_id).code == "PCS"
    assert db.scalar(select(Unit).where(Unit.code == "2000")) is None


def test_unit_change_protects_stock_history_and_invalid_numeric_edits(db):
    content = workbook()
    preview = service.preview(db, "consumables", "synthetic.xlsx", content)
    saved = service.save_upload(db, "consumables", "synthetic.xlsx", content,
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    row = saved.records[0]
    def payload(value):
        return MasterCatalogSave(review_token=saved.review_token, updates=[{"id": row["id"], "changes": {"output_unit": value}}])
    with patch.object(service, "has_stock_records", return_value=True), pytest.raises(ApplicationError) as blocked:
        service.save_changes(db, "consumables", payload("KG"), identity(db))
    assert blocked.value.code == "UNIT_CHANGE_CONFLICT"
    assert db.get(Consumable, UUID(row["id"])).output_unit == "PCS"
    with pytest.raises(ApplicationError) as invalid:
        service.save_changes(db, "consumables", payload("2000"), identity(db))
    assert invalid.value.code == "INVALID_MASTER_UNIT"
    assert db.scalar(select(Unit).where(Unit.code == "2000")) is None


def test_unassigned_unit_blocks_projection_instead_of_assuming_pcs(db):
    from types import SimpleNamespace
    from app.services import projection
    with patch.object(projection.repository, "material", return_value=SimpleNamespace(unit_id=None)), \
            patch.object(projection.repository, "version", return_value=object()), pytest.raises(ApplicationError) as blocked:
        projection.context(db, "synthetic-material", "synthetic-version")
    assert blocked.value.code == "STOCK_UNIT_REQUIRED"


def test_distinct_assembly_processes_and_later_upload_unit_are_preserved(db):
    book = Workbook()
    sheet = book.active
    sheet.append(["Consumable Description", "Part No. of saleable item", "Process name", "Stock Unit"])
    sheet.append(["Synthetic tool", "PART-1", "Drilling", None])
    sheet.append(["Synthetic tool", "PART-1", "Retapping", None])
    output = BytesIO()
    book.save(output)
    content = output.getvalue()
    preview = service.preview(db, "assemble", "synthetic.xlsx", content)
    saved = service.save_upload(db, "assemble", "synthetic.xlsx", content,
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    assert saved.row_count == 2
    assert {row["process_name"] for row in saved.records} == {"DRILLING", "RETAPPING"}
    original_ids = {row["id"] for row in saved.records}
    material = db.scalar(select(Consumable))
    assert material.unit_id is None
    service.delete_catalog(db, "assemble", ExistingMasterDelete(review_token=saved.review_token, reason="Review units"), identity(db))
    sheet.cell(2, 4, "PCS")
    sheet.cell(3, 4, "PCS")
    replacement = BytesIO()
    book.save(replacement)
    book.close()
    preview = service.preview(db, "assemble", "synthetic.xlsx", replacement.getvalue())
    saved = service.save_upload(db, "assemble", "synthetic.xlsx", replacement.getvalue(),
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    assert {row["id"] for row in saved.records} == original_ids
    assert db.get(Unit, material.unit_id).code == "PCS"


@pytest.mark.parametrize("category", ["consumables", "assemble"])
def test_partial_inline_update_preserves_unknown_units_and_validates_later_units(db, category):
    content = assembly_workbook() if category == "assemble" else counted_workbook()
    preview = service.preview(db, category, "synthetic.xlsx", content)
    saved = service.save_upload(db, category, "synthetic.xlsx", content,
        MasterCatalogSave(review_token=preview.review_token), identity(db))
    pending = next(row for row in saved.records if row["output_unit"] is None)
    model = service.shared.MODELS[category]
    record = db.get(model, UUID(pending["id"]))
    actor = identity(db)
    update = service.shared.legacy.update_product_master_item
    update(db, category, record.id, {"rate": "7.1234"}, actor.id)
    assert record.rate == Decimal("7.1234") and record.output_unit is None
    material = record if category == "consumables" else db.get(Consumable, record.consumable_id)
    assert material.unit_id is None
    update(db, category, record.id, {"output_unit": "KG"}, actor.id)
    assert record.output_unit == "KG" and db.get(Unit, material.unit_id).code == "KG"
    update(db, category, record.id, {"output_unit": None}, actor.id)
    assert record.output_unit is None
    with pytest.raises(ApplicationError) as invalid:
        update(db, category, record.id, {"output_unit": "2000"}, actor.id)
    assert invalid.value.code == "INVALID_MASTER_UNIT"
    assert db.scalar(select(Unit).where(Unit.code == "2000")) is None
    with pytest.raises(ApplicationError) as precision:
        update(db, category, record.id, {"rate": "7.12345"}, actor.id)
    assert precision.value.code == "INVALID_MASTER_VALUE"
