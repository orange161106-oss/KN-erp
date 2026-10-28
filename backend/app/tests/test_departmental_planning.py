"""Automated test suite verifying departmental and maintenance consumable planning (Tool Room, PMD, NPD, HRD, QAD, Accounts, Sales)."""
from datetime import date, datetime, timezone
from decimal import Decimal
import io
from uuid import uuid4

import openpyxl
import pytest
from sqlalchemy import event, select

from app.models.auth import User
from app.models.inventory_masters import Consumable, Unit
from app.models.prd import PlanningVersion
from app.models.production import Plant
from app.models.requirements import MonthlyRequirementRecord
from app.models.stock_statements import StockStatementBatch, StockStatementLine
from app.services.purchase_workspace import (
    export_purchase_workspace_excel,
    list_purchase_workspace,
)
from app.services.requirements_workspace import (
    batch_update_requirements,
    bulk_delete_requirement_records,
    calculate_workspace_requirements,
    create_departmental_row,
    ensure_operational_units,
    get_monthly_plan_metadata,
    import_monthly_requirements,
    inspect_excel_file,
    recalculate_monthly_requirements,
)
from app.schemas.requirements_workspace import (
    DepartmentalRowCreateRequest,
    RequirementBatchUpdateRequest,
    RequirementImportRequest,
    RequirementUpdateItem,
)
from app.tests.test_permission_foundation import db, actor
from app.tests.test_central_basic_planning import catalog


def test_operational_units_seeding_is_idempotent(db):
    """Department cost centers are seeded as valid operational units in plants table without schema migration."""
    ensure_operational_units(db)
    plants = db.scalars(select(Plant.name)).all()
    plants_upper = {p.upper() for p in plants}

    for expected in ["TOOL ROOM", "PMD", "NPD", "HRD", "QAD", "ACCOUNTS", "SALES"]:
        assert expected in plants_upper

    count_before = len(plants)
    ensure_operational_units(db)
    count_after = len(db.scalars(select(Plant.name)).all())
    assert count_before == count_after


def test_departmental_sheet_ingestion_and_norm_exemption(db, catalog):
    """Departmental sheets parse dynamically and are exempt from finished-good norms."""
    user = catalog[0]

    # 1. Build an in-memory workbook with Tool Room and PMD tabs
    wb = openpyxl.Workbook()
    # Tool Room Sheet
    ws_tr = wb.active
    ws_tr.title = "TOOL ROOM"
    ws_tr.cell(1, 2, "Dept: Tool Room")
    ws_tr.cell(1, 6, "Oct'26")
    ws_tr.cell(2, 2, "S.No")
    ws_tr.cell(2, 3, "Item Id.")
    ws_tr.cell(2, 4, "Description")
    ws_tr.cell(2, 5, "Unit")
    ws_tr.cell(2, 6, "Qty")

    ws_tr.cell(3, 1, "TOOL ROOM")
    ws_tr.cell(3, 2, 1)
    ws_tr.cell(3, 3, "BBG000150QC")
    ws_tr.cell(3, 4, "Anabond 412")
    ws_tr.cell(3, 5, "NOS")
    ws_tr.cell(3, 6, 5)

    ws_tr.cell(4, 1, "TOOL ROOM")
    ws_tr.cell(4, 2, 2)
    ws_tr.cell(4, 3, "GAQ000023QC")
    ws_tr.cell(4, 4, "Cotton Waste - Big size")
    ws_tr.cell(4, 5, "KGS")
    ws_tr.cell(4, 6, 20)

    # PMD Sheet
    ws_pmd = wb.create_sheet(title="PMD")
    ws_pmd.cell(1, 3, "PMD & TOOL MAINTENANCE")
    ws_pmd.cell(1, 6, "Oct'26")
    ws_pmd.cell(2, 1, "To Location")
    ws_pmd.cell(2, 2, "Item Id.")
    ws_pmd.cell(2, 3, "Item Description")
    ws_pmd.cell(2, 4, "Part No.")
    ws_pmd.cell(2, 5, "Unit")
    ws_pmd.cell(2, 6, "Qty")

    ws_pmd.cell(3, 1, "PMD")
    ws_pmd.cell(3, 2, "ABK100003QC")
    ws_pmd.cell(3, 3, "Gloves - Knitted - 12\"")
    ws_pmd.cell(3, 5, "NOS")
    ws_pmd.cell(3, 6, 150)

    stream = io.BytesIO()
    wb.save(stream)
    file_bytes = stream.getvalue()

    # 2. Inspect excel file
    inspected = inspect_excel_file(file_bytes, "oct_dept_plan.xlsx", 10, 2026, db)
    assert inspected.error_count == 0
    assert len(inspected.preview_rows) == 3

    # Check that rows have departmental classification
    tr_row = next(r for r in inspected.preview_rows if r["consumable_code"] == "BBG000150QC")
    assert tr_row["part_name"] == "Tool Room"
    assert tr_row["part_number"] == "DEPARTMENT_INDENT"
    assert tr_row["process_name"] == "DEPARTMENT_INDENT"
    assert tr_row["rule_type"] == "PLANT_REQUEST"
    assert tr_row["scheduled_consumable_qty"] == "5.0000"

    pmd_row = next(r for r in inspected.preview_rows if r["consumable_code"] == "ABK100003QC")
    assert pmd_row["part_name"] == "PMD"
    assert pmd_row["process_name"] == "MAINTENANCE"
    assert pmd_row["scheduled_consumable_qty"] == "150.0000"

    # 3. Import requirements
    count, meta = import_monthly_requirements(
        db,
        user.id,
        RequirementImportRequest(
            planning_month=10,
            planning_year=2026,
            source_filename="oct_dept_plan.xlsx",
            records=inspected.preview_rows,
            overwrite=True,
        ),
    )
    assert count == 3

    # 4. Calculate workspace requirements
    results = calculate_workspace_requirements(db, month=10, year=2026)
    assert len(results) == 3
    for r in results:
        # Must never trigger configuration errors or missing norm errors
        assert r.status in ("Normal", "Calculated", "Low", "Critical shortage")
        assert r.rule_type == "PLANT_REQUEST"
        assert r.part_number == "DEPARTMENT_INDENT"

    # 5. Recalculate requirements: must not error or flag missing norm
    recalc = recalculate_monthly_requirements(db, user.id, 10, 2026)
    assert recalc.record_count == 3
    for r in recalc.records:
        assert r.status != "Configuration required"
        assert "Departmental indent" in (r.remarks or "")


def test_t2_stock_baseline_and_purchase_matrix_consolidation(db, catalog):
    """T-2 (August) stock baseline is sourced and purchase workspace consolidates AT:AX and AY:BF into BG."""
    user = catalog[0]

    # 1. Seed active units and consumables
    u_nos = Unit(id=uuid4(), code="NOS", name="Numbers", is_active=True)
    u_kgs = Unit(id=uuid4(), code="KGS", name="Kilograms", is_active=True)
    db.add_all([u_nos, u_kgs])
    db.flush()

    c_glove = Consumable(
        id=uuid4(),
        code="ABK100003QC",
        name="Gloves - Knitted - 12\"",
        unit_id=u_nos.id,
        is_active=True,
        latest_grn_rate=Decimal("45.5000"),
    )
    db.add(c_glove)
    db.flush()

    # 2. Seed T-2 (August 2026) Stock Statement Line for ABK100003QC
    stmt_batch = StockStatementBatch(
        id=uuid4(),
        statement_month=8,
        statement_year=2026,
        statement_date=date(2026, 8, 31),
        source_filename="august_2026_stock.xlsx",
        file_hash="mock_hash_august_2026",
        status="active",
        created_by=user.id,
    )
    db.add(stmt_batch)
    db.flush()

    stmt_line = StockStatementLine(
        id=uuid4(),
        batch_id=stmt_batch.id,
        source_row=1,
        item_id="ABK100003QC",
        item_description="Gloves - Knitted - 12\"",
        item_class="CONSUMABLE",
        unit="NOS",
        closing_qty=Decimal("80.0000"),
        closing_value=Decimal("3640.0000"),
    )
    db.add(stmt_line)
    db.flush()

    # 3. Create October 2026 Planning Version and MonthlyRequirementRecords
    # Plant 1 demands 100, Plant 2 demands 50, Tool Room demands 20, PMD demands 30
    p_version = PlanningVersion(
        id=uuid4(),
        planning_period="2026-10",
        version_number=1,
        revision_label="R1",
        source_filename="oct_2026_plan.xlsx",
        status="DRAFT",
        created_by=user.id,
        planning_month=10,
        planning_year=2026,
    )
    db.add(p_version)
    db.flush()

    rows = [
        # Production Plants (AT:AX)
        MonthlyRequirementRecord(
            planning_version_id=p_version.id,
            planning_period="2026-10",
            planning_month=10,
            planning_year=2026,
            row_index=1,
            part_name="Bracket Assy",
            part_number="BRK-001",
            consumable_code="ABK100003QC",
            consumable_name="Gloves - Knitted - 12\"",
            process_name="Welding",
            scheduled_consumable_qty=Decimal("100.0000"),
            gross_requirement=Decimal("100.0000"),
            unit="NOS",
            plant="Plant 1",
            plant_id="PLANT_1",
            status="Calculated",
        ),
        MonthlyRequirementRecord(
            planning_version_id=p_version.id,
            planning_period="2026-10",
            planning_month=10,
            planning_year=2026,
            row_index=2,
            part_name="Frame Assy",
            part_number="FRM-001",
            consumable_code="ABK100003QC",
            consumable_name="Gloves - Knitted - 12\"",
            process_name="Assembly",
            scheduled_consumable_qty=Decimal("50.0000"),
            gross_requirement=Decimal("50.0000"),
            unit="NOS",
            plant="Plant 2",
            plant_id="PLANT_2",
            status="Calculated",
        ),
        # Departmental Indents (AY:BF)
        MonthlyRequirementRecord(
            planning_version_id=p_version.id,
            planning_period="2026-10",
            planning_month=10,
            planning_year=2026,
            row_index=3,
            part_name="Tool Room",
            part_number="DEPARTMENT_INDENT",
            consumable_code="ABK100003QC",
            consumable_name="Gloves - Knitted - 12\"",
            process_name="DEPARTMENT_INDENT",
            scheduled_consumable_qty=Decimal("20.0000"),
            gross_requirement=Decimal("20.0000"),
            unit="NOS",
            plant="Tool Room",
            plant_id="TOOL_ROOM",
            status="CALCULATED",
        ),
        MonthlyRequirementRecord(
            planning_version_id=p_version.id,
            planning_period="2026-10",
            planning_month=10,
            planning_year=2026,
            row_index=4,
            part_name="PMD",
            part_number="DEPARTMENT_INDENT",
            consumable_code="ABK100003QC",
            consumable_name="Gloves - Knitted - 12\"",
            process_name="MAINTENANCE",
            scheduled_consumable_qty=Decimal("30.0000"),
            gross_requirement=Decimal("30.0000"),
            unit="NOS",
            plant="PMD",
            plant_id="PMD",
            status="CALCULATED",
        ),
    ]
    db.add_all(rows)
    db.commit()

    # 4. Check workspace requirements: opening stock should reflect August baseline (80.0000)
    reqs = calculate_workspace_requirements(db, month=10, year=2026)
    assert len(reqs) == 4
    for r in reqs:
        assert r.opening_stock == "80.0000"

    # 5. Check purchase workspace consolidation
    current_u = actor(is_super_admin=True)
    pw = list_purchase_workspace(db, month=10, year=2026, current_user=current_u)
    assert len(pw.records) == 1
    p_rec = pw.records[0]

    # Verify Plant Demands (AT:AX)
    assert p_rec.plant_demands["PLANT_1"] == Decimal("100.0000")
    assert p_rec.plant_demands["PLANT_2"] == Decimal("50.0000")
    assert p_rec.total_plant_qty == Decimal("150.0000")

    # Verify Department Demands (AY:BF)
    assert p_rec.dept_demands["TOOL_ROOM"] == Decimal("20.0000")
    assert p_rec.dept_demands["PMD"] == Decimal("30.0000")
    assert p_rec.total_dept_qty == Decimal("50.0000")

    # Verify Total Gross Requirement (BG = Plants + Departments = 150 + 50 = 200)
    assert p_rec.required_qty == Decimal("200.0000")
    # Stock Qty should be T-2 August statement stock: 80.0000
    assert p_rec.stock_qty == Decimal("80.0000")

    # 6. Verify Excel Export Matrix Structure
    excel_bytes = export_purchase_workspace_excel(db, month=10, year=2026, current_user=current_u)
    wb_export = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    assert "Purchase Recommendations" in wb_export.sheetnames
    assert "Consolidated Matrix (AT-BG)" in wb_export.sheetnames

    ws_mat = wb_export["Consolidated Matrix (AT-BG)"]
    # Check headers in row 4
    headers = [ws_mat.cell(4, c).value for c in range(1, 19)]
    assert "Plant I (AT)" in headers
    assert "Plant II (AU)" in headers
    assert "Total Plant Qty (AT:AX)" in headers
    assert "Tool Room (AY)" in headers
    assert "PMD (BA)" in headers
    assert "Total Dept Qty (AY:BE)" in headers
    assert "Total Gross Requirement (BG)" in headers

    # Check row 5 data
    assert ws_mat.cell(5, 1).value == "ABK100003QC"
    # Plant 1 (AT) is col 4
    assert ws_mat.cell(5, 4).value == Decimal("100.0000")
    # Plant 2 (AU) is col 5
    assert ws_mat.cell(5, 5).value == Decimal("50.0000")
    # Total Plant (AT:AX) is col 9
    assert ws_mat.cell(5, 9).value == Decimal("150.0000")
    # Tool Room (AY) is col 10
    assert ws_mat.cell(5, 10).value == Decimal("20.0000")
    # PMD (BA) is col 12
    assert ws_mat.cell(5, 12).value == Decimal("30.0000")
    # Total Dept is col 17
    assert ws_mat.cell(5, 17).value == Decimal("50.0000")
    # Total Gross Requirement (BG) is col 18
    assert ws_mat.cell(5, 18).value == Decimal("200.0000")


def test_departmental_non_zero_rolling_average_rule(db, catalog):
    """Seed departmental item across 5 months [10, 12, 0, 14, 12]. Verify active count is 4 and avg is 12.0000 (0 excluded)."""
    user = catalog[0]
    qtys = [Decimal("10.0000"), Decimal("12.0000"), Decimal("0.0000"), Decimal("14.0000"), Decimal("12.0000")]
    months = [5, 6, 7, 8, 9]

    for m, q in zip(months, qtys):
        p_ver = PlanningVersion(
            planning_month=m,
            planning_year=2026,
            planning_period=f"2026-{m:02d}",
            version_number=1,
            status="CONFIRMED",
            revision_label="R0",
            source_filename="test.xlsx",
            created_by=user.id,
        )
        db.add(p_ver)
        db.flush()

        rec = MonthlyRequirementRecord(
            planning_version_id=p_ver.id,
            planning_period=f"2026-{m:02d}",
            planning_month=m,
            planning_year=2026,
            plant="HRD",
            part_name="HRD",
            part_number="DOC-A4-01",
            consumable_code="HRD-A4",
            consumable_name="A4 Paper",
            process_name="DEPARTMENT_INDENT",
            unit="NOS",
            scheduled_consumable_qty=q,
            gross_requirement=q,
            stock_qty=Decimal("0.0000"),
            shortage_qty=q,
            po_pending_qty=Decimal("0.0000"),
            msl=Decimal("0.0000"),
            status="Normal",
            is_deleted=False,
            created_by=user.id,
        )
        db.add(rec)
    db.commit()

    reqs = calculate_workspace_requirements(db, month=10, year=2026)
    a4_req = next((r for r in reqs if r.consumable_code == "HRD-A4"), None)
    assert a4_req is not None
    assert a4_req.active_month_count == 4
    assert a4_req.historical_avg_qty == "12.0000"
    assert a4_req.doc_ref == "DOC-A4-01"


def test_departmental_auto_prefill_and_forward_average_update(db, catalog):
    """Nov 2026 auto-populates HRD - A4 Paper with 12.0000; updating Oct 2026 to 22.0000 updates Nov 2026 avg to 14.0000."""
    user = catalog[0]
    test_departmental_non_zero_rolling_average_rule(db, catalog)

    nov_reqs = calculate_workspace_requirements(db, month=11, year=2026)
    nov_a4 = next((r for r in nov_reqs if r.consumable_code == "HRD-A4"), None)
    assert nov_a4 is not None
    assert nov_a4.scheduled_consumable_qty == "12.0000"
    assert nov_a4.historical_avg_qty == "12.0000"
    assert nov_a4.active_month_count == 4

    oct_ver = db.scalar(select(PlanningVersion).where(PlanningVersion.planning_period == "2026-10"))
    if not oct_ver:
        oct_ver = PlanningVersion(
            planning_month=10,
            planning_year=2026,
            planning_period="2026-10",
            version_number=1,
            status="CONFIRMED",
            revision_label="R0",
            source_filename="test.xlsx",
            created_by=user.id,
        )
        db.add(oct_ver)
        db.flush()

    oct_rec = db.scalar(
        select(MonthlyRequirementRecord).where(
            MonthlyRequirementRecord.planning_period == "2026-10",
            MonthlyRequirementRecord.consumable_code == "HRD-A4",
            MonthlyRequirementRecord.is_deleted == False,
        )
    )
    if oct_rec:
        oct_rec.scheduled_consumable_qty = Decimal("22.0000")
        oct_rec.gross_requirement = Decimal("22.0000")
    else:
        oct_rec = MonthlyRequirementRecord(
            planning_version_id=oct_ver.id,
            planning_period="2026-10",
            planning_month=10,
            planning_year=2026,
            plant="HRD",
            part_name="HRD",
            part_number="DOC-A4-01",
            consumable_code="HRD-A4",
            consumable_name="A4 Paper",
            process_name="DEPARTMENT_INDENT",
            unit="NOS",
            scheduled_consumable_qty=Decimal("22.0000"),
            gross_requirement=Decimal("22.0000"),
            stock_qty=Decimal("0.0000"),
            shortage_qty=Decimal("22.0000"),
            po_pending_qty=Decimal("0.0000"),
            msl=Decimal("0.0000"),
            status="Normal",
            is_deleted=False,
            created_by=user.id,
        )
        db.add(oct_rec)
    db.commit()

    nov_reqs_updated = calculate_workspace_requirements(db, month=11, year=2026)
    nov_a4_updated = next((r for r in nov_reqs_updated if r.consumable_code == "HRD-A4"), None)
    assert nov_a4_updated is not None
    assert nov_a4_updated.active_month_count == 5
    assert nov_a4_updated.historical_avg_qty == "14.0000"
    assert nov_a4_updated.scheduled_consumable_qty == "14.0000"


def test_departmental_crud_endpoints_and_purchase_workspace_sync(db, catalog):
    """Test create_departmental_row (POST), batch_update_requirements (PATCH), and bulk_delete (POST delete-batch)."""
    user = catalog[0]

    create_req = DepartmentalRowCreateRequest(
        month=10,
        year=2026,
        plant="Tool Room",
        consumable_code="TOOL-DRILL-01",
        consumable_name="Carbide Drill Bit 6mm",
        part_number="DOC-DR-01",
        unit="NOS",
        scheduled_consumable_qty=Decimal("15.0000"),
        remarks="Special tooling indent",
    )
    created = create_departmental_row(db, user.id, create_req)
    assert created.consumable_code == "TOOL-DRILL-01"
    assert created.scheduled_consumable_qty == "15.0000"
    assert created.rule_type == "PLANT_REQUEST"
    assert created.doc_ref == "DOC-DR-01"

    update_req = RequirementBatchUpdateRequest(
        records=[
            RequirementUpdateItem(
                id=created.id,
                scheduled_consumable_qty="25.0000",
                remarks="Updated drill qty",
            )
        ]
    )
    updated = batch_update_requirements(db, user.id, update_req)
    assert updated.updated_count == 1
    updated_item = next((r for r in updated.records if r.id == created.id), None)
    assert updated_item is not None
    assert updated_item.scheduled_consumable_qty == "25.0000"

    current_u = actor(is_super_admin=True)
    pw = list_purchase_workspace(db, month=10, year=2026, current_user=current_u)
    pw_item = next((r for r in pw.records if r.consumable_code == "TOOL-DRILL-01"), None)
    assert pw_item is not None
    assert pw_item.dept_demands["TOOL_ROOM"] == Decimal("25.0000")
    assert pw_item.required_qty >= Decimal("25.0000")

    count, meta = bulk_delete_requirement_records(
        db,
        record_ids=[created.id],
        current_user_id=user.id,
        reason="Test deletion of drill bit",
        month=10,
        year=2026,
    )
    assert count == 1
    post_delete_reqs = calculate_workspace_requirements(db, month=10, year=2026)
    assert not any(r.id == created.id for r in post_delete_reqs)


def test_query_bound_invariants_for_departmental_requirements(db, catalog):
    """Verify calculate_workspace_requirements <= 4 queries and get_monthly_plan_metadata == 2 queries."""
    user = catalog[0]
    db.expunge_all()

    queries = []
    def capture(_conn, _cursor, statement, *_args):
        if statement.lstrip().upper().startswith("SELECT"):
            queries.append(statement)

    event.listen(db.get_bind(), "before_cursor_execute", capture)
    try:
        queries.clear()
        meta = get_monthly_plan_metadata(db, 10, 2026)
        assert len(queries) == 2

        queries.clear()
        records = calculate_workspace_requirements(db, month=10, year=2026)
        assert len(queries) <= 4
    finally:
        event.remove(db.get_bind(), "before_cursor_execute", capture)

