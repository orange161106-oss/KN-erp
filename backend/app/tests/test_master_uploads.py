"""Focused monthly Product Master workflow checks, isolated from company data."""
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch
from uuid import UUID

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.errors import ApplicationError
from app.models.audit import AuditLog
from app.models.inventory_masters import Consumable
from app.models.master_uploads import MasterUploadBatch
from app.schemas.master_uploads import ExistingMasterDelete, MasterPeriodDelete, MasterPeriodSave
from app.services import master_uploads as service
from app.tests.test_stock_statement_delete import db, identity, client_for


def workbook(category="consumables", description="Synthetic material", rate="12.3456"):
    book = Workbook()
    sheet = book.active
    rows = {
        "consumables": [["item_id", "description", "output_unit", "rate"], ["TEST-MAT-1", description, "PCS", rate]],
        "sales_parts": [["part_no", "item_id", "part_description", "uom"], ["PART-1", "ITEM-1", description, "PCS"]],
        "assemble": [["consumable_item", "part_no_saleable_item", "process_name", "no_of_process_per_part", "rate"],
                     ["TEST-MAT-1", "PART-1", "ASSEMBLY", "0", rate]],
    }
    for row in rows[category]:
        sheet.append(row)
    output = BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def upload(db, category="consumables", month=9, year=2026, **kwargs):
    return service.import_period(db, category, month, year, "synthetic.xlsx", workbook(category, **kwargs), identity(db))


@pytest.mark.parametrize("category", ["consumables", "sales_parts", "assemble"])
def test_upload_export_delete_are_scoped_and_preserve_canonical_master(db, category):
    selected = upload(db, category)
    other_month = upload(db, category, month=10)
    other_year = upload(db, category, year=2025)
    exported = load_workbook(BytesIO(service.export_period(db, category, 9, 2026)))
    assert exported.active.max_row == 2
    assert exported.active.cell(1, 1).value == next(iter(service.FIELDS[category]))
    exported.close()
    model = service.MODELS[category]
    canonical_id = UUID(selected.records[0]["id"])
    request = MasterPeriodDelete(batch_id=selected.batch_id, version=selected.version, reason="Wrong month")
    service.delete_period(db, category, 9, 2026, request, identity(db))
    assert service.get_period(db, category, 9, 2026).batch_id is None
    assert service.get_period(db, category, 10, 2026).batch_id == other_month.batch_id
    assert service.get_period(db, category, 9, 2025).batch_id == other_year.batch_id
    assert db.get(model, canonical_id).is_active
    assert service.delete_period(db, category, 9, 2026, request, identity(db))["replayed"]
    assert db.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "DELETE")) == 1
    assert db.get(MasterUploadBatch, selected.batch_id).records


def test_save_updates_decimal_and_canonical_master_and_rejects_stale_or_protected_edits(db):
    batch = upload(db)
    row_id = batch.records[0]["id"]
    request = MasterPeriodSave(batch_id=batch.batch_id, version=1,
        updates=[{"id": row_id, "changes": {"description": "Reviewed description", "rate": "20.0001"}}])
    saved = service.save_period(db, "consumables", 9, 2026, request, identity(db))
    assert saved.version == 2
    assert db.get(Consumable, UUID(row_id)).rate == Decimal("20.0001")
    assert db.get(Consumable, UUID(row_id)).description == "Reviewed description"
    with pytest.raises(ApplicationError, match="changed"):
        service.save_period(db, "consumables", 9, 2026, request, identity(db))
    protected = request.model_copy(update={"version": 2, "updates": [request.updates[0].model_copy(update={"changes": {"item_id": "RENAME"}})]})
    with pytest.raises(ApplicationError) as failure:
        service.save_period(db, "consumables", 9, 2026, protected, identity(db))
    assert failure.value.status_code == 422
    upload(db, month=10, description="Newer catalog description")
    with pytest.raises(ApplicationError) as stale:
        service.save_period(db, "consumables", 9, 2026, request.model_copy(update={"version": 2}), identity(db))
    assert stale.value.code == "MASTER_CHANGED"


def test_delete_is_required_before_replacement_and_failed_delete_rolls_back(db):
    original = upload(db)
    with pytest.raises(ApplicationError) as duplicate:
        upload(db, description="Must delete first")
    assert duplicate.value.code == "MASTER_PERIOD_EXISTS"
    request = MasterPeriodDelete(batch_id=original.batch_id, version=1, reason="Replace workbook")
    with patch.object(db, "commit", side_effect=RuntimeError("synthetic failure")), pytest.raises(RuntimeError):
        service.delete_period(db, "consumables", 9, 2026, request, identity(db))
    assert service.get_period(db, "consumables", 9, 2026).batch_id == original.batch_id
    old_id = UUID(original.records[0]["id"])
    assert db.get(Consumable, old_id).is_active
    service.delete_period(db, "consumables", 9, 2026, request, identity(db))
    assert not db.get(Consumable, old_id).is_active
    assert not service.get_existing(db, "consumables").records
    replacement = upload(db, description="Replacement workbook")
    assert replacement.batch_id != original.batch_id
    assert replacement.records[0]["id"] == str(old_id)
    assert db.get(Consumable, old_id).is_active
    assert db.get(Consumable, old_id).description == "Replacement workbook"


@pytest.mark.parametrize("category", ["consumables", "sales_parts", "assemble"])
def test_existing_data_can_be_reviewed_and_removed_before_dated_upload(db, category):
    # Use the established pre-period importer to represent older/manual catalog data.
    functions = {"consumables": service.legacy.bulk_upsert_consumables,
                 "sales_parts": service.legacy.bulk_upsert_sales_parts, "assemble": service.legacy.bulk_upsert_assemble}
    result = functions[category](db, workbook(category), identity(db).id)
    assert result["upserted"] == 1, result
    review = service.get_existing(db, category)
    assert review.row_count == 1
    old_id = UUID(review.records[0]["id"])
    with pytest.raises(ApplicationError) as failure:
        upload(db, category)
    assert failure.value.code == "EXISTING_MASTER_DATA"
    request = ExistingMasterDelete(review_token=review.review_token, reason="Replace older upload")
    result = service.delete_existing(db, category, request, identity(db))
    assert result["deactivated_count"] == 1
    assert not db.get(service.MODELS[category], old_id).is_active
    assert service.get_existing(db, category).row_count == 0
    assert service.delete_existing(db, category, request, identity(db))["replayed"]
    imported = upload(db, category)
    assert imported.records[0]["id"] == str(old_id)
    assert db.get(service.MODELS[category], old_id).is_active


def test_existing_cleanup_rejects_changes_and_leaves_dated_rows_untouched(db):
    dated = upload(db)
    unit = db.get(Consumable, UUID(dated.records[0]["id"])).unit_id
    old = Consumable(item_id="LEGACY-ONLY", name="Old data", unit_id=unit, output_unit="PCS")
    db.add(old)
    db.commit()
    review = service.get_existing(db, "consumables")
    assert review.row_count == 1 and review.records[0]["id"] == str(old.id)
    old.description = "Changed after review"
    db.commit()
    request = ExistingMasterDelete(review_token=review.review_token, reason="Replace old data")
    with pytest.raises(ApplicationError) as failure:
        service.delete_existing(db, "consumables", request, identity(db))
    assert failure.value.code == "MASTER_CHANGED"
    request.review_token = service.get_existing(db, "consumables").review_token
    service.delete_existing(db, "consumables", request, identity(db))
    assert db.get(Consumable, UUID(dated.records[0]["id"])).is_active
    assert service.get_period(db, "consumables", 9, 2026).batch_id == dated.batch_id


def test_http_permissions_and_delete_route_match_period_not_item_id(db):
    batch = upload(db)
    endpoint = "/api/v1/masters/product-master/consumables/period?month=9&year=2026"
    body = {"batch_id": str(batch.batch_id), "version": 1, "reason": "Wrong period"}
    with client_for(db, identity(db, masters_read=True)) as client:
        assert client.get(endpoint).status_code == 200
        assert client.post(endpoint, files={"file": ("synthetic.xlsx", workbook())}).status_code == 403
        assert client.patch(endpoint, json={"batch_id": str(batch.batch_id), "version": 1,
            "updates": [{"id": batch.records[0]["id"], "changes": {"rate": "1"}}]}).status_code == 403
        assert client.request("DELETE", endpoint, json=body).status_code == 403
        existing = client.get("/api/v1/masters/product-master/consumables/existing-data")
        assert existing.status_code == 200
        assert client.request("DELETE", "/api/v1/masters/product-master/consumables/existing-data",
            json={"review_token": existing.json()["review_token"], "reason": "Not allowed"}).status_code == 403
    with client_for(db, identity(db, masters_delete=True)) as client:
        assert client.request("DELETE", endpoint.replace("month=9", "month=10"), json=body).status_code == 404
        assert client.request("DELETE", endpoint, json=body).status_code == 200
