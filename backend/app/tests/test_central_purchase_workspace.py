"""Central demand must remain visible without becoming an unreviewed purchase order."""
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import openpyxl
import pytest
from sqlalchemy import func, select

from app.core.errors import ApplicationError
from app.models.purchase_approval import PurchaseApproval
from app.models.purchase_order import PurchaseDemandEvidence
from app.models.requirements import MonthlyRequirementRecord
from app.services.purchase_workspace import list_purchase_workspace, submit_purchase_workspace_records, export_purchase_workspace_excel
from app.tests.test_permission_foundation import db, actor
from app.tests.test_central_basic_planning import catalog, calculated, approved, imported, rules


def show(db, user=None):
    return list_purchase_workspace(db, month=10, year=2026, current_user=user)


def test_saved_central_demand_appears_without_monthly_copies_or_guessed_purchase_quantities(db, catalog):
    identity, _ = calculated(db, catalog)
    result = show(db)
    assert len(result.records) == 1
    row = result.records[0]
    assert row.plan_id == str(identity) and row.consumable_code == 'DRILL-A'
    assert row.required_qty == Decimal('230') and row.source_row_count == 2
    assert row.status == 'REQUIREMENT_REVIEW_REQUIRED'
    assert row.stock_qty is row.po_pending_qty is row.shortage_qty is row.system_recommended_qty is None
    assert row.supplier_id is None  # No arbitrary supplier selection for central plans.
    assert 'approve the central plan' in result.notice
    exported = openpyxl.load_workbook(BytesIO(export_purchase_workspace_excel(db, month=10, year=2026))).active
    assert exported.cell(5, 7).value == 'Unavailable'
    assert 'No stock statement for 2026-09' in exported.cell(5, 7).comment.text
    assert 'menu beside Recalculate' in exported.cell(3, 1).value
    assert db.scalar(select(func.count(MonthlyRequirementRecord.id))) == 0
    assert db.scalar(select(func.count(PurchaseApproval.id))) == 0


def test_configuration_issues_do_not_hide_valid_demand_or_create_zero_demand(db, catalog):
    norm = db.scalar(select(rules.ConsumptionNorm).where(rules.ConsumptionNorm.product_id == catalog[4].id))
    norm.is_active = False
    db.commit()
    calculated(db, catalog)
    result = show(db)
    assert result.records[0].required_qty == Decimal('200')
    assert '1 requirement configuration issues' in result.notice
    assert result.records[0].system_recommended_qty is None


def test_newer_uncalculated_revision_does_not_reuse_older_purchase_demand(db, catalog):
    calculated(db, catalog)
    latest, _, _ = imported(db, catalog, revision='R4')
    result = show(db)
    assert result.plan_metadata.plan_id == str(latest)
    assert not result.records
    assert 'no saved consumable requirements' in result.notice


def test_approved_central_plan_needs_assessment_and_cannot_use_legacy_submit(db, catalog):
    approved(db, catalog)
    user = actor(is_super_admin=True)
    result = show(db, user)
    assert result.records[0].status == 'MONTHLY_INPUTS_REQUIRED'
    assert result.records[0].system_recommended_qty is None
    with pytest.raises(ApplicationError) as error:
        submit_purchase_workspace_records(db, month=10, year=2026, record_ids=[result.records[0].id], reason='Synthetic', current_user=user)
    assert error.value.code == 'CENTRAL_PURCHASE_ASSESSMENT_REQUIRED'
    assert db.scalar(select(func.count(PurchaseApproval.id))) == 0


def test_traced_assessment_quantities_are_visible_and_export_explains_missing_values(db, catalog):
    identity = approved(db, catalog)
    before = openpyxl.load_workbook(BytesIO(export_purchase_workspace_excel(db, month=10, year=2026))).active
    assert before.cell(5, 6).value == 230
    assert all(before.cell(5, column).value == 'Unavailable' for column in (7, 9, 10))
    assert before.cell(5, 8).value == 'Not used by monthly formula'
    assert 'No stock statement for 2026-09' in before.cell(5, 10).comment.text
    assert before.cell(5, 16).value == 'Not entered'
    assert before.cell(5, 3).value == 'Not selected'
    assert before.cell(5, 12).value == 'Not submitted'
    assert before.cell(4, 7).value == 'Stock Qty (see calculation basis)'
    assert before.cell(4, 9).value == 'Raw Purchase Need'
    assert before.freeze_panes == 'C5' and before.auto_filter.ref == 'A4:Z5'
    assert before['A3'].alignment.wrap_text
    assert 'Central requirements are approved' in before.cell(3, 1).value
    approval = PurchaseApproval(planning_version_id=identity, consumable_id=catalog[5].id, supplier_id=catalog[6].id,
        raw_calculated_qty=Decimal('100'), system_recommended_qty=Decimal('125'), uom='PCS', requested_by=catalog[0].id)
    db.add(approval)
    db.flush()
    db.add(PurchaseDemandEvidence(approval_id=approval.id, submission_key='synthetic-display', payload_hash='synthetic',
        submitted_by=catalog[0].id, report={'current_stock':'100', 'confirmed_incoming_before_receipt':'50'}))
    db.commit()
    row = show(db).records[0]
    assert row.status == 'SUBMITTED' and row.approval_id == approval.id
    assert row.required_qty == Decimal('230') and row.stock_qty == Decimal('100')
    assert row.po_pending_qty == Decimal('50') and row.shortage_qty == Decimal('100')
    assert row.system_recommended_qty == Decimal('125')
    assert row.supplier_id == catalog[6].id
    after = openpyxl.load_workbook(BytesIO(export_purchase_workspace_excel(db, month=10, year=2026))).active
    assert [after.cell(5, col).value for col in (6, 7, 8, 9, 10)] == [230, 100, 50, 100, 125]
    assert after.cell(5, 7).comment is None
    approval.raw_calculated_qty = approval.system_recommended_qty = Decimal('0')
    db.scalar(select(PurchaseDemandEvidence)).report = {'current_stock':'0', 'confirmed_incoming_before_receipt':'0'}
    db.commit()
    zero = openpyxl.load_workbook(BytesIO(export_purchase_workspace_excel(db, month=10, year=2026))).active
    assert all(zero.cell(5, col).value == 0 for col in (7, 8, 9, 10))


def test_central_purchase_scope_and_filters_are_enforced(db, catalog):
    calculated(db, catalog)
    user = actor(plant_ids=[uuid4()])
    with pytest.raises(ApplicationError) as error:
        show(db, user)
    assert error.value.status_code == 403
    with pytest.raises(ApplicationError):
        export_purchase_workspace_excel(db, month=10, year=2026, current_user=user)
    result = list_purchase_workspace(db, month=10, year=2026, search='DRILL-A', status_filter='REQUIREMENT_REVIEW_REQUIRED')
    assert len(result.records) == 1
    result = list_purchase_workspace(db, month=10, year=2026, status_filter='READY_FOR_APPROVAL')
    assert not result.records and result.notice
