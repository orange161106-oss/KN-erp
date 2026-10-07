from functools import wraps
from decimal import Decimal, InvalidOperation
import io
from uuid import UUID, uuid4
from datetime import datetime, timezone

from sqlalchemy import select, or_, and_, func, update
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.core.errors import ApplicationError
from app.domain.purchase_engine.orders import pending_quantity
from app.models.grn import GRN, GRNItem, GoodsReceiptRecord
from app.models.inventory import StockSnapshot, StockTransaction
from app.repositories import grn as repository, purchase_order as orders
from app.schemas.grn import (
    GRNResponse,
    ReceiptLineResponse,
    GoodsReceiptRecordResponse,
    WorkspaceSaveRequest,
    WorkspaceSaveResponse,
    ExcelInspectResponse,
    ExcelSheetInspectInfo,
    ExcelImportSheetResponse,
)
from app.services import inventory
from app.services.projection import digest
from app.services.purchase_order import audit


def transaction(function):
    @wraps(function)
    def wrapped(session, *args, **kwargs):
        try:
            return function(session, *args, **kwargs)
        except IntegrityError:
            session.rollback()
            raise ApplicationError('GRN_SOURCE_CONFLICT', 'Source receipt, stock event or snapshot conflicts. Retry unchanged only for a concurrent import.', 409) from None
        except SQLAlchemyError as error:
            session.rollback()
            if getattr(getattr(error, 'orig', None), 'sqlstate', None) in ('40001', '40P01'):
                raise ApplicationError('GRN_RETRY_REQUIRED', 'A concurrent receipt changed this order. Retry the unchanged export.', 409) from None
            raise ApplicationError('DATABASE_UNAVAILABLE', 'Database is unavailable.', 503) from None
        except Exception:
            session.rollback()
            raise
    return wrapped


def response(session, row, replayed=False):
    order = orders.order(session, row.purchase_order_id)
    po_items = {item.id: item for item in orders.items(session, order.id)}
    lines = []
    for line in repository.items(session, row.id):
        material = po_items[line.purchase_order_item_id].material_snapshot
        values = {field: getattr(line, field) for field in (
            'id', 'source_line_id', 'purchase_order_item_id', 'consumable_id', 'unit_id',
            'received_quantity', 'accepted_quantity', 'rejected_quantity', 'stock_transaction_id', 'stock_snapshot_id')}
        lines.append(ReceiptLineResponse(**values, consumable_code=material['code'],
                     consumable_name=material['name'], unit_code=material['unit_code']))
    return GRNResponse(**{field: getattr(row, field) for field in (
        'id', 'source_grn_id', 'purchase_order_id', 'supplier_id', 'stock_batch_id', 'event_at',
        'source_actor', 'imported_by', 'imported_at', 'reason')},
        po_number=order.po_number, supplier_name=order.supplier_snapshot['name'], items=lines, replayed=replayed)


@transaction
def import_grn(session, data, actor, *, enabled):
    if not enabled:
        raise ApplicationError('INVENTORY_IMPORT_DISABLED', 'Source import is disabled until the ERP mapping is verified.', 409)
    normalized = data.model_dump(mode='json')
    normalized['items'].sort(key=lambda row: row['source_line_id'])
    normalized['stock']['movements'].sort(key=lambda row: row['source_event_id'])
    normalized['stock']['snapshots'].sort(key=lambda row: row['source_snapshot_id'])
    payload_hash = digest(normalized)
    # Lock the parent before reading totals. Every receipt for this PO uses this
    # lock; stock stage then locks materials in UUID order. SERIALIZABLE also
    # protects against stale snapshots after waiting for the lock.
    order = orders.order(session, data.purchase_order_id, lock=True)
    if order is None:
        raise ApplicationError('PO_NOT_FOUND', 'Purchase order does not exist.', 404)
    previous = repository.by_source(session, data.source_grn_id)
    if previous:
        if previous.payload_hash != payload_hash:
            raise ApplicationError('GRN_KEY_CONFLICT', 'This source GRN already has different content. Imported history cannot be replaced.', 409)
        return response(session, previous, True)
    if order.status != 'ISSUED':
        raise ApplicationError('GRN_PO_STATE', 'Only issued purchase orders can receive imported GRNs.', 409)
    if order.supplier_id != data.supplier_id:
        raise ApplicationError('GRN_SUPPLIER_MISMATCH', 'Source supplier must match the purchase order.', 409)
    if data.event_at.date() < order.po_date:
        raise ApplicationError('GRN_EVENT_TIME', 'Receipt cannot precede the purchase order date.', 422)
    lines = {row.id: row for row in orders.items(session, order.id)}
    for line in data.items:
        item = lines.get(line.purchase_order_item_id)
        if item is None or item.consumable_id != line.consumable_id or item.unit_id != line.unit_id:
            raise ApplicationError('GRN_ITEM_MISMATCH', 'Receipt item, material and stock unit must match this PO.', 409)
        _, accepted, _ = repository.totals(session, item.id)
        # Conservative hold: an import exceeding outstanding physical quantity
        # requires an approved over-receipt policy, even if some is rejected.
        if line.received_quantity > pending_quantity(item.ordered_quantity, Decimal('0'), accepted):
            raise ApplicationError('GRN_OVER_RECEIPT_TBD', 'Receipt exceeds pending quantity. Over-receipt policy requires KNL confirmation; nothing was imported.', 409)
    stock = inventory.stage_source(session, data.stock, actor, enabled=enabled)
    row = GRN(source_grn_id=data.source_grn_id, payload_hash=payload_hash, purchase_order_id=order.id,
              supplier_id=data.supplier_id, stock_batch_id=stock.id, event_at=data.event_at,
              source_actor=data.source_actor, imported_by=actor, reason=data.reason)
    session.add(row)
    session.flush()
    snapshots = {source.consumable_id: session.scalar(select(StockSnapshot).where(
        StockSnapshot.source_snapshot_id == source.source_snapshot_id)) for source in data.stock.snapshots}
    for line in data.items:
        movement = session.scalar(select(StockTransaction).where(StockTransaction.source_event_id == line.source_event_id)) if line.source_event_id else None
        session.add(GRNItem(grn_id=row.id, **line.model_dump(exclude={'source_event_id'}),
                            stock_transaction_id=movement.id if movement else None,
                            stock_snapshot_id=snapshots[line.consumable_id].id))
    session.flush()
    audit(session, actor, row.id, 'IMPORT_GRN', data.reason,
          {'source_grn_id': row.source_grn_id, 'purchase_order_id': str(order.id), 'stock_batch_id': str(stock.id),
           'payload_hash': payload_hash, 'source': 'EXISTING_ERP'}, entity='grns')
    audit(session, actor, order.id, 'IMPORT_PO_RECEIPT', data.reason,
          {'grn_id': str(row.id), 'source_grn_id': row.source_grn_id, 'fulfilment_basis': 'ACCEPTED_USABLE_ONLY'})
    session.flush()
    result = response(session, row)
    session.commit()
    return result


@transaction
def get(session, identity):
    row = session.get(GRN, identity)
    if row is None:
        raise ApplicationError('GRN_NOT_FOUND', 'Imported GRN does not exist.', 404)
    return response(session, row)


@transaction
def listing(session, limit, offset, purchase_order_id=None):
    query = select(GRN)
    if purchase_order_id:
        query = query.where(GRN.purchase_order_id == purchase_order_id)
    rows = session.scalars(query.order_by(GRN.event_at.desc(), GRN.id).limit(limit).offset(offset)).all()
    return [response(session, row) for row in rows]


def to_record_response(row: GoodsReceiptRecord) -> GoodsReceiptRecordResponse:
    return GoodsReceiptRecordResponse(
        id=row.id,
        row_index=row.row_index,
        part_number=row.part_number,
        item_id=row.item_id,
        description=row.description,
        quantity=format(row.quantity, '.4f') if row.quantity is not None else '0.0000',
        unit=row.unit or 'Nos',
        po_number=row.po_number,
        supplier_name=row.supplier_name,
        status=row.status or 'SAVED',
        source_grn_id=row.source_grn_id,
        plant=row.plant,
        grn_date=row.grn_date or (row.created_at.strftime('%Y-%m-%d') if row.created_at else None),
        notes=row.notes,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


@transaction
def list_workspace_records(session, search: str | None = None, status: str | None = None, limit: int = 500, offset: int = 0) -> list[GoodsReceiptRecordResponse]:
    query = select(GoodsReceiptRecord).where(GoodsReceiptRecord.is_deleted.is_(False))

    if status:
        query = query.where(GoodsReceiptRecord.status == status.upper())

    if search:
        s = f"%{search.strip()}%"
        query = query.where(
            or_(
                GoodsReceiptRecord.part_number.ilike(s),
                GoodsReceiptRecord.item_id.ilike(s),
                GoodsReceiptRecord.description.ilike(s),
                GoodsReceiptRecord.po_number.ilike(s),
                GoodsReceiptRecord.supplier_name.ilike(s),
                GoodsReceiptRecord.unit.ilike(s),
            )
        )

    rows = session.scalars(query.order_by(GoodsReceiptRecord.row_index.asc().nulls_last(), GoodsReceiptRecord.created_at.asc()).limit(limit).offset(offset)).all()

    # If no workspace records exist yet, seed from existing GRN items if any exist
    if not rows and not search and not status and offset == 0:
        grns = session.scalars(select(GRN).order_by(GRN.event_at.desc())).all()
        idx = 1
        created_rows = []
        for grn_row in grns:
            order = orders.order(session, grn_row.purchase_order_id)
            po_items = {item.id: item for item in orders.items(session, order.id)} if order else {}
            for item in repository.items(session, grn_row.id):
                material = po_items.get(item.purchase_order_item_id)
                mat_snap = material.material_snapshot if material else {}
                rec = GoodsReceiptRecord(
                    id=uuid4(),
                    row_index=idx,
                    part_number=mat_snap.get('code', 'PART-' + str(idx)),
                    item_id=item.source_line_id,
                    description=mat_snap.get('name', 'Consumable Item'),
                    quantity=item.received_quantity,
                    unit=mat_snap.get('unit_code', 'Nos'),
                    po_number=order.po_number if order else None,
                    supplier_name=order.supplier_snapshot.get('name') if order else None,
                    status='SAVED',
                    source_grn_id=grn_row.source_grn_id,
                    created_by=grn_row.imported_by,
                )
                session.add(rec)
                created_rows.append(rec)
                idx += 1
        if created_rows:
            session.commit()
            return [to_record_response(r) for r in created_rows]

    return [to_record_response(row) for row in rows]


@transaction
def save_workspace_records(session, data: WorkspaceSaveRequest, actor: UUID) -> WorkspaceSaveResponse:
    saved_records = []

    # Process deletions
    for del_id in data.deleted_ids:
        row = session.get(GoodsReceiptRecord, del_id)
        if row and not row.is_deleted:
            row.is_deleted = True
            audit(session, actor, row.id, 'DELETE_ROW', data.reason,
                  {'id': str(row.id), 'part_number': row.part_number, 'item_id': row.item_id},
                  old={'status': row.status, 'is_deleted': False}, entity='goods_receipt_records')

    # Process additions / updates
    for item in data.records:
        try:
            qty = Decimal(str(item.quantity).strip())
            if qty < 0:
                raise ValueError("Quantity must be non-negative")
        except (InvalidOperation, ValueError):
            raise ApplicationError('INVALID_QUANTITY', f"Invalid quantity '{item.quantity}' for item {item.item_id or item.part_number}.", 422)

        existing = session.get(GoodsReceiptRecord, item.id) if item.id else None
        if existing and not existing.is_deleted:
            old_vals = {
                'part_number': existing.part_number,
                'item_id': existing.item_id,
                'description': existing.description,
                'quantity': str(existing.quantity),
                'unit': existing.unit,
                'status': existing.status,
            }
            existing.part_number = item.part_number.strip()
            existing.item_id = item.item_id.strip()
            existing.description = item.description.strip()
            existing.quantity = qty
            existing.unit = item.unit.strip() or 'Nos'
            existing.po_number = item.po_number.strip() if item.po_number else None
            existing.supplier_name = item.supplier_name.strip() if item.supplier_name else None
            if item.source_grn_id:
                existing.source_grn_id = item.source_grn_id.strip()
            existing.plant = item.plant.strip() if item.plant else None
            existing.grn_date = item.grn_date.strip() if item.grn_date else None
            existing.status = item.status or 'SAVED'
            existing.notes = item.notes
            existing.row_index = item.row_index
            existing.updated_at = datetime.now(timezone.utc)
            saved_records.append(existing)
            audit(session, actor, existing.id, 'UPDATE_ROW', data.reason,
                  {'id': str(existing.id), 'part_number': item.part_number, 'quantity': str(qty)},
                  old=old_vals, entity='goods_receipt_records')
        else:
            new_row = GoodsReceiptRecord(
                id=item.id or uuid4(),
                row_index=item.row_index,
                part_number=item.part_number.strip(),
                item_id=item.item_id.strip(),
                description=item.description.strip(),
                quantity=qty,
                unit=item.unit.strip() or 'Nos',
                po_number=item.po_number.strip() if item.po_number else None,
                supplier_name=item.supplier_name.strip() if item.supplier_name else None,
                source_grn_id=item.source_grn_id.strip() if item.source_grn_id else None,
                plant=item.plant.strip() if item.plant else None,
                grn_date=item.grn_date.strip() if item.grn_date else None,
                status='SAVED',
                notes=item.notes,
                created_by=actor,
            )
            session.add(new_row)
            saved_records.append(new_row)
            audit(session, actor, new_row.id, 'CREATE_ROW', data.reason,
                  {'id': str(new_row.id), 'part_number': item.part_number, 'quantity': str(qty)},
                  entity='goods_receipt_records')

    session.commit()
    return WorkspaceSaveResponse(
        saved_count=len(saved_records),
        deleted_count=len(data.deleted_ids),
        records=[to_record_response(r) for r in saved_records]
    )


@transaction
def delete_workspace_record(session, identity: UUID, actor: UUID, reason: str = 'Row deleted') -> None:
    row = session.get(GoodsReceiptRecord, identity)
    if not row or row.is_deleted:
        raise ApplicationError('RECORD_NOT_FOUND', 'Goods receipt record not found.', 404)
    row.is_deleted = True
    audit(session, actor, row.id, 'DELETE_ROW', reason,
          {'id': str(row.id), 'part_number': row.part_number, 'item_id': row.item_id},
          old={'is_deleted': False}, entity='goods_receipt_records')
    session.commit()


@transaction
def bulk_delete_workspace_records(
    session,
    ids: list[UUID] | None = None,
    delete_all_matching: bool = False,
    search: str | None = None,
    status: str | None = None,
    actor: UUID | None = None,
    reason: str = 'Bulk delete from workspace',
) -> int:
    query = select(GoodsReceiptRecord).where(GoodsReceiptRecord.is_deleted.is_(False))
    if delete_all_matching:
        if status and status.upper() != 'ALL':
            query = query.where(GoodsReceiptRecord.status == status.upper())
        if search and search.strip():
            s = f"%{search.strip()}%"
            query = query.where(
                or_(
                    GoodsReceiptRecord.part_number.ilike(s),
                    GoodsReceiptRecord.item_id.ilike(s),
                    GoodsReceiptRecord.description.ilike(s),
                    GoodsReceiptRecord.po_number.ilike(s),
                    GoodsReceiptRecord.supplier_name.ilike(s),
                    GoodsReceiptRecord.unit.ilike(s),
                    GoodsReceiptRecord.plant.ilike(s),
                    GoodsReceiptRecord.source_grn_id.ilike(s),
                )
            )
    elif ids:
        query = query.where(GoodsReceiptRecord.id.in_(ids))
    else:
        return 0

    records_to_delete = session.scalars(query).all()
    count = len(records_to_delete)
    if count == 0:
        return 0

    target_ids = [r.id for r in records_to_delete]
    session.execute(
        update(GoodsReceiptRecord)
        .where(GoodsReceiptRecord.id.in_(target_ids))
        .values(is_deleted=True, updated_at=datetime.now(timezone.utc))
    )

    if actor:
        audit(
            session, actor, uuid4(), 'BULK_DELETE_ROWS', reason,
            {'deleted_count': count, 'deleted_ids': [str(i) for i in target_ids[:50]]},
            entity='goods_receipt_records'
        )

    session.commit()
    return count


def inspect_excel_file(file_bytes: bytes, filename: str) -> ExcelInspectResponse:
    if not file_bytes:
        raise ApplicationError('EMPTY_FILE', 'The uploaded file is empty.', 422)

    # Handle CSV format
    if filename.lower().endswith('.csv'):
        import csv
        try:
            text = file_bytes.decode('utf-8', errors='replace')
            csv_rows = list(csv.reader(io.StringIO(text)))
        except Exception as exc:
            raise ApplicationError('INVALID_CSV', f'Could not read CSV file: {str(exc)}', 422)

        header_row = [str(c).strip() for c in csv_rows[0]] if csv_rows else []
        data_rows = csv_rows[1:] if len(csv_rows) > 1 else []
        sample = []
        for d_row in data_rows[:5]:
            row_dict = {}
            for idx, h in enumerate(header_row):
                val = d_row[idx] if idx < len(d_row) else ''
                row_dict[h] = str(val)
            sample.append(row_dict)

        return ExcelInspectResponse(
            filename=filename,
            sheets=[ExcelSheetInspectInfo(
                name='CSV Data',
                row_count=len(data_rows),
                column_count=len(header_row),
                headers=header_row,
                sample_rows=sample
            )]
        )

    # Handle Excel format (.xlsx)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as exc:
            raise ApplicationError('INVALID_EXCEL', f'Could not read Excel file: {str(exc)}', 422)

    sheets_info: list[ExcelSheetInspectInfo] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        # Skip Chartsheets and non-worksheet objects
        if not hasattr(sheet, 'iter_rows'):
            continue

        try:
            rows = list(sheet.iter_rows(values_only=True))
        except Exception:
            rows = []

        if not rows:
            sheets_info.append(ExcelSheetInspectInfo(
                name=sheet_name,
                row_count=0,
                column_count=0,
                headers=[],
                sample_rows=[]
            ))
            continue

        header_row = None
        data_rows = []
        for r in rows:
            if any(cell is not None and str(cell).strip() != '' for cell in r):
                if header_row is None:
                    header_row = [str(cell).strip() if cell is not None else f'Col_{i+1}' for i, cell in enumerate(r)]
                else:
                    data_rows.append(r)

        headers = header_row or []
        col_count = len(headers)
        row_count = len(data_rows)

        sample = []
        for d_row in data_rows[:5]:
            row_dict = {}
            for idx, h in enumerate(headers):
                val = d_row[idx] if idx < len(d_row) else ''
                row_dict[h] = str(val) if val is not None else ''
            sample.append(row_dict)

        sheets_info.append(ExcelSheetInspectInfo(
            name=sheet_name,
            row_count=row_count,
            column_count=col_count,
            headers=headers,
            sample_rows=sample
        ))

    if not sheets_info and getattr(workbook, 'sheetnames', None):
        for s_name in workbook.sheetnames:
            sheets_info.append(ExcelSheetInspectInfo(
                name=s_name,
                row_count=0,
                column_count=0,
                headers=[],
                sample_rows=[]
            ))

    if not sheets_info:
        raise ApplicationError('NO_DATA_SHEETS', 'Workbook contains no readable data worksheets.', 422)

    return ExcelInspectResponse(filename=filename, sheets=sheets_info)


@transaction
def import_excel_sheet(session, file_bytes: bytes, filename: str, sheet_name: str, actor: UUID, mode: str = 'APPEND', reason: str = 'Excel sheet import') -> ExcelImportSheetResponse:
    if not file_bytes:
        raise ApplicationError('EMPTY_FILE', 'The uploaded file is empty.', 422)

    # Handle CSV format
    if filename.lower().endswith('.csv'):
        import csv
        try:
            text = file_bytes.decode('utf-8', errors='replace')
            rows = list(csv.reader(io.StringIO(text)))
        except Exception as exc:
            raise ApplicationError('INVALID_CSV', f'Could not read CSV file: {str(exc)}', 422)
    else:
        try:
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception:
            try:
                workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
            except Exception as exc:
                raise ApplicationError('INVALID_EXCEL', f'Could not read Excel file: {str(exc)}', 422)

        if sheet_name not in workbook.sheetnames:
            raise ApplicationError('SHEET_NOT_FOUND', f"Sheet '{sheet_name}' not found in workbook.", 404)

        sheet = workbook[sheet_name]
        if not hasattr(sheet, 'iter_rows'):
            raise ApplicationError('INVALID_SHEET', f"Sheet '{sheet_name}' is a chart sheet and contains no tabular data.", 422)

        rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        raise ApplicationError('EMPTY_SHEET', f"Sheet '{sheet_name}' is empty.", 422)

    non_empty_rows = [r for r in rows if any(cell is not None and str(cell).strip() != '' for cell in r)]
    if not non_empty_rows:
        raise ApplicationError('EMPTY_DATA', f"Sheet '{sheet_name}' has no data rows.", 422)

    # Detect best header row from the first 10 rows
    header_idx_map: dict[str, int] = {}
    best_header_idx = None
    best_matches = 0

    for r_idx, r in enumerate(non_empty_rows[:10]):
        temp_map: dict[str, int] = {}
        for idx, cell in enumerate(r):
            if cell is None:
                continue
            key = str(cell).strip().lower().replace('_', ' ').replace('-', ' ')

            # Check specific multi-word tokens first
            if any(w in key for w in ['part description', 'item description', 'material description', 'product description', 'particulars']):
                temp_map.setdefault('description', idx)
            elif any(w in key for w in ['part number', 'part no', 'part code', 'material code', 'item code', 'product code', 'part #']):
                temp_map.setdefault('part_number', idx)
            elif any(w in key for w in ['item id', 'line id', 'line no', 'sl no', 'sl.no', 'sr no', 'sr.no', 'serial no', 'serial']):
                temp_map.setdefault('item_id', idx)
            elif any(w in key for w in ['grn date', 'receipt date', 'mrn date', 'invoice date', 'dc date', 'date']):
                temp_map.setdefault('grn_date', idx)
            elif any(w in key for w in ['grn no', 'grn number', 'grn #', 'receipt no', 'receipt number', 'mrn no', 'mrn number', 'mrn']):
                temp_map.setdefault('source_grn_id', idx)
            elif any(w in key for w in ['po number', 'po no', 'po #', 'purchase order', 'order no']):
                temp_map.setdefault('po_number', idx)
            elif any(w in key for w in ['supplier name', 'supplier', 'vendor name', 'vendor', 'party name', 'party']):
                temp_map.setdefault('supplier_name', idx)
            elif any(w in key for w in ['plant code', 'plant name', 'plant', 'factory', 'location']):
                temp_map.setdefault('plant', idx)
            elif any(w in key for w in ['received qty', 'accepted qty', 'rec qty', 'receipt qty', 'quantity', 'qty']):
                temp_map.setdefault('quantity', idx)
            elif any(w in key for w in ['unit of measure', 'unit', 'uom', 'measure']):
                temp_map.setdefault('unit', idx)
            elif any(w in key for w in ['remarks', 'remark', 'notes', 'note', 'comments', 'comment']):
                temp_map.setdefault('notes', idx)
            elif any(w in key for w in ['description', 'material', 'item name', 'product name', 'details', 'name']):
                temp_map.setdefault('description', idx)
            elif 'part' in key:
                temp_map.setdefault('part_number', idx)
            elif 'item' in key:
                temp_map.setdefault('item_id', idx)

        if len(temp_map) > best_matches:
            best_matches = len(temp_map)
            best_header_idx = r_idx
            header_idx_map = temp_map

    if best_header_idx is None:
        best_header_idx = 0
        header_row = non_empty_rows[0]
        if len(header_row) > 0: header_idx_map.setdefault('part_number', 0)
        if len(header_row) > 1: header_idx_map.setdefault('item_id', 1)
        if len(header_row) > 2: header_idx_map.setdefault('description', 2)
        if len(header_row) > 3: header_idx_map.setdefault('quantity', 3)
        if len(header_row) > 4: header_idx_map.setdefault('unit', 4)

    data_rows = non_empty_rows[best_header_idx + 1:]
    if not data_rows:
        data_rows = non_empty_rows

    if mode == 'REPLACE':
        active = session.scalars(select(GoodsReceiptRecord).where(GoodsReceiptRecord.is_deleted.is_(False))).all()
        for r in active:
            r.is_deleted = True

    imported_rows: list[GoodsReceiptRecord] = []
    base_idx = session.scalar(select(func.coalesce(func.max(GoodsReceiptRecord.row_index), 0)).where(GoodsReceiptRecord.is_deleted.is_(False))) or 0

    for i, row in enumerate(data_rows, start=1):
        if not any(cell is not None and str(cell).strip() != '' for cell in row):
            continue

        part_no = str(row[header_idx_map['part_number']]).strip() if 'part_number' in header_idx_map and header_idx_map['part_number'] < len(row) and row[header_idx_map['part_number']] is not None else f'P-{i}'
        item_id = str(row[header_idx_map['item_id']]).strip() if 'item_id' in header_idx_map and header_idx_map['item_id'] < len(row) and row[header_idx_map['item_id']] is not None else str(i)

        raw_desc = str(row[header_idx_map['description']]).strip() if 'description' in header_idx_map and header_idx_map['description'] < len(row) and row[header_idx_map['description']] is not None else ''
        desc = raw_desc if raw_desc else f'{part_no} Item'

        raw_qty = row[header_idx_map['quantity']] if 'quantity' in header_idx_map and header_idx_map['quantity'] < len(row) else 1
        try:
            qty = Decimal(str(raw_qty).strip())
        except (InvalidOperation, ValueError, TypeError):
            qty = Decimal('0')

        unit = str(row[header_idx_map['unit']]).strip() if 'unit' in header_idx_map and header_idx_map['unit'] < len(row) and row[header_idx_map['unit']] is not None else 'Nos'
        po_no = str(row[header_idx_map['po_number']]).strip() if 'po_number' in header_idx_map and header_idx_map['po_number'] < len(row) and row[header_idx_map['po_number']] is not None else None
        supplier = str(row[header_idx_map['supplier_name']]).strip() if 'supplier_name' in header_idx_map and header_idx_map['supplier_name'] < len(row) and row[header_idx_map['supplier_name']] is not None else None
        grn_no = str(row[header_idx_map['source_grn_id']]).strip() if 'source_grn_id' in header_idx_map and header_idx_map['source_grn_id'] < len(row) and row[header_idx_map['source_grn_id']] is not None else f'EXCEL-{sheet_name}'
        grn_dt = str(row[header_idx_map['grn_date']]).strip() if 'grn_date' in header_idx_map and header_idx_map['grn_date'] < len(row) and row[header_idx_map['grn_date']] is not None else None
        plant_val = str(row[header_idx_map['plant']]).strip() if 'plant' in header_idx_map and header_idx_map['plant'] < len(row) and row[header_idx_map['plant']] is not None else None
        notes_val = str(row[header_idx_map['notes']]).strip() if 'notes' in header_idx_map and header_idx_map['notes'] < len(row) and row[header_idx_map['notes']] is not None else None

        rec = GoodsReceiptRecord(
            id=uuid4(),
            row_index=base_idx + i,
            part_number=part_no,
            item_id=item_id,
            description=desc,
            quantity=qty,
            unit=unit,
            po_number=po_no,
            supplier_name=supplier,
            status='SAVED',
            source_grn_id=grn_no,
            plant=plant_val,
            grn_date=grn_dt,
            notes=notes_val,
            created_by=actor,
        )
        session.add(rec)
        imported_rows.append(rec)

    audit(session, actor, uuid4(), 'IMPORT_EXCEL_SHEET', reason,
          {'sheet_name': sheet_name, 'filename': filename, 'imported_count': len(imported_rows)},
          entity='goods_receipt_records')

    session.commit()
    return ExcelImportSheetResponse(
        sheet_name=sheet_name,
        imported_count=len(imported_rows),
        records=[to_record_response(r) for r in imported_rows]
    )


def export_workspace_excel(session, search: str | None = None, status: str | None = None, record_ids: list[UUID] | None = None) -> bytes:
    query = select(GoodsReceiptRecord).where(GoodsReceiptRecord.is_deleted.is_(False))
    if record_ids:
        query = query.where(GoodsReceiptRecord.id.in_(record_ids))
    elif status and status.upper() != 'ALL':
        query = query.where(GoodsReceiptRecord.status == status.upper())
    if search:
        s = f"%{search.strip()}%"
        query = query.where(
            or_(
                GoodsReceiptRecord.part_number.ilike(s),
                GoodsReceiptRecord.item_id.ilike(s),
                GoodsReceiptRecord.description.ilike(s),
                GoodsReceiptRecord.po_number.ilike(s),
                GoodsReceiptRecord.supplier_name.ilike(s),
                GoodsReceiptRecord.unit.ilike(s),
                GoodsReceiptRecord.plant.ilike(s),
                GoodsReceiptRecord.source_grn_id.ilike(s),
            )
        )
    rows = session.scalars(query.order_by(GoodsReceiptRecord.row_index.asc().nulls_last(), GoodsReceiptRecord.created_at.asc())).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Goods Receipts"

    header_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = ["#", "Part Number", "Item ID", "Description", "Qty", "Unit", "PO Number", "Supplier", "GRN No.", "GRN Date", "Plant", "Status", "Remarks"]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center" if col_num in (1, 5, 6, 10, 12) else "left", vertical="center")

    ws.row_dimensions[1].height = 24

    for r_idx, row in enumerate(rows, start=1):
        qty_val = float(row.quantity) if row.quantity is not None else 0.0
        grn_dt = row.grn_date or (row.created_at.strftime('%Y-%m-%d') if row.created_at else '')
        row_data = [
            row.row_index or r_idx,
            row.part_number,
            row.item_id,
            row.description,
            qty_val,
            row.unit,
            row.po_number or '',
            row.supplier_name or '',
            row.source_grn_id or '',
            grn_dt,
            row.plant or '',
            row.status,
            row.notes or '',
        ]
        ws.append(row_data)
        current_row = r_idx + 1
        ws.row_dimensions[current_row].height = 20
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if col_num in (1, 6, 10, 12):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '#,##0.0000'
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

