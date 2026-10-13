import calendar
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import re
import time
from typing import Any, Optional
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import ApplicationError
from app.models.audit import AuditLog
from app.models.auth import User
from app.models.inventory_masters import Consumable, Unit
from app.models.masters import Product
from app.models.stock_statements import StockStatementBatch, StockStatementLine
from app.schemas.auth import CurrentUser
from app.schemas.stock_statements import (
    ComparisonSummary,
    GroupedWarningResponse,
    RowIssue,
    StockStatementConfirmRequest,
    StockStatementBatchUpdateRequest,
    StockStatementDeleteRequest,
    StockStatementDeleteResponse,
    StockStatementHistoryItem,
    StockStatementLineResponse,
    StockStatementMetadata,
    StockStatementPreviewResponse,
    StockStatementResponse,
    StockStatementSummary,
)
from app.services.audit import log_audit_event


# In-memory storage for preview sessions (token -> payload)
_PREVIEW_STORE: dict[str, dict[str, Any]] = {}
PREVIEW_TTL_SECONDS: int = 3600  # 1 hour preview lifetime


def _cleanup_stale_preview_entries() -> None:
    """Evict expired preview entries to prevent unbounded memory growth."""
    now = time.time()
    stale = [k for k, v in _PREVIEW_STORE.items() if now - v.get("created_at", 0) > PREVIEW_TTL_SECONDS]
    for k in stale:
        _PREVIEW_STORE.pop(k, None)


def normalize_item_id(val: Any) -> str:
    """Safely normalize item IDs handling unicode spaces, tabs, and line breaks without altering code meaning."""
    if val is None:
        return ""
    # Collapse any form of whitespace (including non-breaking space \u00a0 and zero-width \ufeff)
    return re.sub(r"[\s\u00a0\ufeff]+", " ", str(val)).strip()


def normalize_item_id_key(val: Any) -> str:
    """Uppercase normalized item id for deterministic case-insensitive dictionary lookup."""
    return normalize_item_id(val).upper()


# Exact KNL column specification for Stock Statement
EXPECTED_STATEMENT_COLUMNS = [
    "Item Id",
    "Item Description",
    "Cust. Part No.",
    "Unit",
    "Opening Qty.",
    "Opening Value",
    "Receipt Qty.",
    "Receipt Value",
    "Issue Qty.",
    "Issue Value",
    "Closing Qty.",
    "Closing Value",
]

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December"
]


def normalize_header(text: Any) -> str:
    """Normalize header text: strip, lowercase, collapse spaces, remove periods."""
    if not text:
        return ""
    clean = re.sub(r"\s+", " ", str(text)).strip().lower()
    return clean.replace(".", "")


def parse_raw_decimal(value: Any, row_num: int, col_name: str, issues: list[RowIssue]) -> Optional[Decimal]:
    """Parse raw cell value into Decimal handling Indian grouping, brackets, negatives, and '-'."""
    if value is None:
        return None

    # Already numeric
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and (value != value or value == float("inf") or value == float("-inf")):
            issues.append(RowIssue(row=row_num, column=col_name, description="Invalid float or NaN/Inf", severity="ERROR", code="INVALID_DECIMAL"))
            return None
        return Decimal(str(value))

    val_str = str(value).strip()
    if not val_str:
        return None

    # Check for excel errors
    if val_str.startswith("#") or val_str in ("#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NUM!"):
        issues.append(RowIssue(row=row_num, column=col_name, description=f"Excel calculation error: {val_str}", severity="ERROR", code="EXCEL_CALC_ERROR"))
        return None

    # Zero indicators
    if val_str in ("-", "--", "—", "NIL", "nil", "Nil"):
        return Decimal("0.0000")

    # Parenthesized negative e.g. (1,234.56)
    is_negative = False
    if val_str.startswith("(") and val_str.endswith(")"):
        is_negative = True
        val_str = val_str[1:-1].strip()
    elif val_str.startswith("-"):
        is_negative = True
        val_str = val_str[1:].strip()

    # Remove commas and spaces
    clean_val = val_str.replace(",", "").replace(" ", "").replace("₹", "").replace("Rs", "")
    try:
        dec = Decimal(clean_val)
        return -dec if is_negative else dec
    except InvalidOperation:
        issues.append(RowIssue(row=row_num, column=col_name, description=f"Non-numeric value '{value}' in numeric column", severity="ERROR", code="INVALID_DECIMAL"))
        return None


SHORT_MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def detect_month_in_text(text: str) -> Optional[tuple[int, Optional[int]]]:
    """Check text for month names (full or abbreviated) and optional year."""
    if not text:
        return None
    for idx, name in enumerate(MONTH_NAMES, start=1):
        if re.search(r"\b" + re.escape(name) + r"\b", text, re.IGNORECASE):
            year_match = re.search(r"\b(20\d\d)\b", text)
            if not year_match:
                year_match = re.search(r"['’](\d{2})\b", text)
                if year_match:
                    return (idx, 2000 + int(year_match.group(1)))
            year = int(year_match.group(1)) if year_match else None
            return (idx, year)

    for idx, short_name in enumerate(SHORT_MONTHS, start=1):
        if re.search(r"\b" + re.escape(short_name) + r"(?:['’\s\d_-]|$)", text, re.IGNORECASE):
            year_match = re.search(r"\b(20\d\d)\b", text)
            if not year_match:
                year_match = re.search(r"['’](\d{2})\b", text)
                if year_match:
                    return (idx, 2000 + int(year_match.group(1)))
            year = int(year_match.group(1)) if year_match else None
            return (idx, year)

    return None


def get_statement_end_date(year: int, month: int) -> date:
    """Return the last calendar day of the given period."""
    _, last_day = calendar.monthrange(year, month)
    return date(year, month, last_day)


def parse_and_preview_stock_statement(
    session: Session,
    file_bytes: bytes = b"",
    filename: str = "",
    selected_month: int = 1,
    selected_year: int = 2026,
    content: Optional[bytes] = None,
    month: Optional[int] = None,
    year: Optional[int] = None,
) -> StockStatementPreviewResponse:
    """Inspect and validate an uploaded stock statement Excel file."""
    if content is not None:
        file_bytes = content
    if month is not None:
        selected_month = month
    if year is not None:
        selected_year = year
    statement_end = get_statement_end_date(selected_year, selected_month)
    file_hash = hashlib.sha256(file_bytes).hexdigest()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ApplicationError("INVALID_EXCEL_FILE", f"Failed to parse Excel workbook: {str(e)}", 400)

    sheet = wb.active
    if not sheet:
        raise ApplicationError("INVALID_EXCEL_FILE", "The workbook contains no active sheets.", 400)

    # 1. Detect header row & title rows
    header_row_idx = -1
    col_mapping: dict[str, int] = {}  # normalized key -> col 1-indexed
    title_text: Optional[str] = None

    expected_normalized = {
        "item id": "item_id",
        "item description": "item_description",
        "cust part no": "cust_part_no",
        "customer part no": "cust_part_no",
        "unit": "unit",
        "uom": "unit",
        "opening qty": "opening_qty",
        "opening quantity": "opening_qty",
        "opening value": "opening_value",
        "receipt qty": "receipt_qty",
        "receipt quantity": "receipt_qty",
        "receipt value": "receipt_value",
        "issue qty": "issue_qty",
        "issue quantity": "issue_qty",
        "issue value": "issue_value",
        "closing qty": "closing_qty",
        "closing quantity": "closing_qty",
        "closing value": "closing_value",
    }

    errors: list[RowIssue] = []
    warnings: list[RowIssue] = []

    for r in range(1, min(25, sheet.max_row + 1)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        row_normalized = [normalize_header(v) for v in row_vals]

        # Check if title row
        first_non_empty = next((str(v) for v in row_vals if v is not None and str(v).strip()), "")
        if first_non_empty and ("stock statement" in first_non_empty.lower() or "statement" in first_non_empty.lower()):
            title_text = first_non_empty

        # Check if matches header
        matches = {col_idx + 1: expected_normalized[val] for col_idx, val in enumerate(row_normalized) if val in expected_normalized}
        if "item_id" in matches.values() and ("closing_qty" in matches.values() or "opening_qty" in matches.values()):
            header_row_idx = r
            # Map canonical keys to column indexes
            for c_idx, canon in matches.items():
                if canon not in col_mapping:
                    col_mapping[canon] = c_idx
            break

    if header_row_idx == -1:
        errors.append(RowIssue(
            row=1,
            column="Header",
            description="Could not locate valid header row with required columns ('Item Id', 'Closing Qty.', etc.).",
            severity="ERROR",
        ))
        return StockStatementPreviewResponse(
            statement_month=selected_month,
            statement_year=selected_year,
            statement_date=statement_end.isoformat(),
            detected_title=title_text,
            source_filename=filename,
            file_token="",
            errors=errors,
            warnings=warnings,
            can_confirm=False,
        )

    # Check for missing required columns
    required_keys = ["item_id", "opening_qty", "receipt_qty", "issue_qty", "closing_qty"]
    for rk in required_keys:
        if rk not in col_mapping:
            errors.append(RowIssue(row=header_row_idx, column=rk, description=f"Required column '{rk}' is missing from header row.", severity="ERROR"))

    # Title Date Mismatch Check
    if title_text:
        detected_month_year = detect_month_in_text(title_text)
        if detected_month_year:
            det_month, det_year = detected_month_year
            if det_month != selected_month or (det_year and det_year != selected_year):
                det_label = f"{MONTH_NAMES[det_month - 1]} {det_year or ''}".strip()
                sel_label = f"{MONTH_NAMES[selected_month - 1]} {selected_year}"
                errors.append(RowIssue(
                    row=1,
                    column="Title",
                    description=f"File title specifies '{det_label}', but selected period is '{sel_label}'. Upload blocked due to month mismatch.",
                    severity="ERROR",
                ))

    # Pre-fetch Consumables master for verification
    consumables = session.execute(
        select(Consumable.id, Consumable.code, Consumable.name, Consumable.unit_id)
    ).all()
    consumable_by_code = {
        normalize_item_id_key(c.code): c
        for c in consumables
        if c.code
    }

    # Pre-fetch Products master to identify parts/products in whole-store statement
    products = session.execute(
        select(Product.id, Product.code, Product.part_number, Product.item_id)
    ).all()
    product_keys: set[str] = set()
    for p in products:
        if p.code:
            product_keys.add(normalize_item_id_key(p.code))
        if p.part_number:
            product_keys.add(normalize_item_id_key(p.part_number))
        if p.item_id:
            product_keys.add(normalize_item_id_key(p.item_id))

    units = session.execute(select(Unit.id, Unit.code)).all()
    unit_code_by_id = {u.id: u.code for u in units}

    seen_item_ids: set[str] = set()
    unknown_items: set[str] = set()
    parsed_lines: list[dict[str, Any]] = []

    # Iterate rows
    for r in range(header_row_idx + 1, sheet.max_row + 1):
        # Check if entire row is empty (ignore empty string or whitespace cells)
        row_has_data = any(
            sheet.cell(row=r, column=c).value is not None
            and str(sheet.cell(row=r, column=c).value).strip() != ""
            for c in range(1, sheet.max_column + 1)
        )
        if not row_has_data:
            continue

        item_id_raw = sheet.cell(row=r, column=col_mapping.get("item_id", 1)).value
        item_id_str = normalize_item_id(item_id_raw)
        if not item_id_str:
            # Missing Item ID on non-empty row
            errors.append(RowIssue(row=r, column="Item Id", description="Row has data but Item Id is empty.", severity="ERROR", code="EMPTY_ITEM_ID"))
            continue

        item_id_upper = item_id_str.upper()

        # Skip total / subtotal summary rows
        if "TOTAL" in item_id_upper or "GRAND TOTAL" in item_id_upper or "SUB TOTAL" in item_id_upper:
            continue

        # Skip banner / section category headers (e.g. 'KNITVEL NEEDLES', 'Loc. Id. : CONSUMABLES STORE', 'Ashok Leyland - CHILD PARTS', 'BENDING TOOL', etc.)
        # Section banners have no unit and no opening/receipt/issue/closing quantities or values
        unit_raw = sheet.cell(row=r, column=col_mapping["unit"]).value if "unit" in col_mapping else None
        has_unit = unit_raw is not None and str(unit_raw).strip() != ""
        qty_val_keys = ["opening_qty", "opening_value", "receipt_qty", "receipt_value", "issue_qty", "issue_value", "closing_qty", "closing_value"]
        has_any_qty_or_val = any(
            (sheet.cell(row=r, column=col_mapping[k]).value is not None and str(sheet.cell(row=r, column=col_mapping[k]).value).strip() != "")
            for k in qty_val_keys if k in col_mapping
        )
        if not has_unit and not has_any_qty_or_val:
            continue
        if item_id_upper.startswith("LOC. ID") or "KNITVEL" in item_id_upper:
            continue

        # Classify item into CONSUMABLE, PRODUCT_OR_OTHER, UNKNOWN
        master_consumable = consumable_by_code.get(item_id_upper)
        if master_consumable:
            item_class = "CONSUMABLE"
        elif item_id_upper in product_keys:
            item_class = "PRODUCT_OR_OTHER"
        else:
            item_class = "UNKNOWN"
            unknown_items.add(item_id_str)

        # Duplicate Item ID check within file
        if item_id_upper in seen_item_ids:
            errors.append(RowIssue(
                row=r,
                column="Item Id",
                description=f"Duplicate Item Id '{item_id_str}' in upload file.",
                severity="ERROR",
                code="DUPLICATE_ITEM_ID",
                item_id=item_id_str,
                item_class=item_class,
            ))
        seen_item_ids.add(item_id_upper)

        line_issues: list[str] = []

        if item_class == "UNKNOWN":
            warnings.append(RowIssue(
                row=r,
                column="Item Id",
                description=f"Item Id '{item_id_str}' was not found in Consumable Master. Classified as UNKNOWN.",
                severity="WARNING",
                code="ITEM_UNKNOWN_IN_MASTER",
                item_id=item_id_str,
                item_class=item_class,
            ))
            line_issues.append("ITEM_UNKNOWN_IN_MASTER")

        # Extract values
        desc_val = sheet.cell(row=r, column=col_mapping["item_description"]).value if "item_description" in col_mapping else None
        cust_part_val = sheet.cell(row=r, column=col_mapping["cust_part_no"]).value if "cust_part_no" in col_mapping else None
        unit_val = sheet.cell(row=r, column=col_mapping["unit"]).value if "unit" in col_mapping else None

        open_qty = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["opening_qty"]).value if "opening_qty" in col_mapping else None, r, "Opening Qty.", errors)
        open_val = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["opening_value"]).value if "opening_value" in col_mapping else None, r, "Opening Value", errors)
        rcpt_qty = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["receipt_qty"]).value if "receipt_qty" in col_mapping else None, r, "Receipt Qty.", errors)
        rcpt_val = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["receipt_value"]).value if "receipt_value" in col_mapping else None, r, "Receipt Value", errors)
        iss_qty = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["issue_qty"]).value if "issue_qty" in col_mapping else None, r, "Issue Qty.", errors)
        iss_val = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["issue_value"]).value if "issue_value" in col_mapping else None, r, "Issue Value", errors)
        cls_qty = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["closing_qty"]).value if "closing_qty" in col_mapping else None, r, "Closing Qty.", errors)
        cls_val = parse_raw_decimal(sheet.cell(row=r, column=col_mapping["closing_value"]).value if "closing_value" in col_mapping else None, r, "Closing Value", errors)

        # Check for negative receipts or issues (blocking errors)
        if rcpt_qty is not None and rcpt_qty < Decimal("0"):
            errors.append(RowIssue(row=r, column="Receipt Qty.", description="Receipt Qty. cannot be negative.", severity="ERROR", code="NEGATIVE_RECEIPT", item_id=item_id_str, item_class=item_class))
        if rcpt_val is not None and rcpt_val < Decimal("0"):
            errors.append(RowIssue(row=r, column="Receipt Value", description="Receipt Value cannot be negative.", severity="ERROR", code="NEGATIVE_RECEIPT_VAL", item_id=item_id_str, item_class=item_class))
        if iss_qty is not None and iss_qty < Decimal("0"):
            errors.append(RowIssue(row=r, column="Issue Qty.", description="Issue Qty. cannot be negative.", severity="ERROR", code="NEGATIVE_ISSUE", item_id=item_id_str, item_class=item_class))
        if iss_val is not None and iss_val < Decimal("0"):
            errors.append(RowIssue(row=r, column="Issue Value", description="Issue Value cannot be negative.", severity="ERROR", code="NEGATIVE_ISSUE_VAL", item_id=item_id_str, item_class=item_class))

        # Arithmetic warnings
        diff_qty = None
        if open_qty is not None and rcpt_qty is not None and iss_qty is not None and cls_qty is not None:
            expected_cls_qty = open_qty + rcpt_qty - iss_qty
            diff_q = cls_qty - expected_cls_qty
            if abs(diff_q) > Decimal("0.0001"):
                diff_qty = diff_q
                w_msg = f"Closing Qty ({cls_qty}) != Opening ({open_qty}) + Receipt ({rcpt_qty}) - Issue ({iss_qty}). Diff: {diff_q:+f}"
                warnings.append(RowIssue(row=r, column="Closing Qty.", description=w_msg, severity="WARNING", code="QTY_MISMATCH", item_id=item_id_str, item_class=item_class))
                line_issues.append("QTY_MISMATCH")

        diff_val = None
        if open_val is not None and rcpt_val is not None and iss_val is not None and cls_val is not None:
            expected_cls_val = open_val + rcpt_val - iss_val
            diff_v = cls_val - expected_cls_val
            if abs(diff_v) > Decimal("1.0000"):  # 1.0 tolerance for paise/rounding
                diff_val = diff_v
                w_msg = f"Closing Value ({cls_val}) != Opening ({open_val}) + Receipt ({rcpt_val}) - Issue ({iss_val}). Diff: {diff_v:+f}"
                warnings.append(RowIssue(row=r, column="Closing Value", description=w_msg, severity="WARNING", code="VAL_MISMATCH", item_id=item_id_str, item_class=item_class))
                line_issues.append("VAL_MISMATCH")

        # Negative opening or closing quantity warnings
        if open_qty is not None and open_qty < Decimal("0"):
            warnings.append(RowIssue(row=r, column="Opening Qty.", description=f"Negative Opening Qty ({open_qty}).", severity="WARNING", code="NEGATIVE_OPENING", item_id=item_id_str, item_class=item_class))
            line_issues.append("NEGATIVE_OPENING")
        if cls_qty is not None and cls_qty < Decimal("0"):
            warnings.append(RowIssue(row=r, column="Closing Qty.", description=f"Negative Closing Qty ({cls_qty}).", severity="WARNING", code="NEGATIVE_CLOSING", item_id=item_id_str, item_class=item_class))
            line_issues.append("NEGATIVE_CLOSING")

        # Unit mismatch warning (consumables only)
        if item_class == "CONSUMABLE" and master_consumable and unit_val:
            master_unit_code = unit_code_by_id.get(master_consumable.unit_id)
            if master_unit_code and str(unit_val).strip().upper() != master_unit_code.upper():
                warnings.append(RowIssue(
                    row=r,
                    column="Unit",
                    description=f"Unit '{unit_val}' differs from Consumable Master unit '{master_unit_code}'.",
                    severity="WARNING",
                    code="UNIT_MISMATCH",
                    item_id=item_id_str,
                    item_class=item_class,
                ))
                line_issues.append("UNIT_MISMATCH")

        # Description mismatch warning (consumables only, info)
        if item_class == "CONSUMABLE" and master_consumable and desc_val and master_consumable.name:
            if str(desc_val).strip().upper() != master_consumable.name.strip().upper():
                line_issues.append("DESC_MISMATCH")

        # Empty cell warning where value expected
        for col_n, col_v in [
            ("Opening Qty.", open_qty), ("Opening Value", open_val),
            ("Receipt Qty.", rcpt_qty), ("Receipt Value", rcpt_val),
            ("Issue Qty.", iss_qty), ("Issue Value", iss_val),
            ("Closing Qty.", cls_qty), ("Closing Value", cls_val),
        ]:
            if col_v is None:
                code_str = f"EMPTY_{col_n.upper().replace(' ', '_').replace('.', '')}"
                warnings.append(RowIssue(
                    row=r,
                    column=col_n,
                    description=f"Empty cell in {col_n} (treated as null).",
                    severity="WARNING",
                    code=code_str,
                    item_id=item_id_str,
                    item_class=item_class,
                ))
                line_issues.append(code_str)

        parsed_lines.append({
            "source_row": r,
            "item_id": item_id_str,
            "item_class": item_class,
            "consumable_id": str(master_consumable.id) if master_consumable else None,
            "item_description": str(desc_val).strip() if desc_val else None,
            "cust_part_no": str(cust_part_val).strip() if cust_part_val else None,
            "unit": str(unit_val).strip() if unit_val else None,
            "opening_qty": open_qty,
            "opening_value": open_val,
            "receipt_qty": rcpt_qty,
            "receipt_value": rcpt_val,
            "issue_qty": iss_qty,
            "issue_value": iss_val,
            "closing_qty": cls_qty,
            "closing_value": cls_val,
            "diff_closing_qty": diff_qty,
            "diff_closing_value": diff_val,
            "warning_codes": line_issues,
        })

    # Class counts
    consumable_count = sum(1 for p in parsed_lines if p["item_class"] == "CONSUMABLE")
    product_or_other_count = sum(1 for p in parsed_lines if p["item_class"] == "PRODUCT_OR_OTHER")
    unknown_count = sum(1 for p in parsed_lines if p["item_class"] == "UNKNOWN")

    # Group warnings by code
    grouped_map: dict[str, dict[str, Any]] = {}
    for w in warnings:
        c_code = w.code or "OTHER_WARNING"
        if c_code not in grouped_map:
            grouped_map[c_code] = {
                "code": c_code,
                "description": w.description,
                "count": 0,
                "sample_rows": [],
            }
        grouped_map[c_code]["count"] += 1
        if len(grouped_map[c_code]["sample_rows"]) < 10 and w.row not in grouped_map[c_code]["sample_rows"]:
            grouped_map[c_code]["sample_rows"].append(w.row)

    grouped_warnings = [
        GroupedWarningResponse(
            code=v["code"],
            description=v["description"],
            count=v["count"],
            sample_rows=v["sample_rows"],
        )
        for v in sorted(grouped_map.values(), key=lambda x: -x["count"])
    ]

    # Check for existing active batch
    existing = session.scalar(
        select(StockStatementBatch).where(
            StockStatementBatch.statement_month == selected_month,
            StockStatementBatch.statement_year == selected_year,
            StockStatementBatch.status == "active",
        )
    )

    _cleanup_stale_preview_entries()

    can_confirm = len(errors) == 0 and len(parsed_lines) > 0

    token = str(uuid.uuid4())
    all_issues = list(errors) + list(warnings)

    _PREVIEW_STORE[token] = {
        "month": selected_month,
        "year": selected_year,
        "date": statement_end,
        "filename": filename,
        "file_hash": file_hash,
        "lines": parsed_lines,
        "issues": all_issues,
        "errors": list(errors),
        "warnings": list(warnings),
        "warnings_count": len(warnings),
        "errors_count": len(errors),
        "consumable_count": consumable_count,
        "product_or_other_count": product_or_other_count,
        "unknown_count": unknown_count,
        "can_confirm": can_confirm,
        "created_at": time.time(),
    }

    preview_lines = [
        StockStatementLineResponse(
            id=f"preview-{p['source_row']}",
            item_id=p["item_id"],
            consumable_id=p["consumable_id"],
            item_class=p["item_class"],
            item_description=p["item_description"],
            cust_part_no=p["cust_part_no"],
            unit=p["unit"],
            opening_qty=p["opening_qty"],
            opening_value=p["opening_value"],
            receipt_qty=p["receipt_qty"],
            receipt_value=p["receipt_value"],
            issue_qty=p["issue_qty"],
            issue_value=p["issue_value"],
            closing_qty=p["closing_qty"],
            closing_value=p["closing_value"],
            source_row=p["source_row"],
            warning_codes=p["warning_codes"],
            diff_closing_qty=p["diff_closing_qty"],
            diff_closing_value=p["diff_closing_value"],
        )
        for p in parsed_lines[:50]
    ]

    return StockStatementPreviewResponse(
        statement_month=selected_month,
        statement_year=selected_year,
        statement_date=statement_end.isoformat(),
        detected_title=title_text,
        source_filename=filename,
        file_token=token,
        is_existing_batch=existing is not None,
        existing_batch_filename=existing.source_filename if existing else None,
        existing_batch_date=existing.uploaded_at if existing else None,
        errors=errors,
        warnings=warnings,
        grouped_warnings=grouped_warnings,
        unknown_items=sorted(list(unknown_items)),
        total_rows=len(parsed_lines),
        warning_count=len(warnings),
        consumable_count=consumable_count,
        product_or_other_count=product_or_other_count,
        unknown_count=unknown_count,
        lines_preview=preview_lines,
        can_confirm=can_confirm,
    )


def confirm_stock_statement_import(
    session: Session,
    payload: Optional[StockStatementConfirmRequest] = None,
    user: Optional[CurrentUser] = None,
    request: Optional[StockStatementConfirmRequest] = None,
    current_user: Optional[CurrentUser] = None,
) -> StockStatementResponse:
    """Commit the parsed stock statement in one atomic transaction, marking existing batches superseded."""
    req = payload or request
    usr = user or current_user
    if not req or not usr:
        raise ApplicationError("INVALID_REQUEST", "Missing confirmation payload or user.", 400)

    token_data = _PREVIEW_STORE.get(req.file_token)
    if not token_data:
        raise ApplicationError("EXPIRED_TOKEN", "The preview session has expired or is invalid. Please re-upload the file.", 404)

    # Check TTL expiry
    if time.time() - token_data.get("created_at", 0) > PREVIEW_TTL_SECONDS:
        _PREVIEW_STORE.pop(req.file_token, None)
        raise ApplicationError("EXPIRED_TOKEN", "The preview session has expired. Please re-upload the file.", 404)

    # Strictly block confirmation if there are any unresolved blocking errors
    if not token_data.get("can_confirm", False):
        raise ApplicationError(
            "VALIDATION_FAILED",
            "Cannot confirm stock statement import with unresolved blocking validation errors.",
            422,
        )

    if (
        token_data["month"] != req.statement_month
        or token_data["year"] != req.statement_year
    ):
        raise ApplicationError("PERIOD_MISMATCH", "Token period does not match confirmed month/year.", 400)

    try:
        # 1. Existing batch check
        existing_batches = session.scalars(
            select(StockStatementBatch).where(
                StockStatementBatch.statement_month == req.statement_month,
                StockStatementBatch.statement_year == req.statement_year,
                StockStatementBatch.status == "active",
            ).with_for_update().execution_options(populate_existing=True)
        ).all()

        if existing_batches and not req.replace_confirmed:
            raise ApplicationError(
                "REPLACE_CONFIRMATION_REQUIRED",
                f"An active stock statement already exists for {MONTH_NAMES[req.statement_month - 1]} {req.statement_year}. Confirmation required to replace.",
                409,
            )

        # Mark existing active batches as superseded
        now_utc = datetime.now(timezone.utc)
        for old_batch in existing_batches:
            old_batch.status = "superseded"
            old_batch.updated_at = now_utc
            old_batch.updated_by = usr.id
            log_audit_event(
                session=session,
                entity_type="StockStatementBatch",
                entity_id=old_batch.id,
                action="UPDATE",
                actor_id=usr.id,
                old_values={"status": "active"},
                new_values={"status": "superseded"},
                reason=f"Superseded by replacement upload '{token_data['filename']}' for period {req.statement_month}/{req.statement_year}.",
            )

        # 2. Insert new batch
        new_batch = StockStatementBatch(
            statement_month=req.statement_month,
            statement_year=req.statement_year,
            statement_date=token_data["date"],
            source_filename=token_data["filename"],
            file_hash=token_data["file_hash"],
            status="active",
            row_count=len(token_data["lines"]),
            warning_count=token_data["warnings_count"],
            uploaded_by=usr.id,
            uploaded_at=now_utc,
            created_by=usr.id,
            created_at=now_utc,
            updated_by=usr.id,
            updated_at=now_utc,
        )
        session.add(new_batch)
        session.flush()

        # 3. Insert lines
        line_objects: list[StockStatementLine] = []
        for line in token_data["lines"]:
            line_obj = StockStatementLine(
                batch_id=new_batch.id,
                item_id=line["item_id"],
                consumable_id=uuid.UUID(line["consumable_id"]) if line["consumable_id"] else None,
                item_class=line.get("item_class", "CONSUMABLE"),
                item_description=line["item_description"],
                cust_part_no=line["cust_part_no"],
                unit=line["unit"],
                opening_qty=line["opening_qty"],
                opening_value=line["opening_value"],
                receipt_qty=line["receipt_qty"],
                receipt_value=line["receipt_value"],
                issue_qty=line["issue_qty"],
                issue_value=line["issue_value"],
                closing_qty=line["closing_qty"],
                closing_value=line["closing_value"],
                source_row=line["source_row"],
                warning_codes=line["warning_codes"],
                created_at=now_utc,
                updated_at=now_utc,
            )
            line_objects.append(line_obj)

        session.add_all(line_objects)

        log_audit_event(
            session=session,
            entity_type="StockStatementBatch",
            entity_id=new_batch.id,
            action="CREATE",
            actor_id=usr.id,
            old_values=None,
            new_values={
                "period": f"{req.statement_month}/{req.statement_year}",
                "filename": token_data["filename"],
                "row_count": len(line_objects),
                "warning_count": token_data["warnings_count"],
            },
            reason=f"Uploaded monthly stock statement for {MONTH_NAMES[req.statement_month - 1]} {req.statement_year}.",
        )

        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        # Consume token so it cannot be re-used to submit duplicate batches
        _PREVIEW_STORE.pop(req.file_token, None)

    return get_stock_statement(session, req.statement_month, req.statement_year)


def batch_update_stock_statement(session: Session, *, payload: StockStatementBatchUpdateRequest,
                                 user: CurrentUser) -> StockStatementResponse:
    try:
        batch = session.scalar(select(StockStatementBatch).where(
            StockStatementBatch.id == payload.batch_id,
            StockStatementBatch.statement_month == payload.month,
            StockStatementBatch.statement_year == payload.year,
            StockStatementBatch.status == "active",
        ).with_for_update().execution_options(populate_existing=True))
        if batch is None:
            raise ApplicationError("STATEMENT_CHANGED", "The active statement has changed. Refresh before saving.", 409)
        for update in payload.updates:
            line = session.get(StockStatementLine, update.id)
            if line is None or line.batch_id != batch.id:
                raise ApplicationError("INVALID_STATEMENT_ROW", "A row does not belong to the selected statement.", 422)
            changes = update.model_dump(exclude_unset=True, exclude={"id"})
            if not changes:
                continue
            before = {key: str(getattr(line, key)) if getattr(line, key) is not None else None for key in changes}
            for key, value in changes.items():
                setattr(line, key, value)
            line.updated_at = datetime.now(timezone.utc)
            log_audit_event(session=session, entity_type="StockStatementLine", entity_id=line.id,
                action="UPDATE", actor_id=user.id, old_values=before,
                new_values={key: str(value) if value is not None else None for key, value in changes.items()},
                reason=f"Reviewed statement correction for {payload.month}/{payload.year}.")
        batch.updated_at = datetime.now(timezone.utc)
        batch.updated_by = user.id
        session.commit()
    except Exception:
        session.rollback()
        raise
    return get_stock_statement(session, payload.month, payload.year, item_class_filter="ALL")


def delete_stock_statement(
    session: Session,
    *,
    month: int,
    year: int,
    payload: StockStatementDeleteRequest,
    user: CurrentUser,
) -> StockStatementDeleteResponse:
    """Deactivate one reviewed monthly batch and retain its lines and audit history."""
    try:
        batch = session.scalar(
            select(StockStatementBatch).where(
                StockStatementBatch.id == payload.batch_id,
                StockStatementBatch.statement_month == month,
                StockStatementBatch.statement_year == year,
            ).with_for_update().execution_options(populate_existing=True)
        )
        if batch is None:
            raise ApplicationError(
                "STATEMENT_NOT_FOUND", "The selected upload does not belong to this month and year. Refresh the statement.", 404,
            )
        if batch.status not in {"active", "deleted"}:
            raise ApplicationError(
                "STATEMENT_CHANGED", "This statement has been replaced. Refresh and review the current upload before deleting.", 409,
            )
        replayed = batch.status == "deleted"
        result = StockStatementDeleteResponse(
            batch_id=batch.id,
            statement_month=month,
            statement_year=year,
            deleted_row_count=batch.row_count,
            replayed=replayed,
        )
        if not replayed:
            batch.status = "deleted"
            batch.updated_at = datetime.now(timezone.utc)
            batch.updated_by = user.id
            log_audit_event(
                session=session,
                entity_type="StockStatementBatch",
                entity_id=batch.id,
                action="DELETE",
                actor_id=user.id,
                old_values={
                    "status": "active", "statement_month": month, "statement_year": year,
                    "source_filename": batch.source_filename, "row_count": batch.row_count,
                },
                new_values={"status": "deleted", "statement_month": month, "statement_year": year},
                reason=payload.reason,
            )
        session.commit()
        return result
    except Exception:
        session.rollback()
        raise


def get_stock_statement(
    session: Session,
    month: int,
    year: int,
    search: Optional[str] = None,
    unit_filter: Optional[str] = None,
    item_class_filter: Optional[str] = "CONSUMABLE",
    has_warnings_only: bool = False,
) -> StockStatementResponse:
    """Retrieve the active stock statement for the specified month and year."""
    statement_end = get_statement_end_date(year, month)

    batch = session.scalar(
        select(StockStatementBatch).where(
            StockStatementBatch.statement_month == month,
            StockStatementBatch.statement_year == year,
            StockStatementBatch.status == "active",
        )
    )

    if not batch:
        meta = StockStatementMetadata(
            has_statement=False,
            statement_month=month,
            statement_year=year,
            statement_date=statement_end.isoformat(),
            row_count=0,
            warning_count=0,
            consumable_count=0,
            product_or_other_count=0,
            unknown_count=0,
        )
        summary = StockStatementSummary()
        return StockStatementResponse(metadata=meta, summary=summary, records=[])

    creator = session.get(User, batch.created_by) if batch.created_by else None
    updater = session.get(User, batch.updated_by) if batch.updated_by else None

    # Calculate class counts across the whole batch
    class_count_rows = session.execute(
        select(StockStatementLine.item_class, func.count(StockStatementLine.id))
        .where(StockStatementLine.batch_id == batch.id)
        .group_by(StockStatementLine.item_class)
    ).all()
    class_counts = {c: count for c, count in class_count_rows}
    consumable_count = class_counts.get("CONSUMABLE", 0)
    product_or_other_count = class_counts.get("PRODUCT_OR_OTHER", 0)
    unknown_count = class_counts.get("UNKNOWN", 0)

    # Query lines
    stmt = select(StockStatementLine).where(StockStatementLine.batch_id == batch.id)
    if item_class_filter and item_class_filter.upper() != "ALL":
        stmt = stmt.where(StockStatementLine.item_class == item_class_filter.upper())

    if search and search.strip():
        term = f"%{search.strip()}%"
        stmt = stmt.where(
            (StockStatementLine.item_id.ilike(term))
            | (StockStatementLine.item_description.ilike(term))
            | (StockStatementLine.cust_part_no.ilike(term))
        )
    if unit_filter and unit_filter.strip():
        stmt = stmt.where(StockStatementLine.unit == unit_filter.strip())

    stmt = stmt.order_by(StockStatementLine.source_row)
    lines = session.scalars(stmt).all()

    # Calculate summary over all lines in batch
    totals = session.execute(
        select(
            func.count(StockStatementLine.id),
            func.coalesce(func.sum(StockStatementLine.opening_value), 0),
            func.coalesce(func.sum(StockStatementLine.receipt_value), 0),
            func.coalesce(func.sum(StockStatementLine.issue_value), 0),
            func.coalesce(func.sum(StockStatementLine.closing_value), 0),
        ).where(StockStatementLine.batch_id == batch.id)
    ).one()

    warning_rows_count = sum(1 for l in lines if l.warning_codes and len(l.warning_codes) > 0)

    records: list[StockStatementLineResponse] = []
    for l in lines:
        diff_q = None
        if l.opening_qty is not None and l.receipt_qty is not None and l.issue_qty is not None and l.closing_qty is not None:
            diff_q = l.closing_qty - (l.opening_qty + l.receipt_qty - l.issue_qty)
        diff_v = None
        if l.opening_value is not None and l.receipt_value is not None and l.issue_value is not None and l.closing_value is not None:
            diff_v = l.closing_value - (l.opening_value + l.receipt_value - l.issue_value)

        records.append(
            StockStatementLineResponse(
                id=str(l.id),
                item_id=l.item_id,
                consumable_id=str(l.consumable_id) if l.consumable_id else None,
                item_class=l.item_class,
                item_description=l.item_description,
                cust_part_no=l.cust_part_no,
                unit=l.unit,
                opening_qty=l.opening_qty,
                opening_value=l.opening_value,
                receipt_qty=l.receipt_qty,
                receipt_value=l.receipt_value,
                issue_qty=l.issue_qty,
                issue_value=l.issue_value,
                closing_qty=l.closing_qty,
                closing_value=l.closing_value,
                source_row=l.source_row,
                warning_codes=l.warning_codes or [],
                diff_closing_qty=diff_q,
                diff_closing_value=diff_v,
            )
        )

    if has_warnings_only:
        records = [r for r in records if r.warning_codes and len(r.warning_codes) > 0]

    meta = StockStatementMetadata(
        has_statement=True,
        batch_id=str(batch.id),
        statement_month=month,
        statement_year=year,
        statement_date=batch.statement_date.isoformat(),
        source_filename=batch.source_filename,
        status=batch.status,
        created_at=batch.created_at,
        created_by_name=creator.full_name or creator.username if creator else None,
        updated_at=batch.updated_at,
        updated_by_name=updater.full_name or updater.username if updater else None,
        row_count=batch.row_count,
        warning_count=batch.warning_count,
        consumable_count=consumable_count,
        product_or_other_count=product_or_other_count,
        unknown_count=unknown_count,
    )

    summary = StockStatementSummary(
        item_count=int(totals[0]),
        consumable_count=consumable_count,
        product_or_other_count=product_or_other_count,
        unknown_count=unknown_count,
        total_opening_value=Decimal(str(totals[1])),
        total_receipt_value=Decimal(str(totals[2])),
        total_issue_value=Decimal(str(totals[3])),
        total_closing_value=Decimal(str(totals[4])),
        warning_rows_count=warning_rows_count,
    )

    return StockStatementResponse(metadata=meta, summary=summary, records=records)


def list_stock_statement_history(
    session: Session,
    month: int,
    year: int,
) -> list[StockStatementHistoryItem]:
    """List historical uploads for a month/year, including deleted batches."""
    batches = session.scalars(
        select(StockStatementBatch)
        .where(
            StockStatementBatch.statement_month == month,
            StockStatementBatch.statement_year == year,
        )
        .order_by(StockStatementBatch.created_at.desc())
    ).all()

    items: list[StockStatementHistoryItem] = []
    for b in batches:
        uploader = session.get(User, b.uploaded_by) if b.uploaded_by else None
        items.append(
            StockStatementHistoryItem(
                batch_id=str(b.id),
                statement_month=b.statement_month,
                statement_year=b.statement_year,
                statement_date=b.statement_date.isoformat(),
                source_filename=b.source_filename,
                status=b.status,
                row_count=b.row_count,
                warning_count=b.warning_count,
                uploaded_at=b.uploaded_at,
                uploaded_by_name=uploader.full_name or uploader.username if uploader else None,
            )
        )
    return items


def export_stock_statement_excel(
    session: Session,
    month: int,
    year: int,
) -> tuple[bytes, str]:
    """Export the active stock statement to an Excel workbook matching KNL columns."""
    res = get_stock_statement(session, month, year, item_class_filter="ALL")
    if not res.metadata.has_statement:
        raise ApplicationError("NOT_FOUND", f"No stock statement found for {MONTH_NAMES[month - 1]} {year}.", 404)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Stock Statement {month}-{year}"

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    num_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Title line
    ws.append([f"KNL Stock Statement for the month of {MONTH_NAMES[month - 1]} {year} (As of {res.metadata.statement_date})"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPECTED_STATEMENT_COLUMNS))
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    # Header row
    ws.append(EXPECTED_STATEMENT_COLUMNS)
    for col_idx in range(1, len(EXPECTED_STATEMENT_COLUMNS) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data rows
    for r in res.records:
        ws.append([
            r.item_id,
            r.item_description or "",
            r.cust_part_no or "",
            r.unit or "",
            float(r.opening_qty) if r.opening_qty is not None else "",
            float(r.opening_value) if r.opening_value is not None else "",
            float(r.receipt_qty) if r.receipt_qty is not None else "",
            float(r.receipt_value) if r.receipt_value is not None else "",
            float(r.issue_qty) if r.issue_qty is not None else "",
            float(r.issue_value) if r.issue_value is not None else "",
            float(r.closing_qty) if r.closing_qty is not None else "",
            float(r.closing_value) if r.closing_value is not None else "",
        ])
        curr_row = ws.max_row
        for col_idx in range(1, len(EXPECTED_STATEMENT_COLUMNS) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.border = thin_border
            if col_idx in (5, 6, 7, 8, 9, 10, 11, 12):
                cell.alignment = num_alignment
                cell.number_format = "#,##0.0000" if col_idx % 2 != 0 else "#,##0.00"
            elif col_idx in (1, 4):
                cell.alignment = center_alignment

    # Auto column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"KNL_Stock_Statement_{year}_{month:02d}.xlsx"
    return buf.getvalue(), filename


def export_error_report_excel(issues: list[RowIssue], filename: str) -> tuple[bytes, str]:
    """Generate a cleanly formatted Excel error/warning report from validation issues."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validation Issues"

    headers = ["Excel Row", "Column", "Severity", "Code", "Item Id", "Item Class", "Description"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    error_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    warning_fill = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")
    error_font = Font(name="Calibri", size=10, color="C62828", bold=True)
    warning_font = Font(name="Calibri", size=10, color="F57F17", bold=True)
    regular_font = Font(name="Calibri", size=10)

    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    for idx, issue in enumerate(issues, start=2):
        sev = str(issue.severity).upper()
        code_val = getattr(issue, "code", "") or ""
        item_id_val = getattr(issue, "item_id", "") or ""
        item_class_val = getattr(issue, "item_class", "") or ""

        ws.append([
            issue.row,
            issue.column,
            sev,
            code_val,
            item_id_val,
            item_class_val,
            issue.description,
        ])

        row_fill = error_fill if sev == "ERROR" else (warning_fill if sev == "WARNING" else None)
        sev_font = error_font if sev == "ERROR" else (warning_font if sev == "WARNING" else regular_font)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=idx, column=col_idx)
            cell.border = thin_border
            if row_fill:
                cell.fill = row_fill
            if col_idx in (1, 3, 4, 6):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

            if col_idx == 3:
                cell.font = sev_font
            else:
                cell.font = regular_font

    ws.row_dimensions[1].height = 24
    for r in range(2, len(issues) + 2):
        ws.row_dimensions[r].height = 20

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buf = io.BytesIO()
    wb.save(buf)

    clean_name = filename[:-5] if filename.lower().endswith(".xlsx") else (
        filename[:-4] if filename.lower().endswith(".xls") else filename
    )
    return buf.getvalue(), f"Error_Report_{clean_name}.xlsx"

