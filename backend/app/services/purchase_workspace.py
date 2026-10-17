import calendar
from datetime import datetime, timezone
from decimal import Decimal
import io
from typing import Optional
from uuid import UUID

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import ApplicationError
from app.models.inventory_masters import Consumable, Supplier, SupplierConsumable, Unit
from app.models.prd import PlanningVersion
from app.models.purchase_approval import PurchaseApproval
from app.models.purchase_order import PurchaseDemandEvidence
from app.models.purchase_plan import PurchasePlanDraft
from app.models.audit import AuditLog
from app.models.requirements import CalculatedRequirement, MonthlyRequirementRecord
from app.models.stock_statements import (
    StockParameter,
    StockParameterBatch,
    StockStatementBatch,
    StockStatementLine,
)
from app.schemas.auth import CurrentUser
from app.schemas.purchase_approval import SubmitPurchaseApprovalRequest
from app.schemas.purchase_workspace import (
    PurchaseSubmitResponse,
    PurchaseWorkspaceRecord,
    PurchaseWorkspaceResponse,
)
from app.services.purchase_approval import submit_purchase_approval
from app.services.requirements_workspace import get_monthly_plan_metadata
from app.services.stock_management_service import resolve_stock_parameter_for_item


def _filter_records(records, search, status_filter):
    if search:
        query = search.casefold()
        records = [row for row in records if any(query in (value or '').casefold() for value in
            (row.consumable_code, row.consumable_name, row.supplier_code, row.supplier_name))]
    if status_filter:
        records = [row for row in records if row.status == status_filter]
    return sorted(records, key=lambda row: (row.status != 'READY_FOR_APPROVAL', row.consumable_code))


def _with_plan_edits(session, workspace):
    if not workspace.records:
        return workspace
    material_ids = {row.consumable_id for row in workspace.records if row.consumable_id}
    masters = {identity:(name, code) for identity, name, code in session.execute(select(Consumable.id, Consumable.name, Unit.code)
        .outerjoin(Unit, Unit.id == Consumable.unit_id).where(Consumable.id.in_(material_ids)))}
    drafts = {row.consumable_id: (row, code) for row, code in session.execute(select(PurchasePlanDraft, Unit.code)
        .join(Unit, Unit.id == PurchasePlanDraft.unit_id)
        .where(PurchasePlanDraft.planning_version_id == UUID(workspace.plan_metadata.plan_id)))}
    for row in workspace.records:
        if row.consumable_id in masters:
            row.consumable_name, row.master_unit = masters[row.consumable_id]
        value = drafts.get(row.consumable_id)
        if value:
            draft, code = value
            row.requested_purchase_qty = draft.requested_qty
            row.requested_purchase_unit = code
            row.purchase_notes = draft.notes
            row.draft_version = draft.version
    return workspace


def _central_workspace(session, version, meta, *, month, year, search, status_filter, current_user):
    """Show saved central demand without inventing stock or bypassing traced assessment."""
    if current_user and not current_user.is_super_admin and current_user.plant_ids:
        raise ApplicationError('CENTRAL_REVIEW_REQUIRED', 'Central plans are reviewed by Super Admin.', 403)
    totals = session.execute(select(
        Consumable.id, Consumable.code, Consumable.name, CalculatedRequirement.unit_id,
        CalculatedRequirement.uom, func.sum(CalculatedRequirement.calculated_qty),
        func.count(CalculatedRequirement.id), func.min(CalculatedRequirement.created_at),
        func.max(CalculatedRequirement.updated_at),
    ).join(CalculatedRequirement, CalculatedRequirement.consumable_id == Consumable.id)
        .where(CalculatedRequirement.planning_version_id == version.id)
        .group_by(Consumable.id, Consumable.code, Consumable.name,
            CalculatedRequirement.unit_id, CalculatedRequirement.uom)).all()
    approved = version.status == 'APPROVED' and version.approved_at is not None
    submitted = {}
    if approved:
        # Only the existing traced assessment can supply central purchase quantities.
        for approval, report, supplier in session.execute(select(PurchaseApproval, PurchaseDemandEvidence.report, Supplier)
            .join(PurchaseDemandEvidence, PurchaseDemandEvidence.approval_id == PurchaseApproval.id)
            .join(Supplier, Supplier.id == PurchaseApproval.supplier_id)
            .where(PurchaseApproval.planning_version_id == version.id)
            .order_by(PurchaseApproval.created_at.desc(), PurchaseApproval.id.desc())):
            submitted.setdefault((approval.consumable_id, approval.uom), (approval, report, supplier))
    error_count = max(0, meta.record_count - sum(row[6] for row in totals))
    if not totals:
        notice = 'Production is uploaded, but no saved consumable requirements are available for this revision. Open the central requirement plan, calculate and save, then review the results.'
    elif not approved:
        notice = f'{len(totals)} consumables have saved central requirements. '
        if error_count:
            notice += f'{error_count} requirement configuration issues need review. '
        notice += 'Review the requirements and approve the central plan before completing the inventory and supplier assessment. Purchase quantities remain unavailable until that assessment.'
    else:
        notice = 'Central requirements are approved. Complete the central inventory and supplier assessment to calculate purchase quantities. Submitted assessments appear here with their saved stock and purchase evidence.'
    records = []
    for material_id, code, name, unit_id, uom, required, count, created_at, updated_at in totals:
        previous = submitted.get((material_id, uom))
        approval, report, supplier = previous if previous else (None, {}, None)
        status = 'REQUIREMENT_REVIEW_REQUIRED' if not approved else 'ASSESSMENT_REQUIRED'
        remarks = 'Review and approve the central requirement plan.' if not approved else 'Use Central purchase assessment to select verified inventory and supplier policies.'
        if approval:
            status = 'SUBMITTED'
            remarks = f'Latest traced purchase assessment is {approval.status}. Stock and incoming quantities are the saved assessment evidence.'
        records.append(PurchaseWorkspaceRecord(
            id=f'{version.id}:{material_id}:{unit_id}', planning_period=version.planning_period,
            planning_month=month, planning_year=year, plan_id=str(version.id), revision=version.revision_label,
            consumable_id=material_id, consumable_code=code, consumable_name=name, unit=uom,
            required_qty=_q4(required), source_row_count=count, created_at=created_at, updated_at=updated_at,
            supplier_id=supplier.id if supplier else None, supplier_code=supplier.code if supplier else None,
            supplier_name=supplier.name if supplier else None,
            stock_qty=_q4(report['current_stock']) if report.get('current_stock') is not None else None,
            po_pending_qty=_q4(report['confirmed_incoming_before_receipt']) if report.get('confirmed_incoming_before_receipt') is not None else None,
            shortage_qty=approval.raw_calculated_qty if approval else None,
            system_recommended_qty=approval.system_recommended_qty if approval else None,
            approval_id=approval.id if approval else None, approval_status=approval.status if approval else None,
            status=status, remarks=remarks,
        ))
    return PurchaseWorkspaceResponse(plan_metadata=meta,
        records=_filter_records(records, search, status_filter), notice=notice)


def _q4(value: Decimal | int | str | None) -> Decimal:
    if value is None:
        return Decimal("0.0000")
    return Decimal(str(value)).quantize(Decimal("0.0001"))


def _latest_version(session: Session, planning_period: str) -> PlanningVersion | None:
    return session.scalar(
        select(PlanningVersion)
        .where(PlanningVersion.planning_period == planning_period)
        .order_by(PlanningVersion.version_number.desc())
        .limit(1)
    )


def _find_active_consumable(session: Session, code: str) -> tuple[Consumable | None, Unit | None]:
    consumable = session.scalar(
        select(Consumable).where(
            func.upper(Consumable.code) == code.strip().upper(),
            Consumable.is_active == True,  # noqa: E712
        )
    )
    unit = session.get(Unit, consumable.unit_id) if consumable else None
    return consumable, unit


def _supplier_options(session: Session, consumable_id) -> list[tuple[SupplierConsumable, Supplier]]:
    rows = session.execute(
        select(SupplierConsumable, Supplier)
        .join(Supplier, SupplierConsumable.supplier_id == Supplier.id)
        .where(
            SupplierConsumable.consumable_id == consumable_id,
            SupplierConsumable.is_active == True,  # noqa: E712
            Supplier.is_active == True,  # noqa: E712
        )
        .order_by(Supplier.code)
    ).all()
    return list(rows)


def _approval_for(
    session: Session,
    *,
    planning_version_id,
    consumable_id,
    supplier_id,
) -> PurchaseApproval | None:
    return session.scalar(
        select(PurchaseApproval)
        .where(
            PurchaseApproval.planning_version_id == planning_version_id,
            PurchaseApproval.consumable_id == consumable_id,
            PurchaseApproval.supplier_id == supplier_id,
        )
        .order_by(PurchaseApproval.created_at.desc())
        .limit(1)
    )


def list_purchase_workspace(
    session: Session,
    *,
    month: int,
    year: int,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    category: Optional[str] = None,
    current_user: CurrentUser | None = None,
) -> PurchaseWorkspaceResponse:
    planning_period = f"{year:04d}-{month:02d}"
    meta = get_monthly_plan_metadata(session, month, year)
    version = _latest_version(session, planning_period)

    if not version:
        return PurchaseWorkspaceResponse(plan_metadata=meta, records=[])

    if meta.source_workflow == 'BASIC_CENTRAL_V1':
        return _with_plan_edits(session, _central_workspace(session, version, meta, month=month, year=year,
            search=search, status_filter=status_filter, current_user=current_user))

    rows = session.scalars(
        select(MonthlyRequirementRecord).where(
            MonthlyRequirementRecord.planning_version_id == version.id,
            MonthlyRequirementRecord.is_deleted == False,  # noqa: E712
        )
    ).all()

    # Look up previous month's stock statement and current month's stock parameters
    prev_month = 12 if month == 1 else month - 1
    prev_year = year - 1 if month == 1 else year

    prev_statement_batch = session.scalar(
        select(StockStatementBatch).where(
            StockStatementBatch.statement_month == prev_month,
            StockStatementBatch.statement_year == prev_year,
            StockStatementBatch.status == "active",
        )
    )
    stmt_lines_map: dict[str, StockStatementLine] = {}
    if prev_statement_batch:
        lines = session.scalars(
            select(StockStatementLine).where(
                StockStatementLine.batch_id == prev_statement_batch.id,
                StockStatementLine.item_class == "CONSUMABLE",
            )
        ).all()
        for l in lines:
            stmt_lines_map[l.item_id.strip().upper()] = l

    current_param_batch = session.scalar(
        select(StockParameterBatch).where(
            StockParameterBatch.month == month,
            StockParameterBatch.year == year,
            StockParameterBatch.status == "active",
        )
    )
    params_by_item: dict[str, list[StockParameter]] = {}
    if current_param_batch:
        params = session.scalars(
            select(StockParameter)
            .where(StockParameter.batch_id == current_param_batch.id)
            .order_by(StockParameter.source_row)
        ).all()
        for p in params:
            params_by_item.setdefault(p.item_id.strip().upper(), []).append(p)

    grouped: dict[str, dict[str, object]] = {}
    for row in rows:
        key = row.consumable_code.strip().upper() or row.consumable_name.strip().upper()
        bucket = grouped.setdefault(
            key,
            {
                "code": row.consumable_code.strip(),
                "name": row.consumable_name.strip(),
                "unit": row.unit,
                "required": Decimal("0"),
                "stock": Decimal("0"),
                "po_pending": Decimal("0"),
                "shortage": Decimal("0"),
                "count": 0,
                "created_at": row.created_at,
                "updated_at": row.updated_at,
            },
        )
        bucket["required"] = bucket["required"] + row.scheduled_consumable_qty  # type: ignore[operator]
        bucket["stock"] = bucket["stock"] + row.stock_qty  # type: ignore[operator]
        bucket["po_pending"] = bucket["po_pending"] + row.po_pending_qty  # type: ignore[operator]
        bucket["shortage"] = bucket["shortage"] + row.shortage_qty  # type: ignore[operator]
        bucket["count"] = int(bucket["count"]) + 1
        if row.updated_at:
            bucket["updated_at"] = row.updated_at

    records: list[PurchaseWorkspaceRecord] = []
    for key, bucket in grouped.items():
        code = str(bucket["code"])
        consumable, stock_unit = _find_active_consumable(session, code)
        suppliers = _supplier_options(session, consumable.id) if consumable else []
        supplier = suppliers[0][1] if suppliers else None
        shortage = _q4(bucket["shortage"])  # type: ignore[arg-type]
        recommended = shortage if shortage > Decimal("0") else Decimal("0.0000")
        approval = (
            _approval_for(
                session,
                planning_version_id=version.id,
                consumable_id=consumable.id,
                supplier_id=supplier.id,
            )
            if consumable and supplier
            else None
        )

        if not consumable:
            status = "MASTER_REQUIRED"
            remarks = "No active consumable master matches this requirement code."
        elif not supplier:
            status = "SUPPLIER_REQUIRED"
            remarks = "No active supplier mapping exists for this consumable."
        elif recommended <= Decimal("0"):
            status = "NO_PURCHASE_REQUIRED"
            remarks = "No positive shortage is available for approval submission."
        elif approval:
            status = "SUBMITTED"
            remarks = f"Existing approval queue item is {approval.status}."
        else:
            status = "READY_FOR_APPROVAL"
            remarks = "Ready to submit to the existing purchase approval queue."

        item_key = code.strip().upper()
        flags: list[str] = []
        if prev_statement_batch:
            if item_key in stmt_lines_map:
                stmt_line = stmt_lines_map[item_key]
                if stmt_line.closing_qty is not None:
                    effective_stock = stmt_line.closing_qty
                else:
                    effective_stock = _q4(bucket["stock"])
            else:
                effective_stock = _q4(bucket["stock"])
                flags.append(f"Not in {prev_month:02d}/{prev_year} statement")
        else:
            effective_stock = _q4(bucket["stock"])
            flags.append(f"No stock statement for {prev_month:02d}/{prev_year}")

        if not current_param_batch:
            flags.append(f"No stock parameters for {month:02d}/{year}")
        else:
            item_params = params_by_item.get(item_key, [])
            resolved_param, param_flag = resolve_stock_parameter_for_item(
                item_params,
                category=category,
            )
            if param_flag:
                flags.append(param_flag)
            elif resolved_param is not None:
                if (
                    recommended > Decimal("0")
                    and resolved_param.min_order_qty is not None
                    and recommended < resolved_param.min_order_qty
                ):
                    recommended = _q4(resolved_param.min_order_qty)
                flags.append(
                    f"Limits ({resolved_param.category or 'Default'}): "
                    f"Min={_q4(resolved_param.min_stock_level)}, "
                    f"Max={_q4(resolved_param.max_stock_level)}, "
                    f"MOQ={_q4(resolved_param.min_order_qty)}, "
                    f"Lead={resolved_param.lead_time_days if resolved_param.lead_time_days is not None else 0}d"
                )

        if flags:
            remarks = f"{remarks} [{'; '.join(flags)}]" if remarks else "; ".join(flags)

        record = PurchaseWorkspaceRecord(
            id=f"{version.id}:{key}",
            planning_period=planning_period,
            planning_month=month,
            planning_year=year,
            plan_id=str(version.id),
            revision=version.revision_label,
            consumable_id=consumable.id if consumable else None,
            consumable_code=code,
            consumable_name=str(bucket["name"]),
            supplier_id=supplier.id if supplier else None,
            supplier_code=supplier.code if supplier else None,
            supplier_name=supplier.name if supplier else None,
            unit=stock_unit.code if stock_unit else str(bucket["unit"]),
            required_qty=_q4(bucket["required"]),  # type: ignore[arg-type]
            stock_qty=_q4(effective_stock),
            po_pending_qty=_q4(bucket["po_pending"]),  # type: ignore[arg-type]
            shortage_qty=shortage,
            system_recommended_qty=recommended,
            approval_status=approval.status if approval else None,
            approval_id=approval.id if approval else None,
            status=status,
            remarks=remarks,
            source_row_count=int(bucket["count"]),
            created_at=bucket["created_at"],  # type: ignore[arg-type]
            updated_at=bucket["updated_at"],  # type: ignore[arg-type]
        )
        records.append(record)

    records = _filter_records(records, search, status_filter)
    return _with_plan_edits(session, PurchaseWorkspaceResponse(plan_metadata=meta, records=records))


def save_purchase_plan(session, *, month, year, request, current_user):
    from app.schemas.inventory_masters import ConsumableUpdate
    from app.services import inventory_masters
    elevated = current_user.is_super_admin or 'ADMIN' in current_user.roles
    if not (elevated or current_user.purchase_update or current_user.purchase_orders_update):
        raise ApplicationError('PERMISSION_DENIED', 'Purchase update permission is required to edit this plan.', 403)
    try:
        workspace = list_purchase_workspace(session, month=month, year=year, current_user=current_user)
        if not workspace.plan_metadata.plan_id:
            raise ApplicationError('PLAN_NOT_FOUND', 'Load an existing purchase plan before saving.', 404)
        version = session.scalar(select(PlanningVersion).where(PlanningVersion.id == UUID(workspace.plan_metadata.plan_id)).with_for_update())
        if version.id != _latest_version(session, version.planning_period).id:
            raise ApplicationError('LATEST_REVISION_REQUIRED', 'Reload the latest purchase planning revision.', 409)
        available = {row.id:row for row in workspace.records}
        seen = set()
        for edit in sorted(request.edits, key=lambda item: item.record_id):
            row = available.get(edit.record_id)
            if not row or not row.consumable_id or row.consumable_id in seen:
                raise ApplicationError('INVALID_PLAN_ROW', 'Select one valid row per consumable in the selected period.', 409)
            seen.add(row.consumable_id)
            draft = session.scalar(select(PurchasePlanDraft).where(PurchasePlanDraft.planning_version_id == version.id,
                PurchasePlanDraft.consumable_id == row.consumable_id).with_for_update())
            if (draft.version if draft else 0) != edit.expected_draft_version:
                raise ApplicationError('STALE_PURCHASE_EDIT', 'This row changed since it was loaded. Reload before editing again.', 409)
            material = session.scalar(select(Consumable).where(Consumable.id == row.consumable_id).with_for_update())
            unit = session.get(Unit, material.unit_id) if material.unit_id else None
            if material.name != edit.expected_name or (unit.code if unit else '') != edit.expected_master_unit:
                raise ApplicationError('STALE_PURCHASE_EDIT', 'The master details changed. Reload before saving.', 409)
            master_changes = {}
            if 'consumable_name' in edit.model_fields_set:
                if not edit.consumable_name or not edit.consumable_name.strip():
                    raise ApplicationError('INVALID_MASTER_NAME', 'Consumable name cannot be blank.', 422)
                master_changes['name'] = edit.consumable_name.strip()
                master_changes['description'] = master_changes['name']
            if 'master_unit' in edit.model_fields_set:
                chosen = session.scalar(select(Unit).where(func.upper(Unit.code) == (edit.master_unit or '').strip().upper(), Unit.is_active.is_(True)))
                if chosen is None:
                    raise ApplicationError('INVALID_UNIT', 'Choose an existing active stock unit, such as PCS or KG.', 422)
                master_changes['unit_id'] = chosen.id
            if master_changes:
                if not (elevated or current_user.masters_update or 'masters.consumables.write' in current_user.permissions):
                    raise ApplicationError('PERMISSION_DENIED', 'Master update permission is required to change consumable names or units.', 403)
                resource = next(item for item in inventory_masters.RESOURCES if item.model is Consumable)
                inventory_masters.apply_record_update(session, resource, material.id,
                    ConsumableUpdate(**master_changes, change_reason=request.reason), current_user.id)
            if {'requested_purchase_qty', 'purchase_notes'} & edit.model_fields_set:
                if session.scalar(select(PurchaseApproval.id).where(PurchaseApproval.planning_version_id == version.id,
                        PurchaseApproval.consumable_id == material.id).limit(1)):
                    raise ApplicationError('PURCHASE_REVIEW_PROTECTED', 'This material is already in the approval queue. Review its quantity there to preserve approval history.', 409)
                source_unit = session.scalar(select(Unit).where(Unit.code == row.unit))
                if not source_unit:
                    raise ApplicationError('INVALID_UNIT', 'Review this requirement stock unit before entering a purchase quantity.', 409)
                before = {'requested_qty': str(draft.requested_qty) if draft and draft.requested_qty is not None else None,
                    'notes': draft.notes if draft else None, 'version': draft.version if draft else 0}
                if not draft:
                    draft = PurchasePlanDraft(planning_version_id=version.id, consumable_id=material.id,
                        unit_id=source_unit.id, updated_by=current_user.id, version=0)
                    session.add(draft)
                if 'requested_purchase_qty' in edit.model_fields_set:
                    draft.requested_qty = edit.requested_purchase_qty
                    draft.unit_id = source_unit.id
                if 'purchase_notes' in edit.model_fields_set:
                    draft.notes = edit.purchase_notes.strip() if edit.purchase_notes else None
                draft.version += 1
                draft.updated_by = current_user.id
                session.flush()
                session.add(AuditLog(actor_id=current_user.id, action='EDIT_PURCHASE_PLAN', entity_type='purchase_plan_draft',
                    entity_id=draft.id, old_values=before, new_values={'requested_qty':str(draft.requested_qty) if draft.requested_qty is not None else None,
                    'notes':draft.notes, 'version':draft.version, 'unit_id':str(draft.unit_id)}, reason=request.reason))
        session.commit()
    except Exception:
        session.rollback()
        raise
    return list_purchase_workspace(session, month=month, year=year, current_user=current_user)


def submit_purchase_workspace_records(
    session: Session,
    *,
    month: int,
    year: int,
    record_ids: list[str],
    reason: Optional[str],
    current_user: CurrentUser,
) -> PurchaseSubmitResponse:
    if not record_ids:
        raise ApplicationError("EMPTY_SUBMISSION", "Select at least one purchase recommendation.", 400)

    workspace = list_purchase_workspace(session, month=month, year=year, current_user=current_user)
    if workspace.plan_metadata.source_workflow == 'BASIC_CENTRAL_V1':
        raise ApplicationError('CENTRAL_PURCHASE_ASSESSMENT_REQUIRED',
            'Submit central purchase demand through Central purchase assessment after requirement approval and inventory/supplier review.', 409)
    selected = [row for row in workspace.records if row.id in set(record_ids)]
    if not selected:
        raise ApplicationError("NO_MATCHING_RECORDS", "No selected purchase records were found.", 404)

    submitted = 0
    skipped = 0
    for row in selected:
        if (
            row.status != "READY_FOR_APPROVAL"
            or not row.consumable_id
            or not row.supplier_id
            or row.system_recommended_qty is None
            or row.system_recommended_qty <= Decimal("0")
        ):
            skipped += 1
            continue

        submit_purchase_approval(
            session,
            SubmitPurchaseApprovalRequest(
                consumable_id=row.consumable_id,
                supplier_id=row.supplier_id,
                planning_version_id=UUID(row.plan_id) if row.plan_id else None,
                raw_calculated_qty=row.shortage_qty,
                system_recommended_qty=row.system_recommended_qty,
                uom=row.unit,
                reason=reason or "Submitted from monthly Purchase workspace.",
            ),
            current_user,
        )
        submitted += 1

    refreshed = list_purchase_workspace(session, month=month, year=year)
    return PurchaseSubmitResponse(
        message=f"Submitted {submitted} purchase recommendation(s) to approval queue.",
        submitted_count=submitted,
        skipped_count=skipped,
        records=refreshed.records,
        plan_metadata=refreshed.plan_metadata,
    )


def export_purchase_workspace_excel(
    session: Session,
    *,
    month: int,
    year: int,
    search: Optional[str] = None,
    status_filter: Optional[str] = None,
    record_ids: Optional[list[str]] = None,
    current_user: CurrentUser | None = None,
) -> bytes:
    workspace = list_purchase_workspace(
        session,
        month=month,
        year=year,
        search=search,
        status_filter=status_filter,
        current_user=current_user,
    )
    records = workspace.records
    if record_ids:
        ids = set(record_ids)
        records = [r for r in records if r.id in ids]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Purchase Recommendations"

    navy_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
    steel_fill = PatternFill(start_color="2E5B88", end_color="2E5B88", fill_type="solid")
    title_font = Font(name="Calibri", size=14, bold=True, color="1B3A5C")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    period_label = f"{calendar.month_name[month]} {year}"
    ws.cell(row=1, column=1, value="KNL PURCHASE RECOMMENDATIONS").font = title_font
    ws.cell(
        row=2,
        column=1,
        value=f"Planning Period: {period_label} | Revision: {workspace.plan_metadata.revision_label} | Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}",
    )
    if workspace.notice:
        ws.cell(row=3, column=1, value=workspace.notice)

    headers = [
        "Consumable Code",
        "Consumable Name",
        "Supplier Code",
        "Supplier Name",
        "Unit",
        "Required Qty",
        "Stock Qty",
        "PO Pending Qty",
        "Shortage Qty",
        "System Recommended Qty",
        "Status",
        "Approval Status",
        "Source Rows",
        "Remarks",
        "Stock Unit (Master)",
        "Requested Purchase Qty",
        "Requested Purchase Unit",
        "Purchase Remarks",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.fill = navy_fill if col <= 5 else steel_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, record in enumerate(records, start=5):
        values = [
            record.consumable_code,
            record.consumable_name,
            record.supplier_code or "",
            record.supplier_name or "",
            record.unit,
            record.required_qty,
            record.stock_qty,
            record.po_pending_qty,
            record.shortage_qty,
            record.system_recommended_qty,
            record.status,
            record.approval_status or "",
            record.source_row_count,
            record.remarks or "",
            record.master_unit or "",
            record.requested_purchase_qty,
            record.requested_purchase_unit or record.unit,
            record.purchase_notes or "",
        ]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font = data_font
            cell.border = thin_border
            if col in (6, 7, 8, 9, 10, 16):
                cell.number_format = "#,##0.0000"
                cell.alignment = Alignment(horizontal="right")

    for col in ws.columns:
        width = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(12, min(width + 3, 48))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
