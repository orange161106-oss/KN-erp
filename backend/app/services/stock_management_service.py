from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import io
import re
import time
from typing import Any, Optional
import uuid

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import case, func, select
from sqlalchemy.orm import Session, joinedload

from app.core.errors import ApplicationError
from app.models.audit import AuditLog
from app.models.auth import User
from app.models.inventory_masters import Consumable
from app.models.stock_statements import StockParameter, StockParameterBatch
from app.schemas.auth import CurrentUser
from app.schemas.stock_statements import (
    ComparisonSummary,
    RowIssue,
    RowResolutionInput,
    StockManagementMetadata,
    StockManagementResponse,
    StockManagementSummary,
    StockParameterBatchUpdateRequest,
    StockParameterConfirmRequest,
    StockParameterDeleteRequest,
    StockParameterDeleteResponse,
    StockParameterHistoryItem,
    StockParameterPreviewResponse,
    StockParameterRecord,
)
from app.services.audit import log_audit_event


# ── Single Configurable Uniqueness Rule for Stock Management Parameters ──
# Options:
#   "ITEM_AND_CATEGORY": Unique per (Item Id, Category, Type) within the month.
#                        Allows the same Item Id under different categories/departments.
#   "ITEM_ONLY":         Strictly one row per Item Id within the month.
STOCK_PARAMETER_UNIQUENESS_MODE: str = "ITEM_AND_CATEGORY"

_PARAM_PREVIEW_STORE: dict[str, dict[str, Any]] = {}
PARAM_PREVIEW_TTL_SECONDS: int = 3600


def _cleanup_stale_param_preview_entries() -> None:
    now = time.time()
    stale = [k for k, v in _PARAM_PREVIEW_STORE.items() if now - v.get("created_at", 0) > PARAM_PREVIEW_TTL_SECONDS]
    for k in stale:
        _PARAM_PREVIEW_STORE.pop(k, None)


def normalize_item_id(val: Any) -> str:
    if val is None:
        return ""
    return re.sub(r"[\s\u00a0\ufeff]+", " ", str(val)).strip()


def normalize_item_id_key(val: Any) -> str:
    return normalize_item_id(val).upper()


def normalize_category_key(val: Any) -> str:
    if val is None:
        return ""
    return re.sub(r"[\s\u00a0\ufeff]+", " ", str(val)).strip().upper()


def normalize_description_for_match(val: Any) -> str:
    """Normalize description ignoring differences in whitespace, case, and quote characters."""
    if val is None:
        return ""
    s = str(val).strip().lower()
    return re.sub(r'[\s\u00a0\ufeff"\'’”″`]+', "", s)


def _norm_dec(val: Optional[Decimal]) -> Optional[Decimal]:
    if val is None:
        return None
    return Decimal(str(val)).quantize(Decimal("0.0001"))


def _limits_tuple(
    min_val: Optional[Decimal],
    max_val: Optional[Decimal],
    moq_val: Optional[Decimal],
    lt_val: Optional[int],
) -> tuple[Optional[Decimal], Optional[Decimal], Optional[Decimal], Optional[int]]:
    return (_norm_dec(min_val), _norm_dec(max_val), _norm_dec(moq_val), lt_val)


def resolve_stock_parameter_for_item(
    params_for_item: list[StockParameter],
    category: Optional[str] = None,
) -> tuple[Optional[StockParameter], Optional[str]]:
    """Resolve the authoritative StockParameter for an item that may have 1 or more category rows.

    Rules:
    - If 0 rows: returns (None, "No stock parameters set for item")
    - If category context is provided and matches a row: returns (matched_row, None)
    - If 1 row: returns (params_for_item[0], None)
    - If multiple rows and no category match:
        - If all rows have identical (min, max, moq, lead_time): returns (params_for_item[0], None)
        - If rows have DIFFERENT limits: returns (None, "Ambiguous stock limits across categories — review required")
          and does NOT choose one silently.
    """
    if not params_for_item:
        return None, "No stock parameters set for item"

    if category and category.strip():
        target_cat = normalize_category_key(category)
        for p in params_for_item:
            if normalize_category_key(p.category) == target_cat:
                return p, None

    if len(params_for_item) == 1:
        return params_for_item[0], None

    # Multiple rows without a specific category match: check if all limits are identical
    first = params_for_item[0]
    first_limits = _limits_tuple(
        first.min_stock_level,
        first.max_stock_level,
        first.min_order_qty,
        first.lead_time_days,
    )
    all_identical = all(
        _limits_tuple(p.min_stock_level, p.max_stock_level, p.min_order_qty, p.lead_time_days) == first_limits
        for p in params_for_item[1:]
    )
    if all_identical:
        return first, None

    cats = ", ".join(f"'{p.category or 'Uncategorized'}'" for p in params_for_item)
    return (
        None,
        f"Ambiguous stock limits across {len(params_for_item)} categories ({cats}) — review required",
    )


EXPECTED_PARAMETER_COLUMNS = [
    "Item Id",
    "Category",
    "Type of material",
    "Description",
    "Min stock level",
    "Max stock level",
    "Minimum order qty.",
    "Lead time to supply in days",
]

MONTH_NAMES = [
    "January", "February", "March", "April", "May", "June",
    "July", "August", "September", "October", "November", "December"
]


def normalize_header(text: Any) -> str:
    if not text:
        return ""
    clean = re.sub(r"\s+", " ", str(text)).strip().lower()
    return clean.replace(".", "")


def parse_raw_decimal(
    value: Any,
    row_num: int,
    col_name: str,
    issues: list[RowIssue],
    row_content: Optional[dict[str, Any]] = None,
) -> Optional[Decimal]:
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and (value != value or value == float("inf") or value == float("-inf")):
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Invalid float or NaN/Inf in {col_name}.",
                    severity="ERROR",
                    code="INVALID_NUMBER",
                    suggested_action=f"Enter a valid non-negative number in {col_name} or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        dec = Decimal(str(value))
        if dec < Decimal("0"):
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Negative value ({dec}) not allowed in {col_name}.",
                    severity="ERROR",
                    code="NEGATIVE_VALUE",
                    suggested_action=f"Correct the negative value in {col_name} or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        return dec

    val_str = str(value).strip()
    if not val_str:
        return None

    if val_str.startswith("#") or val_str in ("#REF!", "#VALUE!", "#DIV/0!", "#N/A"):
        issues.append(
            RowIssue(
                row=row_num,
                column=col_name,
                description=f"Row {row_num}: Excel calculation error '{val_str}' in {col_name}.",
                severity="ERROR",
                code="EXCEL_FORMULA_ERROR",
                suggested_action="Fix the formula error in Excel or skip this row.",
                row_content=row_content,
                resolvable=True,
                allowed_actions=["SKIP_ROW"],
            )
        )
        return None

    if val_str in ("-", "--", "—", "NIL", "nil", "Nil"):
        return Decimal("0.0000")

    clean_val = val_str.replace(",", "").replace(" ", "")
    try:
        dec = Decimal(clean_val)
        if dec < Decimal("0"):
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Negative value ({dec}) not allowed in {col_name}.",
                    severity="ERROR",
                    code="NEGATIVE_VALUE",
                    suggested_action=f"Correct the negative value in {col_name} or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        return dec
    except InvalidOperation:
        issues.append(
            RowIssue(
                row=row_num,
                column=col_name,
                description=f"Row {row_num}: Non-numeric value '{value}' in {col_name}.",
                severity="ERROR",
                code="INVALID_NUMBER",
                suggested_action=f"Enter a valid numeric value in {col_name} or skip this row.",
                row_content=row_content,
                resolvable=True,
                allowed_actions=["SKIP_ROW"],
            )
        )
        return None


def parse_raw_int(
    value: Any,
    row_num: int,
    col_name: str,
    issues: list[RowIssue],
    row_content: Optional[dict[str, Any]] = None,
) -> Optional[int]:
    if value is None:
        return None
    if isinstance(value, int) and not isinstance(value, bool):
        if value < 0:
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Negative lead time ({value}) not allowed.",
                    severity="ERROR",
                    code="INVALID_LEAD_TIME",
                    suggested_action="Provide a non-negative whole number of days or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        return value
    if isinstance(value, float):
        if value < 0 or not value.is_integer():
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Lead time must be a non-negative whole number of days, got {value}.",
                    severity="ERROR",
                    code="INVALID_LEAD_TIME",
                    suggested_action="Provide a whole number of days or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        return int(value)

    val_str = str(value).strip()
    if not val_str:
        return None
    if val_str in ("-", "–", "—"):
        return 0
    try:
        val_f = float(val_str)
        if val_f < 0 or not val_f.is_integer():
            issues.append(
                RowIssue(
                    row=row_num,
                    column=col_name,
                    description=f"Row {row_num}: Lead time must be a non-negative whole number of days, got '{value}'.",
                    severity="ERROR",
                    code="INVALID_LEAD_TIME",
                    suggested_action="Provide a whole number of days or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )
            return None
        return int(val_f)
    except ValueError:
        issues.append(
            RowIssue(
                row=row_num,
                column=col_name,
                description=f"Row {row_num}: Invalid lead time days '{value}'.",
                severity="ERROR",
                code="INVALID_LEAD_TIME",
                suggested_action="Provide a valid whole number of days or skip this row.",
                row_content=row_content,
                resolvable=True,
                allowed_actions=["SKIP_ROW"],
            )
        )
        return None


SHORT_MONTHS = ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"]


def detect_month_in_text(text: str) -> Optional[tuple[int, Optional[int]]]:
    if not text:
        return None
    for idx, name in enumerate(MONTH_NAMES, start=1):
        if re.search(r"\b" + re.escape(name) + r"\b", text, re.IGNORECASE):
            year_match = re.search(r"\b(20\d\d)\b", text)
            if not year_match:
                year_match = re.search(r"['’](\d{2})\b", text)
                if year_match:
                    return (idx, 2000 + int(year_match.group(1)))
            year = int(year_match.group(1)) if year_match else None
            return (idx, year)

    for idx, short_name in enumerate(SHORT_MONTHS, start=1):
        if re.search(r"\b" + re.escape(short_name) + r"(?:['’\s\d_-]|$)", text, re.IGNORECASE):
            year_match = re.search(r"\b(20\d\d)\b", text)
            if not year_match:
                year_match = re.search(r"['’](\d{2})\b", text)
                if year_match:
                    return (idx, 2000 + int(year_match.group(1)))
            year = int(year_match.group(1)) if year_match else None
            return (idx, year)

    return None


def get_previous_month(month: int, year: int) -> tuple[int, int]:
    if month == 1:
        return (12, year - 1)
    return (month - 1, year)


def _format_row_summary(rc: dict[str, Any]) -> str:
    """Format a compact plain-language summary of row content."""
    parts = []
    if rc.get("item_id"):
        parts.append(f"Item Id: {rc['item_id']}")
    if rc.get("category"):
        parts.append(f"Category: '{rc['category']}'")
    if rc.get("type_of_material"):
        parts.append(f"Type: '{rc['type_of_material']}'")
    if rc.get("description"):
        parts.append(f"Desc: '{rc['description']}'")
    limits = []
    if rc.get("min_stock_level") is not None:
        limits.append(f"Min={rc['min_stock_level']}")
    if rc.get("max_stock_level") is not None:
        limits.append(f"Max={rc['max_stock_level']}")
    if rc.get("min_order_qty") is not None:
        limits.append(f"MOQ={rc['min_order_qty']}")
    if rc.get("lead_time_days") is not None:
        limits.append(f"Lead={rc['lead_time_days']}d")
    if limits:
        parts.append(", ".join(limits))
    return "; ".join(parts) if parts else "Empty"


def parse_and_preview_stock_parameters(
    session: Session,
    file_bytes: bytes = b"",
    filename: str = "",
    selected_month: int = 1,
    selected_year: int = 2026,
    content: Optional[bytes] = None,
    month: Optional[int] = None,
    year: Optional[int] = None,
    uniqueness_mode: Optional[str] = None,
) -> StockParameterPreviewResponse:
    """Validate stock parameters workbook, classify duplicates/missing IDs, and compare against previous month."""
    if content is not None:
        file_bytes = content
    if month is not None:
        selected_month = month
    if year is not None:
        selected_year = year

    mode = uniqueness_mode or STOCK_PARAMETER_UNIQUENESS_MODE
    effective_date = date(selected_year, selected_month, 1)
    file_hash = hashlib.sha256(file_bytes).hexdigest()

    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    except Exception as e:
        raise ApplicationError("INVALID_EXCEL_FILE", f"Failed to parse Excel workbook: {str(e)}", 400)

    sheet = wb.active
    if not sheet:
        raise ApplicationError("INVALID_EXCEL_FILE", "The workbook contains no active sheets.", 400)

    header_row_idx = -1
    col_mapping: dict[str, int] = {}
    title_text: Optional[str] = None

    expected_normalized = {
        "item id": "item_id",
        "item code": "item_id",
        "category": "category",
        "type of material": "type_of_material",
        "material type": "type_of_material",
        "type": "type_of_material",
        "description": "description",
        "item description": "description",
        "min stock level": "min_stock_level",
        "minimum stock level": "min_stock_level",
        "msl": "min_stock_level",
        "max stock level": "max_stock_level",
        "maximum stock level": "max_stock_level",
        "minimum order qty": "min_order_qty",
        "min order qty": "min_order_qty",
        "minimum order quantity": "min_order_qty",
        "minimum order quanity": "min_order_qty",
        "min order quantity": "min_order_qty",
        "min order quanity": "min_order_qty",
        "moq": "min_order_qty",
        "lead time to supply in days": "lead_time_days",
        "lead time in days": "lead_time_days",
        "lead time to supply": "lead_time_days",
        "lead time": "lead_time_days",
    }

    errors: list[RowIssue] = []
    warnings: list[RowIssue] = []

    for r in range(1, min(25, sheet.max_row + 1)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, sheet.max_column + 1)]
        row_normalized = [normalize_header(v) for v in row_vals]

        first_non_empty = next((str(v) for v in row_vals if v is not None and str(v).strip()), "")
        if first_non_empty and (
            "consumables requirement" in first_non_empty.lower()
            or "stock management" in first_non_empty.lower()
            or "stock control" in first_non_empty.lower()
        ):
            title_text = first_non_empty

        matches = {col_idx + 1: expected_normalized[val] for col_idx, val in enumerate(row_normalized) if val in expected_normalized}
        if "item_id" in matches.values() and ("min_stock_level" in matches.values() or "max_stock_level" in matches.values()):
            header_row_idx = r
            for c_idx, canon in matches.items():
                if canon not in col_mapping:
                    col_mapping[canon] = c_idx
            break

    if header_row_idx == -1:
        errors.append(RowIssue(
            row=1,
            column="Header",
            description="Could not locate valid header row with required columns ('Item Id', 'Min stock level', etc.).",
            severity="ERROR",
            code="MISSING_HEADER",
            suggested_action="Verify the Excel sheet has header columns: Item Id, Category, Type of material, Description, Minimum stock level, Maximum stock level, Minimum order quantity, Lead time to supply in days.",
            resolvable=False,
        ))
        return StockParameterPreviewResponse(
            month=selected_month,
            year=selected_year,
            effective_date=effective_date.isoformat(),
            detected_title=title_text,
            source_filename=filename,
            file_token="",
            uniqueness_mode=mode,
            errors=errors,
            warnings=warnings,
            issue_counts_by_code={"MISSING_HEADER": 1},
            unresolved_blocking_count=1,
            comparison=ComparisonSummary(),
            can_confirm=False,
        )

    # Check required columns
    if "item_id" not in col_mapping:
        errors.append(RowIssue(row=header_row_idx, column="Item Id", description="Required column 'Item Id' is missing.", severity="ERROR", code="MISSING_COLUMN", resolvable=False))
    if "min_stock_level" not in col_mapping:
        errors.append(RowIssue(row=header_row_idx, column="Min stock level", description="Required column 'Min stock level' is missing.", severity="ERROR", code="MISSING_COLUMN", resolvable=False))
    if "max_stock_level" not in col_mapping:
        errors.append(RowIssue(row=header_row_idx, column="Max stock level", description="Required column 'Max stock level' is missing.", severity="ERROR", code="MISSING_COLUMN", resolvable=False))

    # Title Date Mismatch Check
    if title_text:
        detected_month_year = detect_month_in_text(title_text)
        if detected_month_year:
            det_month, det_year = detected_month_year
            if det_month != selected_month or (det_year and det_year != selected_year):
                det_label = f"{MONTH_NAMES[det_month - 1]} {det_year or ''}".strip()
                sel_label = f"{MONTH_NAMES[selected_month - 1]} {selected_year}"
                errors.append(RowIssue(
                    row=1,
                    column="Title",
                    description=f"File title specifies '{det_label}', but selected period is '{sel_label}'. Upload blocked due to month mismatch.",
                    severity="ERROR",
                    code="PERIOD_MISMATCH",
                    suggested_action=f"Select '{det_label}' as the target period or upload the file for '{sel_label}'.",
                    resolvable=False,
                ))

    # Pre-fetch Consumables master
    consumables = session.execute(
        select(Consumable.id, Consumable.code, Consumable.name, Consumable.description, Consumable.category, Consumable.type_of_material)
        .where(Consumable.is_active == True)
    ).all()
    consumable_by_code = {
        normalize_item_id_key(c.code): c
        for c in consumables
        if c.code
    }
    consumable_options = [
        {
            "item_id": c.code,
            "name": c.name or c.description or "",
            "category": c.category or "",
            "type_of_material": c.type_of_material or "",
        }
        for c in sorted(consumables, key=lambda x: x.code or "")
        if c.code
    ]

    # Index master consumables by normalized description/name for possible-match hints
    master_by_desc: dict[str, list[Any]] = {}
    for c in consumables:
        for txt in (c.name, c.description):
            d_key = normalize_description_for_match(txt)
            if d_key:
                master_by_desc.setdefault(d_key, []).append(c)

    # Fetch previous month parameters for comparison
    prev_month, prev_year = get_previous_month(selected_month, selected_year)
    prev_batch = session.scalar(
        select(StockParameterBatch).where(
            StockParameterBatch.month == prev_month,
            StockParameterBatch.year == prev_year,
            StockParameterBatch.status == "active",
        )
    )
    prev_params_by_item: dict[str, StockParameter] = {}
    if prev_batch:
        prev_lines = session.scalars(
            select(StockParameter).where(StockParameter.batch_id == prev_batch.id)
        ).all()
        prev_params_by_item = {p.item_id.strip().upper(): p for p in prev_lines}

    # ── PASS 1: Extract all non-empty rows and index descriptions in the file ──
    extracted_rows: list[dict[str, Any]] = []
    upload_rows_by_desc: dict[str, list[dict[str, Any]]] = {}

    for r in range(header_row_idx + 1, sheet.max_row + 1):
        row_has_data = any(
            sheet.cell(row=r, column=c).value is not None
            and str(sheet.cell(row=r, column=c).value).strip() != ""
            for c in range(1, sheet.max_column + 1)
        )
        if not row_has_data:
            continue

        item_id_raw = sheet.cell(row=r, column=col_mapping.get("item_id", 1)).value
        item_id_str = normalize_item_id(item_id_raw)
        item_id_upper = item_id_str.upper()

        if item_id_upper and ("TOTAL" in item_id_upper or "GRAND TOTAL" in item_id_upper):
            continue

        cat_raw = sheet.cell(row=r, column=col_mapping["category"]).value if "category" in col_mapping else None
        type_raw = sheet.cell(row=r, column=col_mapping["type_of_material"]).value if "type_of_material" in col_mapping else None
        desc_raw = sheet.cell(row=r, column=col_mapping["description"]).value if "description" in col_mapping else None

        cat_str = re.sub(r"\s+", " ", str(cat_raw)).strip() if cat_raw is not None and str(cat_raw).strip() != "" else None
        type_str = re.sub(r"\s+", " ", str(type_raw)).strip() if type_raw is not None and str(type_raw).strip() != "" else None
        desc_str = re.sub(r"\s+", " ", str(desc_raw)).strip() if desc_raw is not None and str(desc_raw).strip() != "" else None

        min_raw = sheet.cell(row=r, column=col_mapping["min_stock_level"]).value if "min_stock_level" in col_mapping else None
        max_raw = sheet.cell(row=r, column=col_mapping["max_stock_level"]).value if "max_stock_level" in col_mapping else None
        moq_raw = sheet.cell(row=r, column=col_mapping["min_order_qty"]).value if "min_order_qty" in col_mapping else None
        lt_raw = sheet.cell(row=r, column=col_mapping["lead_time_days"]).value if "lead_time_days" in col_mapping else None

        prelim_content = {
            "row": r,
            "item_id": item_id_str or None,
            "category": cat_str,
            "type_of_material": type_str,
            "description": desc_str,
            "min_stock_level": str(min_raw).strip() if min_raw is not None and str(min_raw).strip() != "" else None,
            "max_stock_level": str(max_raw).strip() if max_raw is not None and str(max_raw).strip() != "" else None,
            "min_order_qty": str(moq_raw).strip() if moq_raw is not None and str(moq_raw).strip() != "" else None,
            "lead_time_days": str(lt_raw).strip() if lt_raw is not None and str(lt_raw).strip() != "" else None,
        }

        row_parse_errors: list[RowIssue] = []
        min_val = parse_raw_decimal(min_raw, r, "Min stock level", row_parse_errors, prelim_content)
        max_val = parse_raw_decimal(max_raw, r, "Max stock level", row_parse_errors, prelim_content)
        moq_val = parse_raw_decimal(moq_raw, r, "Minimum order qty.", row_parse_errors, prelim_content)
        lt_val = parse_raw_int(lt_raw, r, "Lead time to supply in days", row_parse_errors, prelim_content)

        row_content = {
            "row": r,
            "item_id": item_id_str or None,
            "category": cat_str,
            "type_of_material": type_str,
            "description": desc_str,
            "min_stock_level": str(min_val) if min_val is not None else prelim_content["min_stock_level"],
            "max_stock_level": str(max_val) if max_val is not None else prelim_content["max_stock_level"],
            "min_order_qty": str(moq_val) if moq_val is not None else prelim_content["min_order_qty"],
            "lead_time_days": lt_val if lt_val is not None else prelim_content["lead_time_days"],
        }

        # Update row_content on any parse errors
        for pe in row_parse_errors:
            pe.row_content = row_content

        entry = {
            "source_row": r,
            "item_id": item_id_str,
            "item_id_upper": item_id_upper,
            "category": cat_str,
            "type_of_material": type_str,
            "description": desc_str,
            "min_stock_level": min_val,
            "max_stock_level": max_val,
            "min_order_qty": moq_val,
            "lead_time_days": lt_val,
            "row_content": row_content,
            "parse_errors": row_parse_errors,
        }
        extracted_rows.append(entry)

        if item_id_str and desc_str:
            d_key = normalize_description_for_match(desc_str)
            if d_key:
                upload_rows_by_desc.setdefault(d_key, []).append(entry)

    # Helper to find possible matches by description
    def _find_possible_matches(desc_text: Optional[str], current_row: int) -> list[dict[str, Any]]:
        d_key = normalize_description_for_match(desc_text)
        if not d_key:
            return []
        matches_list: list[dict[str, Any]] = []
        seen_keys: set[tuple[str, Optional[int]]] = set()

        for u_row in upload_rows_by_desc.get(d_key, []):
            if u_row["source_row"] == current_row:
                continue
            key = (u_row["item_id_upper"], u_row["source_row"])
            if key not in seen_keys:
                seen_keys.add(key)
                matches_list.append({
                    "item_id": u_row["item_id"],
                    "description": u_row["description"],
                    "category": u_row["category"],
                    "type_of_material": u_row["type_of_material"],
                    "source": "UPLOAD_ROW",
                    "source_row": u_row["source_row"],
                    "min_stock_level": str(u_row["min_stock_level"]) if u_row["min_stock_level"] is not None else None,
                    "max_stock_level": str(u_row["max_stock_level"]) if u_row["max_stock_level"] is not None else None,
                    "min_order_qty": str(u_row["min_order_qty"]) if u_row["min_order_qty"] is not None else None,
                    "lead_time_days": u_row["lead_time_days"],
                })

        for m_c in master_by_desc.get(d_key, []):
            c_code = normalize_item_id_key(m_c.code)
            if not any(m["item_id"].upper() == c_code for m in matches_list):
                matches_list.append({
                    "item_id": m_c.code,
                    "description": m_c.name or m_c.description,
                    "category": m_c.category,
                    "type_of_material": m_c.type_of_material,
                    "source": "CONSUMABLE_MASTER",
                    "source_row": None,
                    "min_stock_level": None,
                    "max_stock_level": None,
                    "min_order_qty": None,
                    "lead_time_days": None,
                })

        return matches_list

    # ── PASS 2: Validate rows, classify duplicates & missing Item Ids ──
    seen_by_item: dict[str, list[dict[str, Any]]] = {}
    unknown_items: set[str] = set()
    parsed_lines: list[dict[str, Any]] = []
    all_extracted_by_row: dict[int, dict[str, Any]] = {}

    for entry in extracted_rows:
        r = entry["source_row"]
        item_id_str = entry["item_id"]
        item_id_upper = entry["item_id_upper"]
        cat_str = entry["category"]
        type_str = entry["type_of_material"]
        desc_str = entry["description"]
        min_val = entry["min_stock_level"]
        max_val = entry["max_stock_level"]
        moq_val = entry["min_order_qty"]
        lt_val = entry["lead_time_days"]
        row_content = entry["row_content"]

        all_extracted_by_row[r] = entry

        # Add any numeric parse errors
        if entry["parse_errors"]:
            errors.extend(entry["parse_errors"])

        # 1. Missing Item Id handling
        if not item_id_str:
            has_limits = any(v is not None for v in (min_val, max_val, moq_val, lt_val))
            possible_matches = _find_possible_matches(desc_str, r)

            if not has_limits and not cat_str and not type_str:
                errors.append(
                    RowIssue(
                        row=r,
                        column="Item Id",
                        description=(
                            f"Row {r} is an incomplete line: only description '{desc_str or '—'}' is present "
                            f"with no Item Id, Category, Type, or stock limits."
                        ),
                        severity="ERROR",
                        code="INCOMPLETE_LINE",
                        suggested_action="Skip this incomplete row, or assign a valid Item Id if limits are intended.",
                        row_content=row_content,
                        possible_matches=possible_matches,
                        resolvable=True,
                        allowed_actions=["SKIP_ROW", "ASSIGN_ITEM_ID"],
                    )
                )
            else:
                hint_msg = ""
                if possible_matches:
                    m0 = possible_matches[0]
                    if m0["source"] == "UPLOAD_ROW":
                        hint_msg = (
                            f" Possible match by description: '{m0['item_id']}' in row {m0['source_row']} "
                            f"(Category: '{m0.get('category') or '—'}', Min: {m0.get('min_stock_level')}, "
                            f"Max: {m0.get('max_stock_level')}, MOQ: {m0.get('min_order_qty')}, "
                            f"Lead time: {m0.get('lead_time_days')})."
                        )
                    else:
                        hint_msg = (
                            f" Possible match in Consumable Master: '{m0['item_id']}' "
                            f"(Category: '{m0.get('category') or '—'}')."
                        )
                errors.append(
                    RowIssue(
                        row=r,
                        column="Item Id",
                        description=(
                            f"Row {r} has data ({_format_row_summary(row_content)}) "
                            f"but Item Id is empty.{hint_msg}"
                        ),
                        severity="ERROR",
                        code="MISSING_ITEM_ID",
                        suggested_action="Assign a valid Item Id from Consumable Master or skip this row.",
                        row_content=row_content,
                        possible_matches=possible_matches,
                        resolvable=True,
                        allowed_actions=["ASSIGN_ITEM_ID", "SKIP_ROW"],
                    )
                )
            continue

        # 2. Check Consumable Master
        master_c = consumable_by_code.get(item_id_upper)
        if not master_c:
            unknown_items.add(item_id_str)
            possible_matches = _find_possible_matches(desc_str, r)
            errors.append(
                RowIssue(
                    row=r,
                    column="Item Id",
                    item_id=item_id_str,
                    description=(
                        f"Row {r}: Item Id '{item_id_str}' ({_format_row_summary(row_content)}) "
                        f"was not found in Consumable Master."
                    ),
                    severity="ERROR",
                    code="UNKNOWN_ITEM_ID",
                    suggested_action="Assign a valid Item Id from Consumable Master or skip this row.",
                    row_content=row_content,
                    possible_matches=possible_matches,
                    resolvable=True,
                    allowed_actions=["ASSIGN_ITEM_ID", "SKIP_ROW"],
                )
            )

        effective_cat = cat_str if cat_str else (master_c.category if master_c else None)
        effective_type = type_str if type_str else (master_c.type_of_material if master_c else None)
        effective_desc = desc_str if desc_str else (master_c.name if master_c else None)

        line_dict = {
            "source_row": r,
            "item_id": item_id_str,
            "item_id_upper": item_id_upper,
            "consumable_id": str(master_c.id) if master_c else None,
            "category": effective_cat,
            "type_of_material": effective_type,
            "description": effective_desc,
            "min_stock_level": min_val,
            "max_stock_level": max_val,
            "min_order_qty": moq_val,
            "lead_time_days": lt_val,
            "row_content": row_content,
            "resolution_applied": "NONE",
            "resolution_note": None,
        }
        all_extracted_by_row[r] = line_dict

        # 3. Classify Duplicates
        prior_rows = seen_by_item.get(item_id_upper, [])
        if prior_rows:
            curr_cat_key = normalize_category_key(effective_cat)
            curr_type_key = normalize_category_key(effective_type)
            curr_limits = _limits_tuple(min_val, max_val, moq_val, lt_val)

            same_cat_prior = next(
                (
                    pr for pr in prior_rows
                    if normalize_category_key(pr["category"]) == curr_cat_key
                ),
                None,
            )

            if same_cat_prior is not None:
                prior_limits = _limits_tuple(
                    same_cat_prior["min_stock_level"],
                    same_cat_prior["max_stock_level"],
                    same_cat_prior["min_order_qty"],
                    same_cat_prior["lead_time_days"],
                )
                prior_type_key = normalize_category_key(same_cat_prior["type_of_material"])
                if curr_limits == prior_limits and curr_type_key == prior_type_key:
                    # 2a. EXACT DUPLICATE -> WARNING (default action: Keep first, skip repeated row)
                    warnings.append(
                        RowIssue(
                            row=r,
                            column="Item Id",
                            item_id=item_id_str,
                            description=(
                                f"Row {r} repeats row {same_cat_prior['source_row']} "
                                f"({item_id_str}, same category '{effective_cat or '—'}' and same limits: "
                                f"Min={min_val}, Max={max_val}, MOQ={moq_val}, Lead={lt_val}d). Remove one of them."
                            ),
                            severity="WARNING",
                            code="EXACT_DUPLICATE",
                            suggested_action=f"Keep first (row {same_cat_prior['source_row']}), skip the repeated row {r}.",
                            row_content=row_content,
                            conflicting_row=same_cat_prior["source_row"],
                            conflicting_row_content=same_cat_prior["row_content"],
                            resolvable=True,
                            allowed_actions=["KEEP_FIRST", "KEEP_SECOND"],
                            default_action="KEEP_FIRST",
                        )
                    )
                else:
                    # 2c. CONFLICT: same Item Id AND same Category, but DIFFERENT limits -> BLOCKING ERROR
                    prc = same_cat_prior["row_content"]
                    errors.append(
                        RowIssue(
                            row=r,
                            column="Item Id",
                            item_id=item_id_str,
                            description=(
                                f"Row {r} conflicts with row {same_cat_prior['source_row']} "
                                f"({item_id_str}, category '{effective_cat or '—'}' has different limits: "
                                f"Row {same_cat_prior['source_row']} has Min={prc.get('min_stock_level')}, Max={prc.get('max_stock_level')}, "
                                f"MOQ={prc.get('min_order_qty')}, Lead={prc.get('lead_time_days')}d vs "
                                f"Row {r} has Min={min_val}, Max={max_val}, MOQ={moq_val}, Lead={lt_val}d)."
                            ),
                            severity="ERROR",
                            code="SAME_CATEGORY_CONFLICT",
                            suggested_action=f"Choose whether to keep row {same_cat_prior['source_row']} or row {r}.",
                            row_content=row_content,
                            conflicting_row=same_cat_prior["source_row"],
                            conflicting_row_content=prc,
                            resolvable=True,
                            allowed_actions=["KEEP_FIRST", "KEEP_SECOND", "SKIP_ROW"],
                            default_action=None,
                        )
                    )
            else:
                # Same Item Id, DIFFERENT Category
                first_prior = prior_rows[0]
                if mode == "ITEM_ONLY":
                    prc = first_prior["row_content"]
                    errors.append(
                        RowIssue(
                            row=r,
                            column="Item Id",
                            item_id=item_id_str,
                            description=(
                                f"Row {r} repeats Item Id '{item_id_str}' from row {first_prior['source_row']} "
                                f"(Category '{first_prior['category'] or '—'}' vs '{effective_cat or '—'}')."
                            ),
                            severity="ERROR",
                            code="SAME_CATEGORY_CONFLICT",
                            suggested_action=f"Choose whether to keep row {first_prior['source_row']} or row {r}.",
                            row_content=row_content,
                            conflicting_row=first_prior["source_row"],
                            conflicting_row_content=prc,
                            resolvable=True,
                            allowed_actions=["KEEP_FIRST", "KEEP_SECOND", "SKIP_ROW"],
                            default_action=None,
                        )
                    )
                else:
                    # 2b. SAME ITEM, DIFFERENT CATEGORY -> Valid data, keep both, show warning
                    distinct_cats = list(
                        dict.fromkeys(
                            [pr["category"] or "Uncategorized" for pr in prior_rows]
                            + [effective_cat or "Uncategorized"]
                        )
                    )
                    warnings.append(
                        RowIssue(
                            row=r,
                            column="Category",
                            item_id=item_id_str,
                            description=(
                                f"Item {item_id_str} appears under {len(distinct_cats)} categories "
                                f"(Row {first_prior['source_row']}: '{first_prior['category'] or '—'}', "
                                f"Row {r}: '{effective_cat or '—'}'). Both rows are kept."
                            ),
                            severity="WARNING",
                            code="MULTI_CATEGORY_ITEM",
                            suggested_action="Both category limits will be kept. Choose 'Skip this row' only if this category row is unintended.",
                            row_content=row_content,
                            conflicting_row=first_prior["source_row"],
                            conflicting_row_content=first_prior["row_content"],
                            resolvable=True,
                            allowed_actions=["KEEP_BOTH", "SKIP_ROW"],
                            default_action="KEEP_BOTH",
                        )
                    )

        seen_by_item.setdefault(item_id_upper, []).append(line_dict)

        # 4. Min > Max check (blocking error)
        if min_val is not None and max_val is not None and min_val > max_val:
            errors.append(
                RowIssue(
                    row=r,
                    column="Min stock level",
                    item_id=item_id_str,
                    description=(
                        f"Row {r} ({item_id_str}): Min stock level ({min_val}) cannot be greater than Max stock level ({max_val})."
                    ),
                    severity="ERROR",
                    code="MIN_EXCEEDS_MAX",
                    suggested_action="Correct Min/Max stock levels in Excel or skip this row.",
                    row_content=row_content,
                    resolvable=True,
                    allowed_actions=["SKIP_ROW"],
                )
            )

        # 5. Empty limit warnings
        if min_val is None:
            warnings.append(RowIssue(row=r, column="Min stock level", item_id=item_id_str, description=f"Row {r} ({item_id_str}): Min stock level is empty (stored as null).", severity="WARNING", code="EMPTY_LIMIT", row_content=row_content))
        if max_val is None:
            warnings.append(RowIssue(row=r, column="Max stock level", item_id=item_id_str, description=f"Row {r} ({item_id_str}): Max stock level is empty (stored as null).", severity="WARNING", code="EMPTY_LIMIT", row_content=row_content))
        if moq_val is None:
            warnings.append(RowIssue(row=r, column="Minimum order qty.", item_id=item_id_str, description=f"Row {r} ({item_id_str}): Minimum order qty. is empty (stored as null).", severity="WARNING", code="EMPTY_LIMIT", row_content=row_content))
        if lt_val is None:
            warnings.append(RowIssue(row=r, column="Lead time to supply in days", item_id=item_id_str, description=f"Row {r} ({item_id_str}): Lead time is empty (stored as null).", severity="WARNING", code="EMPTY_LIMIT", row_content=row_content))

        # 6. Master discrepancy warnings
        if master_c:
            if cat_str and master_c.category and normalize_category_key(cat_str) != normalize_category_key(master_c.category):
                warnings.append(
                    RowIssue(
                        row=r,
                        column="Category",
                        item_id=item_id_str,
                        description=f"Row {r} ({item_id_str}): Category '{cat_str}' differs from Master '{master_c.category}'.",
                        severity="WARNING",
                        code="MASTER_CATEGORY_DIFF",
                        row_content=row_content,
                    )
                )
            if type_str and master_c.type_of_material and normalize_category_key(type_str) != normalize_category_key(master_c.type_of_material):
                warnings.append(
                    RowIssue(
                        row=r,
                        column="Type of material",
                        item_id=item_id_str,
                        description=f"Row {r} ({item_id_str}): Type of material '{type_str}' differs from Master '{master_c.type_of_material}'.",
                        severity="WARNING",
                        code="MASTER_TYPE_DIFF",
                        row_content=row_content,
                    )
                )

        # 7. Previous month comparison warnings
        if item_id_upper in prev_params_by_item:
            prev_p = prev_params_by_item[item_id_upper]
            if max_val is not None and prev_p.max_stock_level is not None and prev_p.max_stock_level > Decimal("0"):
                ratio = max_val / prev_p.max_stock_level
                if ratio >= Decimal("2.0"):
                    warnings.append(
                        RowIssue(
                            row=r,
                            column="Max stock level",
                            item_id=item_id_str,
                            description=f"Row {r} ({item_id_str}): Max stock level ({max_val}) is over 100% higher than last month ({prev_p.max_stock_level}).",
                            severity="WARNING",
                            code="LARGE_LIMIT_CHANGE",
                            row_content=row_content,
                        )
                    )
            if moq_val is not None and prev_p.min_order_qty is not None and prev_p.min_order_qty > Decimal("0"):
                ratio_moq = moq_val / prev_p.min_order_qty
                if ratio_moq >= Decimal("2.0"):
                    warnings.append(
                        RowIssue(
                            row=r,
                            column="Minimum order qty.",
                            item_id=item_id_str,
                            description=f"Row {r} ({item_id_str}): MOQ ({moq_val}) is over 100% higher than last month ({prev_p.min_order_qty}).",
                            severity="WARNING",
                            code="LARGE_LIMIT_CHANGE",
                            row_content=row_content,
                        )
                    )

        parsed_lines.append(line_dict)

    # Build comparison summary against previous month
    current_item_set = {p["item_id"].strip().upper() for p in parsed_lines}
    prev_item_set = set(prev_params_by_item.keys())

    new_items_count = len(current_item_set - prev_item_set) if prev_batch else 0
    removed_items_count = len(prev_item_set - current_item_set) if prev_batch else 0
    changed_items_count = 0
    changes_sample: list[dict[str, Any]] = []

    if prev_batch:
        for p in parsed_lines:
            k = p["item_id"].strip().upper()
            if k in prev_params_by_item:
                old = prev_params_by_item[k]
                has_diff = (
                    p["min_stock_level"] != old.min_stock_level
                    or p["max_stock_level"] != old.max_stock_level
                    or p["min_order_qty"] != old.min_order_qty
                    or p["lead_time_days"] != old.lead_time_days
                )
                if has_diff:
                    changed_items_count += 1
                    if len(changes_sample) < 10:
                        changes_sample.append({
                            "item_id": p["item_id"],
                            "min_old": str(old.min_stock_level) if old.min_stock_level is not None else None,
                            "min_new": str(p["min_stock_level"]) if p["min_stock_level"] is not None else None,
                            "max_old": str(old.max_stock_level) if old.max_stock_level is not None else None,
                            "max_new": str(p["max_stock_level"]) if p["max_stock_level"] is not None else None,
                        })

    comp_summary = ComparisonSummary(
        has_previous_month=prev_batch is not None,
        previous_month_label=f"{MONTH_NAMES[prev_month - 1]} {prev_year}" if prev_batch else None,
        new_items_count=new_items_count,
        removed_items_count=removed_items_count,
        changed_items_count=changed_items_count,
        changes_sample=changes_sample,
    )

    # Check for existing active batch
    existing = session.scalar(
        select(StockParameterBatch).where(
            StockParameterBatch.month == selected_month,
            StockParameterBatch.year == selected_year,
            StockParameterBatch.status == "active",
        )
    )

    _cleanup_stale_param_preview_entries()

    all_issues = list(errors) + list(warnings)
    issue_counts_by_code: dict[str, int] = {}
    for iss in all_issues:
        c_key = iss.code or iss.severity
        issue_counts_by_code[c_key] = issue_counts_by_code.get(c_key, 0) + 1

    can_confirm = len(errors) == 0 and len(parsed_lines) > 0

    token = str(uuid.uuid4())
    _PARAM_PREVIEW_STORE[token] = {
        "month": selected_month,
        "year": selected_year,
        "date": effective_date,
        "filename": filename,
        "file_hash": file_hash,
        "uniqueness_mode": mode,
        "lines": parsed_lines,
        "all_extracted_by_row": all_extracted_by_row,
        "issues": all_issues,
        "errors": list(errors),
        "warnings": list(warnings),
        "can_confirm": can_confirm,
        "created_at": time.time(),
    }

    preview_lines = [
        StockParameterRecord(
            id=f"param-prev-{p['source_row']}",
            item_id=p["item_id"],
            consumable_id=p["consumable_id"],
            category=p["category"],
            type_of_material=p["type_of_material"],
            description=p["description"],
            min_stock_level=p["min_stock_level"],
            max_stock_level=p["max_stock_level"],
            min_order_qty=p["min_order_qty"],
            lead_time_days=p["lead_time_days"],
            edited_after_import=False,
            source_row=p["source_row"],
            category_count=len(seen_by_item.get(p["item_id_upper"], [])),
        )
        for p in parsed_lines[:50]
    ]

    return StockParameterPreviewResponse(
        month=selected_month,
        year=selected_year,
        effective_date=effective_date.isoformat(),
        detected_title=title_text,
        source_filename=filename,
        file_token=token,
        uniqueness_mode=mode,
        is_existing_batch=existing is not None,
        existing_batch_filename=existing.source_filename if existing else None,
        errors=errors,
        warnings=warnings,
        issue_counts_by_code=issue_counts_by_code,
        unknown_items=sorted(list(unknown_items)),
        consumable_options=consumable_options,
        comparison=comp_summary,
        total_rows=len(extracted_rows),
        unresolved_blocking_count=len(errors),
        lines_preview=preview_lines,
        can_confirm=can_confirm,
    )


def confirm_stock_parameters_import(
    session: Session,
    payload: Optional[StockParameterConfirmRequest] = None,
    user: Optional[CurrentUser] = None,
    request: Optional[StockParameterConfirmRequest] = None,
    current_user: Optional[CurrentUser] = None,
) -> StockManagementResponse:
    """Apply and re-validate resolutions on the server, then commit stock parameters in one atomic transaction."""
    req = payload or request
    usr = user or current_user
    if not req or not usr:
        raise ApplicationError("INVALID_REQUEST", "Missing confirmation payload or user.", 400)

    token_data = _PARAM_PREVIEW_STORE.get(req.file_token)
    if not token_data:
        raise ApplicationError("EXPIRED_TOKEN", "The preview session has expired or is invalid. Please re-upload the file.", 404)

    if time.time() - token_data.get("created_at", 0) > PARAM_PREVIEW_TTL_SECONDS:
        _PARAM_PREVIEW_STORE.pop(req.file_token, None)
        raise ApplicationError("EXPIRED_TOKEN", "The preview session has expired. Please re-upload the file.", 404)

    if token_data["month"] != req.month or token_data["year"] != req.year:
        raise ApplicationError("PERIOD_MISMATCH", "Token period does not match confirmed month/year.", 400)

    preview_errors: list[RowIssue] = token_data.get("errors", [])
    preview_warnings: list[RowIssue] = token_data.get("warnings", [])
    all_extracted_by_row: dict[int, dict[str, Any]] = {
        r: dict(v) for r, v in token_data.get("all_extracted_by_row", {}).items()
    }
    if not all_extracted_by_row and token_data.get("lines"):
        all_extracted_by_row = {p["source_row"]: dict(p) for p in token_data["lines"]}

    # Check for unresolvable errors (e.g. header missing or month mismatch)
    unresolvable = [e for e in preview_errors if not e.resolvable]
    if unresolvable:
        raise ApplicationError(
            "VALIDATION_FAILED",
            f"Cannot confirm import: {unresolvable[0].description}",
            422,
        )

    # Map user-supplied resolutions by row number
    resolutions_by_row: dict[int, RowResolutionInput] = {
        res.row: res for res in (req.resolutions or [])
    }

    # Verify every blocking error row has a resolution from the user
    unresolved_errors = [e for e in preview_errors if e.row not in resolutions_by_row]
    if unresolved_errors:
        rows_list = ", ".join(str(e.row) for e in unresolved_errors[:10])
        raise ApplicationError(
            "VALIDATION_FAILED",
            f"Cannot confirm stock parameters import with unresolved blocking validation errors (rows: {rows_list}).",
            422,
        )

    # Pre-fetch Consumable Master for validating ASSIGN_ITEM_ID
    consumables = session.execute(
        select(Consumable.id, Consumable.code, Consumable.name, Consumable.description, Consumable.category, Consumable.type_of_material)
        .where(Consumable.is_active == True)
    ).all()
    consumable_by_code = {
        normalize_item_id_key(c.code): c
        for c in consumables
        if c.code
    }

    now_utc = datetime.now(timezone.utc)
    actor_name = getattr(usr, "full_name", None) or usr.username

    rows_to_skip: set[int] = set()
    skipped_rows_summary: list[dict[str, Any]] = []
    resolutions_summary: list[dict[str, Any]] = []

    # Group issues by row for validation
    issues_by_row: dict[int, list[RowIssue]] = {}
    for iss in preview_errors + preview_warnings:
        issues_by_row.setdefault(iss.row, []).append(iss)

    # 1. Process explicit user resolutions
    for row_num, res in sorted(resolutions_by_row.items()):
        row_data = all_extracted_by_row.get(row_num)
        if not row_data:
            raise ApplicationError("INVALID_RESOLUTION", f"Row {row_num} does not exist in previewed file.", 400)

        row_issues = issues_by_row.get(row_num, [])
        primary_issue = row_issues[0] if row_issues else None
        issue_desc = primary_issue.description if primary_issue else "User resolution"
        conflict_row_num = primary_issue.conflicting_row if primary_issue else None

        action = (res.action or "").strip().upper()

        if action in ("SKIP_ROW", "SKIP_DUPLICATE", "SKIP"):
            rows_to_skip.add(row_num)
            res_code = "SKIPPED_DUPLICATE" if (primary_issue and primary_issue.code == "EXACT_DUPLICATE") else "SKIPPED_ROW"
            rec = {
                "row": row_num,
                "action": res_code,
                "reason": issue_desc,
                "note": res.note,
                "row_content": row_data.get("row_content"),
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            }
            skipped_rows_summary.append(rec)
            resolutions_summary.append(rec)
            if conflict_row_num and conflict_row_num in all_extracted_by_row:
                all_extracted_by_row[conflict_row_num]["resolution_applied"] = "KEPT_FIRST"
                all_extracted_by_row[conflict_row_num]["resolution_note"] = f"Kept row {conflict_row_num}; skipped row {row_num}."

        elif action == "KEEP_FIRST":
            if not conflict_row_num or conflict_row_num not in all_extracted_by_row:
                raise ApplicationError("INVALID_RESOLUTION", f"Row {row_num} has no conflicting first row to keep.", 400)
            rows_to_skip.add(row_num)
            res_code = "SKIPPED_DUPLICATE" if (primary_issue and primary_issue.code == "EXACT_DUPLICATE") else "SKIPPED_CONFLICT"
            all_extracted_by_row[conflict_row_num]["resolution_applied"] = "KEPT_FIRST"
            all_extracted_by_row[conflict_row_num]["resolution_note"] = f"Kept first row {conflict_row_num}; skipped row {row_num}."
            rec = {
                "row": row_num,
                "conflicting_row": conflict_row_num,
                "action": "KEPT_FIRST",
                "skipped_row": row_num,
                "reason": issue_desc,
                "note": res.note,
                "row_content": row_data.get("row_content"),
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            }
            skipped_rows_summary.append({**rec, "action": res_code})
            resolutions_summary.append(rec)

        elif action in ("KEEP_SECOND", "KEEP_THIS"):
            if not conflict_row_num or conflict_row_num not in all_extracted_by_row:
                raise ApplicationError("INVALID_RESOLUTION", f"Row {row_num} has no conflicting earlier row to replace.", 400)
            rows_to_skip.add(conflict_row_num)
            first_row_data = all_extracted_by_row[conflict_row_num]
            row_data["resolution_applied"] = "KEPT_SECOND"
            row_data["resolution_note"] = f"Kept row {row_num}; replaced earlier row {conflict_row_num}."
            rec = {
                "row": conflict_row_num,
                "kept_row": row_num,
                "action": "KEPT_SECOND",
                "reason": f"Replaced by row {row_num} ({issue_desc})",
                "note": res.note,
                "row_content": first_row_data.get("row_content"),
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            }
            skipped_rows_summary.append(rec)
            resolutions_summary.append(rec)

        elif action == "KEEP_BOTH":
            if primary_issue and primary_issue.code == "SAME_CATEGORY_CONFLICT":
                raise ApplicationError(
                    "INVALID_RESOLUTION",
                    f"Cannot keep both rows for row {row_num}: same Item Id and Category have conflicting limits.",
                    422,
                )
            row_data["resolution_applied"] = "MULTI_CATEGORY_KEPT"
            row_data["resolution_note"] = f"Confirmed keeping multi-category row {row_num}."
            resolutions_summary.append({
                "row": row_num,
                "action": "KEEP_BOTH",
                "reason": issue_desc,
                "note": res.note,
                "row_content": row_data.get("row_content"),
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            })

        elif action == "ASSIGN_ITEM_ID":
            assigned_clean = normalize_item_id(res.assigned_item_id)
            if not assigned_clean:
                raise ApplicationError(
                    "INVALID_ASSIGNED_ITEM_ID",
                    f"Row {row_num}: Assigned Item Id cannot be empty.",
                    422,
                )
            assigned_upper = assigned_clean.upper()
            master_c = consumable_by_code.get(assigned_upper)
            if not master_c:
                raise ApplicationError(
                    "INVALID_ASSIGNED_ITEM_ID",
                    f"Row {row_num}: Assigned Item Id '{assigned_clean}' was not found in Consumable Master.",
                    422,
                )
            row_data["item_id"] = master_c.code
            row_data["item_id_upper"] = assigned_upper
            row_data["consumable_id"] = str(master_c.id)
            if not row_data.get("category") and master_c.category:
                row_data["category"] = master_c.category
            if not row_data.get("type_of_material") and master_c.type_of_material:
                row_data["type_of_material"] = master_c.type_of_material
            if not row_data.get("description") and master_c.name:
                row_data["description"] = master_c.name
            row_data["resolution_applied"] = "ASSIGNED_ITEM_ID"
            row_data["resolution_note"] = f"Assigned Item Id '{master_c.code}' to row {row_num}."
            if row_data.get("row_content"):
                row_data["row_content"] = {**row_data["row_content"], "assigned_item_id": master_c.code}

            resolutions_summary.append({
                "row": row_num,
                "action": "ASSIGNED_ITEM_ID",
                "assigned_item_id": master_c.code,
                "reason": issue_desc,
                "note": res.note,
                "row_content": row_data.get("row_content"),
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            })
        else:
            raise ApplicationError("INVALID_RESOLUTION", f"Unsupported resolution action '{res.action}' on row {row_num}.", 400)

    # 2. Apply default actions for unresolved EXACT_DUPLICATE warnings ("Keep first, skip repeated row")
    for warn in preview_warnings:
        if warn.code == "EXACT_DUPLICATE" and warn.row not in resolutions_by_row and warn.row not in rows_to_skip:
            rows_to_skip.add(warn.row)
            row_data = all_extracted_by_row.get(warn.row, {})
            if warn.conflicting_row and warn.conflicting_row in all_extracted_by_row:
                all_extracted_by_row[warn.conflicting_row]["resolution_applied"] = "KEPT_FIRST"
                all_extracted_by_row[warn.conflicting_row]["resolution_note"] = (
                    f"Kept first row {warn.conflicting_row}; skipped exact duplicate row {warn.row}."
                )
            rec = {
                "row": warn.row,
                "conflicting_row": warn.conflicting_row,
                "action": "SKIPPED_DUPLICATE",
                "reason": warn.description,
                "row_content": row_data.get("row_content") or warn.row_content,
                "resolved_by": actor_name,
                "resolved_by_id": str(usr.id),
                "resolved_at": now_utc.isoformat(),
            }
            skipped_rows_summary.append(rec)
            resolutions_summary.append(rec)
        elif warn.code == "MULTI_CATEGORY_ITEM" and warn.row not in resolutions_by_row and warn.row not in rows_to_skip:
            if warn.row in all_extracted_by_row and all_extracted_by_row[warn.row].get("resolution_applied") in (None, "NONE"):
                all_extracted_by_row[warn.row]["resolution_applied"] = "MULTI_CATEGORY_KEPT"

    # 3. Build final validated list of rows to persist and verify no conflicts remain
    mode = token_data.get("uniqueness_mode", STOCK_PARAMETER_UNIQUENESS_MODE)
    final_lines: list[dict[str, Any]] = []
    seen_unique_keys: dict[tuple[str, ...], int] = {}

    for r_num in sorted(all_extracted_by_row.keys()):
        if r_num in rows_to_skip:
            continue
        line = all_extracted_by_row[r_num]
        item_id_val = normalize_item_id(line.get("item_id"))
        if not item_id_val:
            raise ApplicationError(
                "VALIDATION_FAILED",
                f"Row {r_num} still has an empty Item Id after applying resolutions.",
                422,
            )
        item_upper = item_id_val.upper()
        if item_upper not in consumable_by_code:
            raise ApplicationError(
                "VALIDATION_FAILED",
                f"Row {r_num}: Item Id '{item_id_val}' is not in Consumable Master.",
                422,
            )
        if line.get("parse_errors"):
            raise ApplicationError(
                "VALIDATION_FAILED",
                f"Row {r_num} has unresolved numeric format errors.",
                422,
            )
        min_v = line.get("min_stock_level")
        max_v = line.get("max_stock_level")
        if min_v is not None and max_v is not None and min_v > max_v:
            raise ApplicationError(
                "VALIDATION_FAILED",
                f"Row {r_num} ({item_id_val}): Min stock level ({min_v}) exceeds Max stock level ({max_v}).",
                422,
            )

        cat_key = normalize_category_key(line.get("category"))
        type_key = normalize_category_key(line.get("type_of_material"))
        u_key = (item_upper,) if mode == "ITEM_ONLY" else (item_upper, cat_key, type_key)

        if u_key in seen_unique_keys:
            prev_r = seen_unique_keys[u_key]
            raise ApplicationError(
                "DUPLICATE_KEY_CONFLICT",
                f"Row {r_num} ({item_id_val}, Category '{line.get('category') or '—'}') still conflicts with row {prev_r}. Resolve or skip one of them.",
                422,
            )
        seen_unique_keys[u_key] = r_num
        final_lines.append(line)

    if not final_lines:
        raise ApplicationError("VALIDATION_FAILED", "No valid parameter rows remain to import after skipping rows.", 422)

    try:
        existing_batches = session.scalars(
            select(StockParameterBatch).where(
                StockParameterBatch.month == req.month,
                StockParameterBatch.year == req.year,
                StockParameterBatch.status == "active",
            ).with_for_update().execution_options(populate_existing=True)
        ).all()

        if existing_batches and not req.replace_confirmed:
            raise ApplicationError(
                "REPLACE_CONFIRMATION_REQUIRED",
                f"Active stock management parameters already exist for {MONTH_NAMES[req.month - 1]} {req.year}. Confirmation required to replace.",
                409,
            )

        for old_batch in existing_batches:
            old_batch.status = "superseded"
            old_batch.updated_at = now_utc
            old_batch.updated_by = usr.id
            log_audit_event(
                session=session,
                entity_type="StockParameterBatch",
                entity_id=old_batch.id,
                action="UPDATE",
                actor_id=usr.id,
                old_values={"status": "active"},
                new_values={"status": "superseded"},
                reason=f"Superseded by replacement upload '{token_data['filename']}' for {req.month}/{req.year}.",
            )

        batch_summary = {
            "uniqueness_mode": mode,
            "total_extracted_rows": len(all_extracted_by_row),
            "imported_rows_count": len(final_lines),
            "skipped_rows_count": len(skipped_rows_summary),
            "resolutions_count": len(resolutions_summary),
            "skipped_rows": skipped_rows_summary,
            "resolutions_applied": resolutions_summary,
        }

        new_batch = StockParameterBatch(
            month=req.month,
            year=req.year,
            effective_date=token_data["date"],
            source_filename=token_data["filename"],
            file_hash=token_data["file_hash"],
            status="active",
            row_count=len(final_lines),
            summary=batch_summary,
            uploaded_by=usr.id,
            uploaded_at=now_utc,
            created_by=usr.id,
            created_at=now_utc,
            updated_by=usr.id,
            updated_at=now_utc,
        )
        session.add(new_batch)
        session.flush()

        param_objects: list[StockParameter] = []
        for line in final_lines:
            res_applied = line.get("resolution_applied") or "NONE"
            was_resolved = res_applied != "NONE"
            param_obj = StockParameter(
                batch_id=new_batch.id,
                item_id=line["item_id"],
                consumable_id=uuid.UUID(line["consumable_id"]) if line.get("consumable_id") else None,
                category=line.get("category"),
                type_of_material=line.get("type_of_material"),
                description=line.get("description"),
                min_stock_level=line.get("min_stock_level"),
                max_stock_level=line.get("max_stock_level"),
                min_order_qty=line.get("min_order_qty"),
                lead_time_days=line.get("lead_time_days"),
                edited_after_import=False,
                source_row=line["source_row"],
                resolution_applied=res_applied,
                resolved_by=usr.id if was_resolved else None,
                resolved_at=now_utc if was_resolved else None,
                resolution_note=line.get("resolution_note"),
                created_at=now_utc,
                updated_at=now_utc,
                updated_by=usr.id,
            )
            param_objects.append(param_obj)

        session.add_all(param_objects)

        log_audit_event(
            session=session,
            entity_type="StockParameterBatch",
            entity_id=new_batch.id,
            action="CREATE",
            actor_id=usr.id,
            old_values=None,
            new_values={
                "period": f"{req.month}/{req.year}",
                "filename": token_data["filename"],
                "row_count": len(param_objects),
                "skipped_rows_count": len(skipped_rows_summary),
                "skipped_rows": [
                    {"row": s["row"], "action": s["action"], "reason": s["reason"]}
                    for s in skipped_rows_summary
                ],
                "resolutions_applied": [
                    {
                        "row": r_item["row"],
                        "action": r_item["action"],
                        "assigned_item_id": r_item.get("assigned_item_id"),
                    }
                    for r_item in resolutions_summary
                ],
            },
            reason=f"Uploaded stock control parameters for {MONTH_NAMES[req.month - 1]} {req.year}.",
        )

        session.commit()
    except Exception:
        session.rollback()
        raise
    finally:
        _PARAM_PREVIEW_STORE.pop(req.file_token, None)

    return get_stock_management(session, req.month, req.year)


def delete_stock_parameters(
    session: Session,
    *,
    month: int,
    year: int,
    payload: StockParameterDeleteRequest,
    user: CurrentUser,
) -> StockParameterDeleteResponse:
    """Deactivate one reviewed parameter batch and retain its records and audit history."""
    try:
        batch = session.scalar(
            select(StockParameterBatch).where(
                StockParameterBatch.id == payload.batch_id,
                StockParameterBatch.month == month,
                StockParameterBatch.year == year,
            ).with_for_update().execution_options(populate_existing=True)
        )
        if batch is None:
            raise ApplicationError(
                "PARAMETERS_NOT_FOUND", "The selected upload does not belong to this month and year. Refresh the parameters.", 404,
            )
        if batch.status not in {"active", "deleted"}:
            raise ApplicationError(
                "PARAMETERS_CHANGED", "These parameters have been replaced. Refresh and review the current upload before deleting.", 409,
            )
        replayed = batch.status == "deleted"
        result = StockParameterDeleteResponse(
            batch_id=batch.id,
            month=month,
            year=year,
            deleted_row_count=batch.row_count,
            replayed=replayed,
        )
        if not replayed:
            batch.status = "deleted"
            batch.updated_at = datetime.now(timezone.utc)
            batch.updated_by = user.id
            log_audit_event(
                session=session,
                entity_type="StockParameterBatch",
                entity_id=batch.id,
                action="DELETE",
                actor_id=user.id,
                old_values={
                    "status": "active", "month": month, "year": year,
                    "source_filename": batch.source_filename, "row_count": batch.row_count,
                },
                new_values={"status": "deleted", "month": month, "year": year},
                reason=payload.reason,
            )
        session.commit()
        return result
    except Exception:
        session.rollback()
        raise


def get_stock_management(
    session: Session,
    month: int,
    year: int,
    search: Optional[str] = None,
    category_filter: Optional[str] = None,
    multi_category_only: bool = False,
) -> StockManagementResponse:
    """Retrieve active stock management parameters for the specified month and year."""
    effective_date = date(year, month, 1)

    batch = session.scalar(
        select(StockParameterBatch).where(
            StockParameterBatch.month == month,
            StockParameterBatch.year == year,
            StockParameterBatch.status == "active",
        ).options(
            joinedload(StockParameterBatch.creator).load_only(User.full_name, User.username),
            joinedload(StockParameterBatch.updater).load_only(User.full_name, User.username),
        )
    )

    if not batch:
        meta = StockManagementMetadata(
            has_parameters=False,
            month=month,
            year=year,
            effective_date=effective_date.isoformat(),
            row_count=0,
        )
        summary = StockManagementSummary()
        return StockManagementResponse(metadata=meta, summary=summary, records=[])

    creator = batch.creator
    updater = batch.updater

    # Fetch all parameters in the batch first to compute multi-category stats accurately
    all_batch_params = session.scalars(
        select(StockParameter)
        .where(StockParameter.batch_id == batch.id)
        .order_by(StockParameter.source_row)
    ).all()

    by_item: dict[str, list[StockParameter]] = {}
    for p in all_batch_params:
        by_item.setdefault(p.item_id.strip().upper(), []).append(p)

    multi_cat_keys: set[str] = set()
    ambiguous_limit_keys: set[str] = set()
    for k, p_list in by_item.items():
        if len(p_list) > 1:
            multi_cat_keys.add(k)
            _, flag_msg = resolve_stock_parameter_for_item(p_list, category=None)
            if flag_msg is not None:
                ambiguous_limit_keys.add(k)

    # Filter for display
    filtered_params = all_batch_params
    if search and search.strip():
        q = search.strip().lower()
        filtered_params = [
            p for p in filtered_params
            if q in (p.item_id or "").lower()
            or q in (p.description or "").lower()
            or q in (p.category or "").lower()
            or q in (p.type_of_material or "").lower()
        ]
    if category_filter and category_filter.strip():
        cf = category_filter.strip().lower()
        filtered_params = [p for p in filtered_params if (p.category or "").strip().lower() == cf]
    if multi_category_only:
        filtered_params = [p for p in filtered_params if p.item_id.strip().upper() in multi_cat_keys]

    missing_limits = sum(
        1 for p in all_batch_params
        if p.min_stock_level is None
        or p.max_stock_level is None
        or p.min_order_qty is None
        or p.lead_time_days is None
    )
    min_gt_max = sum(
        1 for p in all_batch_params
        if p.min_stock_level is not None
        and p.max_stock_level is not None
        and p.min_stock_level > p.max_stock_level
    )
    resolved_rows_count = sum(
        1 for p in all_batch_params
        if (p.resolution_applied or "NONE") != "NONE"
    )

    # Map resolver names
    resolver_ids = {p.resolved_by for p in filtered_params if p.resolved_by}
    resolvers_by_id: dict[uuid.UUID, str] = {}
    if resolver_ids:
        r_users = session.scalars(select(User).where(User.id.in_(resolver_ids))).all()
        for u in r_users:
            resolvers_by_id[u.id] = getattr(u, "full_name", None) or u.username

    records = [
        StockParameterRecord(
            id=str(p.id),
            item_id=p.item_id,
            consumable_id=str(p.consumable_id) if p.consumable_id else None,
            category=p.category,
            type_of_material=p.type_of_material,
            description=p.description,
            min_stock_level=p.min_stock_level,
            max_stock_level=p.max_stock_level,
            min_order_qty=p.min_order_qty,
            lead_time_days=p.lead_time_days,
            edited_after_import=p.edited_after_import,
            source_row=p.source_row,
            resolution_applied=p.resolution_applied or "NONE",
            resolved_by_name=resolvers_by_id.get(p.resolved_by) if p.resolved_by else None,
            resolved_at=p.resolved_at,
            resolution_note=p.resolution_note,
            category_count=len(by_item.get(p.item_id.strip().upper(), [p])),
            has_conflicting_limits=p.item_id.strip().upper() in ambiguous_limit_keys,
        )
        for p in filtered_params
    ]

    batch_summary_dict = batch.summary or {}
    skipped_rows = batch_summary_dict.get("skipped_rows", [])
    resolutions_applied = batch_summary_dict.get("resolutions_applied", [])

    meta = StockManagementMetadata(
        has_parameters=True,
        batch_id=str(batch.id),
        month=month,
        year=year,
        effective_date=batch.effective_date.isoformat(),
        source_filename=batch.source_filename,
        status=batch.status,
        created_at=batch.created_at,
        created_by_name=getattr(creator, "full_name", None) or creator.username if creator else None,
        updated_at=batch.updated_at,
        updated_by_name=getattr(updater, "full_name", None) or updater.username if updater else None,
        row_count=batch.row_count,
        skipped_rows=skipped_rows,
        resolutions_applied=resolutions_applied,
    )

    summary = StockManagementSummary(
        item_count=len(all_batch_params),
        missing_limits_count=missing_limits,
        min_exceeds_max_count=min_gt_max,
        multi_category_items_count=len(multi_cat_keys),
        ambiguous_limits_items_count=len(ambiguous_limit_keys),
        skipped_rows_count=len(skipped_rows),
        resolved_rows_count=resolved_rows_count,
    )

    return StockManagementResponse(metadata=meta, summary=summary, records=records)


def batch_update_stock_parameters(
    session: Session,
    request: Optional[StockParameterBatchUpdateRequest] = None,
    current_user: Optional[CurrentUser] = None,
    payload: Optional[StockParameterBatchUpdateRequest] = None,
    user: Optional[CurrentUser] = None,
) -> StockManagementResponse:
    """Batch-update stock control parameters in one atomic transaction with audit trail."""
    req = request or payload
    usr = current_user or user
    if not req or not usr:
        raise ApplicationError("INVALID_REQUEST", "Missing batch update request or user.", 400)

    try:
        batch = session.scalar(
            select(StockParameterBatch).where(
                StockParameterBatch.month == req.month,
                StockParameterBatch.year == req.year,
                StockParameterBatch.status == "active",
            ).with_for_update().execution_options(populate_existing=True)
        )
        if batch is None:
            raise ApplicationError(
                "NOT_FOUND",
                f"No active stock management batch found for {MONTH_NAMES[req.month - 1]} {req.year}.",
                404,
            )
        if req.batch_id is not None and req.batch_id != batch.id:
            raise ApplicationError("PARAMETERS_CHANGED", "Parameters were replaced. Refresh before saving.", 409)

        now_utc = datetime.now(timezone.utc)
        parameters_by_id = {
            param.id: param
            for param in session.scalars(
                select(StockParameter)
                .where(
                    StockParameter.batch_id == batch.id,
                    StockParameter.id.in_([update.id for update in req.updates]),
                )
                .execution_options(populate_existing=True)
            )
        }
        for update in req.updates:
            param = parameters_by_id.get(update.id)
            if param is None or param.batch_id != batch.id:
                raise ApplicationError(
                    "INVALID_PARAMETER_ROW",
                    "A row does not belong to the selected parameters.",
                    422,
                )
            changes = update.model_dump(exclude_unset=True, exclude={"id"})
            if not changes:
                continue

            min_val = changes.get("min_stock_level", param.min_stock_level)
            max_val = changes.get("max_stock_level", param.max_stock_level)
            if min_val is not None and max_val is not None and min_val > max_val:
                raise ApplicationError(
                    "VALIDATION_ERROR",
                    f"Min stock level ({min_val}) cannot be greater than Max stock level ({max_val}) for {param.item_id}.",
                    400,
                )

            lead_days = changes.get("lead_time_days", param.lead_time_days)
            if lead_days is not None and lead_days < 0:
                raise ApplicationError("VALIDATION_ERROR", f"Lead time cannot be negative for {param.item_id}.", 400)

            before = {key: str(getattr(param, key)) if getattr(param, key) is not None else None for key in changes}
            for key, value in changes.items():
                setattr(param, key, value)
            param.edited_after_import = True
            param.updated_at = now_utc
            param.updated_by = usr.id

            log_audit_event(
                session=session,
                entity_type="StockParameter",
                entity_id=param.id,
                action="UPDATE",
                actor_id=usr.id,
                old_values=before,
                new_values={key: str(value) if value is not None else None for key, value in changes.items()},
                reason=f"Edited stock control limits for {param.item_id} ({param.category or '—'}) in {req.month}/{req.year}.",
            )

        batch.updated_at = now_utc
        batch.updated_by = usr.id

        session.commit()
    except Exception:
        session.rollback()
        raise
    return get_stock_management(session, req.month, req.year)


def list_stock_parameter_history(
    session: Session,
    month: Optional[int] = None,
    year: Optional[int] = None,
    limit: int = 50,
) -> list[StockParameterHistoryItem]:
    """List historical upload batches for stock management."""
    stmt = select(StockParameterBatch)
    if month is not None:
        stmt = stmt.where(StockParameterBatch.month == month)
    if year is not None:
        stmt = stmt.where(StockParameterBatch.year == year)
    stmt = (
        stmt.order_by(StockParameterBatch.created_at.desc())
        .limit(limit)
        .options(joinedload(StockParameterBatch.uploader).load_only(User.full_name, User.username))
    )

    batches = session.scalars(stmt).all()

    items: list[StockParameterHistoryItem] = []
    for b in batches:
        uploader = b.uploader
        items.append(
            StockParameterHistoryItem(
                batch_id=str(b.id),
                month=b.month,
                year=b.year,
                effective_date=b.effective_date.isoformat(),
                source_filename=b.source_filename,
                status=b.status,
                row_count=b.row_count,
                uploaded_at=b.uploaded_at,
                uploaded_by_name=getattr(uploader, "full_name", None) or uploader.username if uploader else None,
            )
        )
    return items


def export_parameter_error_report_excel(
    issues: list[RowIssue],
    filename: str = "stock_management.xlsx",
) -> tuple[bytes, str]:
    """Generate a formatted Excel error/warning report with full row content, conflicts, and hints."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validation Issues"

    headers = [
        "Excel Row",
        "Severity",
        "Issue Type",
        "Column",
        "Item Id",
        "Category",
        "Type of Material",
        "Description",
        "Min Stock Level",
        "Max Stock Level",
        "MOQ",
        "Lead Time (Days)",
        "Conflicts With Row",
        "Problem Description",
        "Suggested Action",
        "Possible Matches",
    ]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="8B0000", end_color="8B0000", fill_type="solid")
    err_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
    warn_fill = PatternFill(start_color="FEF9E7", end_color="FEF9E7", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for iss in issues:
        rc = iss.row_content or {}
        matches_str = "; ".join(
            f"{m.get('item_id')} ({m.get('category') or 'Master'}, Row {m.get('source_row') or '—'}, Min={m.get('min_stock_level')}, Max={m.get('max_stock_level')})"
            for m in (iss.possible_matches or [])
        )
        ws.append([
            iss.row,
            iss.severity,
            iss.code or iss.severity,
            iss.column,
            rc.get("item_id") or iss.item_id or "",
            rc.get("category") or "",
            rc.get("type_of_material") or "",
            rc.get("description") or "",
            rc.get("min_stock_level") if rc.get("min_stock_level") is not None else "",
            rc.get("max_stock_level") if rc.get("max_stock_level") is not None else "",
            rc.get("min_order_qty") if rc.get("min_order_qty") is not None else "",
            rc.get("lead_time_days") if rc.get("lead_time_days") is not None else "",
            iss.conflicting_row if iss.conflicting_row is not None else "",
            iss.description,
            iss.suggested_action or "",
            matches_str,
        ])
        curr_row = ws.max_row
        row_fill = err_fill if iss.severity == "ERROR" else warn_fill
        for col_idx in range(1, len(headers) + 1):
            c = ws.cell(row=curr_row, column=col_idx)
            c.border = thin_border
            c.fill = row_fill

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

    buf = io.BytesIO()
    wb.save(buf)
    clean_base = filename[:-5] if filename.lower().endswith(".xlsx") else filename
    report_name = f"Error_Report_{clean_base}.xlsx"
    return buf.getvalue(), report_name


def export_parameter_import_summary_excel(
    session: Session,
    month: int,
    year: int,
) -> tuple[bytes, str]:
    """Export downloadable Import Summary (.xlsx) listing all skipped/resolved rows and imported parameters."""
    res = get_stock_management(session, month, year)
    if not res.metadata.has_parameters:
        raise ApplicationError("NOT_FOUND", f"No stock parameters found for {MONTH_NAMES[month - 1]} {year}.", 404)

    wb = openpyxl.Workbook()

    # Sheet 1: Skipped & Resolved Rows
    ws_res = wb.active
    ws_res.title = "Skipped & Resolved Rows"

    res_headers = [
        "Excel Row",
        "Action / Resolution",
        "Assigned / Kept Item Id",
        "Original Item Id",
        "Category",
        "Type of Material",
        "Description",
        "Min Stock Level",
        "Max Stock Level",
        "MOQ",
        "Lead Time (Days)",
        "Reason / Details",
        "Resolved By",
        "Resolved At",
    ]
    ws_res.append(res_headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for col_idx in range(1, len(res_headers) + 1):
        cell = ws_res.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for r_item in res.metadata.resolutions_applied:
        rc = r_item.get("row_content") or {}
        ws_res.append([
            r_item.get("row", ""),
            r_item.get("action", ""),
            r_item.get("assigned_item_id") or rc.get("assigned_item_id") or rc.get("item_id") or "",
            rc.get("item_id") or "",
            rc.get("category") or "",
            rc.get("type_of_material") or "",
            rc.get("description") or "",
            rc.get("min_stock_level") if rc.get("min_stock_level") is not None else "",
            rc.get("max_stock_level") if rc.get("max_stock_level") is not None else "",
            rc.get("min_order_qty") if rc.get("min_order_qty") is not None else "",
            rc.get("lead_time_days") if rc.get("lead_time_days") is not None else "",
            r_item.get("reason", ""),
            r_item.get("resolved_by", ""),
            r_item.get("resolved_at", ""),
        ])

    for col in ws_res.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_res.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 50)

    # Sheet 2: Imported Parameters
    ws_imp = wb.create_sheet(title="Imported Parameters")
    imp_headers = EXPECTED_PARAMETER_COLUMNS + ["Excel Row", "Resolution Applied", "Resolved By"]
    ws_imp.append(imp_headers)
    for col_idx in range(1, len(imp_headers) + 1):
        cell = ws_imp.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for r in res.records:
        ws_imp.append([
            r.item_id,
            r.category or "",
            r.type_of_material or "",
            r.description or "",
            float(r.min_stock_level) if r.min_stock_level is not None else "",
            float(r.max_stock_level) if r.max_stock_level is not None else "",
            float(r.min_order_qty) if r.min_order_qty is not None else "",
            r.lead_time_days if r.lead_time_days is not None else "",
            r.source_row,
            r.resolution_applied or "NONE",
            r.resolved_by_name or "",
        ])
        curr_row = ws_imp.max_row
        for col_idx in range(1, len(imp_headers) + 1):
            ws_imp.cell(row=curr_row, column=col_idx).border = thin_border

    for col in ws_imp.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_imp.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 45)

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"KNL_Stock_Parameters_Import_Summary_{year}_{month:02d}.xlsx"
    return buf.getvalue(), filename


def export_stock_management_excel(
    session: Session,
    month: int,
    year: int,
) -> tuple[bytes, str]:
    """Export active stock management parameters to an Excel workbook matching KNL columns."""
    res = get_stock_management(session, month, year)
    if not res.metadata.has_parameters:
        raise ApplicationError("NOT_FOUND", f"No stock parameters found for {MONTH_NAMES[month - 1]} {year}.", 404)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Stock Control {month}-{year}"

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    num_alignment = Alignment(horizontal="right", vertical="center")
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    ws.append([f"Consumables Requirement for the month of {MONTH_NAMES[month - 1]} {year}"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPECTED_PARAMETER_COLUMNS))
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="1F4E79")

    ws.append(EXPECTED_PARAMETER_COLUMNS)
    for col_idx in range(1, len(EXPECTED_PARAMETER_COLUMNS) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in res.records:
        ws.append([
            r.item_id,
            r.category or "",
            r.type_of_material or "",
            r.description or "",
            float(r.min_stock_level) if r.min_stock_level is not None else "",
            float(r.max_stock_level) if r.max_stock_level is not None else "",
            float(r.min_order_qty) if r.min_order_qty is not None else "",
            r.lead_time_days if r.lead_time_days is not None else "",
        ])
        curr_row = ws.max_row
        for col_idx in range(1, len(EXPECTED_PARAMETER_COLUMNS) + 1):
            cell = ws.cell(row=curr_row, column=col_idx)
            cell.border = thin_border
            if col_idx in (5, 6, 7):
                cell.alignment = num_alignment
                cell.number_format = "#,##0.0000"
            elif col_idx in (1, 8):
                cell.alignment = center_alignment

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

    buf = io.BytesIO()
    wb.save(buf)
    filename = f"KNL_Stock_Parameters_{year}_{month:02d}.xlsx"
    return buf.getvalue(), filename
