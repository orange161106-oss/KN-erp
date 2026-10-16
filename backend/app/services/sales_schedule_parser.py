"""KNL Sales Schedule (Production Order from Sales) Two-Pass Parser and Service.

Handles KNL's hierarchical Excel layout ("Pro.Order.pdf"):
- Left columns: S.No, Part No, Item ID, Part Description.
- Quantity columns with 4-row headers: Month label, Schedule given on, Month date, Revision (R0, R1, R2).
- Section / company headings in data rows rather than inline columns.
- Two-pass parser:
    * First block (Ashok Leyland) has NO heading row, identified by closing "ASHOK LEYLAND - TOTAL".
    * Subsequent blocks have opening heading banner and closing "TOTAL" / "- TOTAL" row.
    * Job Work tagged with section_type = 'JOB_WORK'.
    * Closing subtotal reconciliation checks.
    * Grand Total reconciliation check.
- Exact alias matching via customer_aliases table (no fuzzy matching, unknown headings block import).
- Skips INR Value columns ("Value", "Lakhs", "Rs", "INR").
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Optional
from uuid import UUID, uuid4

import openpyxl
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.errors import ApplicationError
from app.models.masters import Customer, CustomerAlias, Product


# ---------------------------------------------------------------------------
# Default Seeding for KNL Customers & Aliases
# ---------------------------------------------------------------------------

DEFAULT_CUSTOMERS = [
    {
        "code": "CUST-AL",
        "name": "Ashok Leyland",
        "aliases": [
            {"alias": "ASHOK LEYLAND - CHILD PARTS", "group": "Child Parts"},
            {"alias": "ASHOK LEYLAND CHILD PARTS", "group": "Child Parts"},
            {"alias": "ASHOK LEYLAND", "group": None},
        ],
    },
    {
        "code": "CUST-TAFE",
        "name": "TAFE",
        "aliases": [
            {"alias": "TAFE", "group": None},
        ],
    },
    {
        "code": "CUST-TAFE-FED",
        "name": "TAFE FED",
        "aliases": [
            {"alias": "TAFE FED Assy Part", "group": "Assy Part"},
            {"alias": "TAFE FED Component", "group": "Component"},
            {"alias": "TAFE FED Child Parts", "group": "Child Parts"},
            {"alias": "TAFE FED", "group": None},
        ],
    },
    {
        "code": "CUST-TVS",
        "name": "TVS Motor Company",
        "aliases": [
            {"alias": "TVS COMPONENT", "group": "Component"},
            {"alias": "TVS CHILD PARTS", "group": "Child Parts"},
            {"alias": "TVS", "group": None},
        ],
    },
    {
        "code": "CUST-MINDA",
        "name": "Minda Corporation",
        "aliases": [
            {"alias": "MINDA", "group": None},
        ],
    },
    {
        "code": "CUST-GREAVES",
        "name": "Greaves Mobility",
        "aliases": [
            {"alias": "GREAVES MOBILITY", "group": None},
        ],
    },
    {
        "code": "CUST-JOBWORK",
        "name": "Job Work",
        "aliases": [
            {"alias": "JOB WORK", "group": "Job Work"},
        ],
    },
]


def normalize_alias(text: str) -> str:
    """Normalize alias text for strict, deterministic matching.
    
    1. Trim whitespace and collapse multiple spaces.
    2. Strip trailing '- TOTAL', 'TOTAL', '-TOTAL' (case-insensitive).
    3. Lowercase.
    """
    if not text:
        return ""
    cleaned = " ".join(str(text).strip().split())
    # Strip trailing total markers
    cleaned = re.sub(r"\s*[-–—]?\s*TOTAL\s*$", "", cleaned, flags=re.IGNORECASE)
    cleaned = " ".join(cleaned.strip().split())
    return cleaned.lower()


def seed_default_customers_and_aliases(session: Session) -> None:
    """Ensure standard KNL customers and verified aliases exist in PostgreSQL."""
    existing_codes = set(session.scalars(select(Customer.code)).all())
    existing_aliases = set(session.scalars(select(CustomerAlias.normalized_alias)).all())

    needed_custs = [c for c in DEFAULT_CUSTOMERS if c["code"] not in existing_codes]
    needed_aliases = []
    seen_norm = set(existing_aliases)
    for c in DEFAULT_CUSTOMERS:
        for a in c["aliases"]:
            norm = normalize_alias(a["alias"])
            if norm not in seen_norm:
                needed_aliases.append((c["code"], a))
                seen_norm.add(norm)
    if not needed_custs and not needed_aliases:
        return

    cust_by_code: dict[str, Customer] = {}
    if existing_codes:
        for c in session.scalars(select(Customer).where(Customer.code.in_(list(existing_codes)))).all():
            cust_by_code[c.code] = c

    for item in needed_custs:
        new_cust = Customer(
            id=uuid4(),
            code=item["code"],
            name=item["name"],
            is_active=True,
        )
        session.add(new_cust)
        cust_by_code[item["code"]] = new_cust
    if needed_custs:
        session.flush()

    for code, alias_data in needed_aliases:
        parent_cust = cust_by_code.get(code)
        if not parent_cust:
            continue
        alias_text = alias_data["alias"]
        norm = normalize_alias(alias_text)
        session.add(
            CustomerAlias(
                id=uuid4(),
                customer_id=parent_cust.id,
                alias_text=alias_text,
                normalized_alias=norm,
                group_label=alias_data["group"],
                is_active=True,
            )
        )
    session.flush()
    session.flush()


# ---------------------------------------------------------------------------
# Date & Period Helpers
# ---------------------------------------------------------------------------

def parse_period_string(value: Any) -> str:
    """Parse various month representations (Oct'26, 01.10.2026, 2026-10) to YYYY-MM."""
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    if not text:
        raise ApplicationError("PLANNING_PERIOD_REQUIRED", "Planning month is missing.", 422)

    # Standard formats
    for fmt in ("%Y-%m", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m")
        except ValueError:
            pass

    # "Oct'26", "Oct-26", "October 2026", "Oct 2026"
    m_match = re.match(r"^([a-zA-Z]{3,9})[\s'’\-_]*([0-9]{2,4})$", text)
    if m_match:
        m_name, y_raw = m_match.groups()
        if len(y_raw) == 2:
            year = 2000 + int(y_raw)
        else:
            year = int(y_raw)
        try:
            dt = datetime.strptime(f"{m_name[:3]} {year}", "%b %Y")
            return dt.strftime("%Y-%m")
        except ValueError:
            pass

    raise ApplicationError("PLANNING_PERIOD_INVALID", f"Could not parse period '{text}'.", 422)


def parse_date_value(value: Any) -> Optional[datetime]:
    """Parse cell value to datetime."""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value or "").strip()
    if not text:
        return None
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%Y-%m"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return None


def is_value_column_header(text: str) -> bool:
    """Check if header denotes an INR monetary value column that must be ignored."""
    lower = str(text or "").lower()
    return any(kw in lower for kw in ("value", "lakhs", "lakh", "inr", "rs", "₹", "rupee"))


# ---------------------------------------------------------------------------
# Sales Schedule Parser
# ---------------------------------------------------------------------------

class QuantityColumnMetadata:
    def __init__(
        self,
        col_index: int,
        month_label: str,
        planning_period: str,
        revision_label: str,
        schedule_given_on: Optional[datetime] = None,
        month_date: Optional[datetime] = None,
    ):
        self.col_index = col_index
        self.month_label = month_label
        self.planning_period = planning_period
        self.revision_label = revision_label.upper()
        self.schedule_given_on = schedule_given_on
        self.month_date = month_date

    @property
    def key(self) -> str:
        return f"{self.planning_period}_{self.revision_label}"

    def to_dict(self) -> dict[str, Any]:
        return {
            "col_index": self.col_index,
            "month_label": self.month_label,
            "planning_period": self.planning_period,
            "revision_label": self.revision_label,
            "schedule_given_on": self.schedule_given_on.strftime("%d.%m.%Y") if self.schedule_given_on else None,
            "month_date": self.month_date.strftime("%d.%m.%Y") if self.month_date else None,
        }


def detect_header_layout(rows: list[tuple[Any, ...]]) -> tuple[int, dict[str, int], list[QuantityColumnMetadata]]:
    """Detect left column positions and all valid revision quantity columns dynamically.

    Does not depend on fixed row 4 or fixed column positions.
    Scans for Part No. and Item ID / Description anywhere in the workbook,
    and inspects the banner above the header dynamically by label and content.
    """
    left_columns: dict[str, int] = {}
    header_row_idx = -1

    def clean_norm(val: Any) -> str:
        return re.sub(r"[^a-z0-9]", "", str(val or "").lower())

    # Scan all rows in the workbook for Part No and Item ID
    for r_idx, row in enumerate(rows):
        norm_cells = [clean_norm(c) for c in row]
        col_map: dict[str, int] = {}
        for c_idx, cell in enumerate(norm_cells):
            if cell in ("sno", "slno", "serialno", "snumber", "no", "sno."):
                col_map["s_no"] = c_idx
            elif any(k in cell for k in ("partno", "partnumber", "productcode", "productpartno", "usedpartno")) or cell == "part":
                col_map["part_number"] = c_idx
            elif any(k in cell for k in ("itemid", "itemcode", "consitemid")) or cell == "item":
                col_map["item_id"] = c_idx
            elif any(k in cell for k in ("partdescription", "description", "itemdescription", "productname", "partdesc", "desc", "usedfor")):
                col_map["description"] = c_idx

        # Found if part_number and item_id both exist, or description with either part_number or item_id
        if ("part_number" in col_map and "item_id" in col_map) or (
            "description" in col_map and ("part_number" in col_map or "item_id" in col_map)
        ):
            header_row_idx = r_idx
            left_columns = col_map
            break

    if header_row_idx == -1:
        # Fallback: scan top 50 rows for description and any part-like header
        for r_idx, row in enumerate(rows[:50]):
            norm_cells = [clean_norm(c) for c in row]
            col_map = {}
            for c_idx, cell in enumerate(norm_cells):
                if "desc" in cell:
                    col_map["description"] = c_idx
                if "part" in cell or "code" in cell:
                    col_map["part_number"] = c_idx
                if "item" in cell:
                    col_map["item_id"] = c_idx
                if "sno" in cell or cell == "no":
                    col_map["s_no"] = c_idx
            if "description" in col_map and ("part_number" in col_map or "item_id" in col_map):
                header_row_idx = r_idx
                left_columns = col_map
                break

    if header_row_idx == -1:
        raise ApplicationError("HEADER_NOT_FOUND", "Could not locate table header with Part No and Item ID / Description.", 422)

    # Defaults for missing left columns if possible
    if "part_number" not in left_columns:
        left_columns["part_number"] = 1
    if "item_id" not in left_columns:
        left_columns["item_id"] = 2
    if "description" not in left_columns:
        left_columns["description"] = 3

    identified_left_cols = set(left_columns.values())
    max_cols = max(len(r) for r in rows) if rows else 0

    quantity_columns: list[QuantityColumnMetadata] = []

    # Detect sheet-wide banner labels from rows above header_row_idx (0..header_row_idx-1)
    banner_rows = rows[:header_row_idx]
    sheet_wide_month_label: Optional[str] = None
    sheet_wide_period: Optional[str] = None
    sheet_wide_schedule_date: Optional[datetime] = None
    sheet_wide_month_date: Optional[datetime] = None

    for b_row in banner_rows:
        for c_val in b_row:
            if c_val is None:
                continue
            str_c = str(c_val).strip()
            if not str_c:
                continue
            try:
                p_cand = parse_period_string(c_val)
                if not sheet_wide_period:
                    sheet_wide_period = p_cand
                    sheet_wide_month_label = str_c
            except ApplicationError:
                pass
            d_cand = parse_date_value(c_val)
            if d_cand:
                if d_cand.day == 1 and not sheet_wide_month_date:
                    sheet_wide_month_date = d_cand
                elif d_cand.day != 1 and not sheet_wide_schedule_date:
                    sheet_wide_schedule_date = d_cand

    # Inspect all candidate columns (not in left_columns)
    for col_idx in range(max_cols):
        if col_idx in identified_left_cols:
            continue

        col_cells = [rows[r][col_idx] if col_idx < len(rows[r]) else None for r in range(header_row_idx + 1)]
        col_text_combined = " ".join(str(c or "") for c in col_cells)

        # Skip value columns (rupees, lakhs, inr)
        if is_value_column_header(col_text_combined):
            continue

        header_val = col_cells[header_row_idx] if header_row_idx < len(col_cells) else None
        header_text = str(header_val or "").strip()

        # Check for revision R0, R1, R2, etc. in header cell or anywhere in this column's banner
        rev_match = re.search(r"\b(R(?:0|[1-9][0-9]*))\b", header_text.upper())
        if not rev_match:
            for c in reversed(col_cells[:header_row_idx]):
                m = re.search(r"\b(R(?:0|[1-9][0-9]*))\b", str(c or "").strip().upper())
                if m:
                    rev_match = m
                    break

        norm_h = clean_norm(header_text)
        is_qty_col = rev_match or "qty" in norm_h or "quantity" in norm_h or "planned" in norm_h
        if not is_qty_col:
            continue

        rev_label = rev_match.group(1) if rev_match else "R0"

        col_month_label: Optional[str] = None
        col_period: Optional[str] = None
        col_sched_date: Optional[datetime] = None
        col_month_date: Optional[datetime] = None

        # Check labeled rows in banner (rows where left columns indicate Month, Schedule date, etc.)
        for r in range(header_row_idx):
            cell_val = col_cells[r]
            if cell_val is None or str(cell_val).strip() == "":
                continue

            left_label_str = " ".join(str(rows[r][lc] or "") for lc in identified_left_cols if lc < len(rows[r])).lower()
            if "schedule" in left_label_str or "given" in left_label_str:
                col_sched_date = parse_date_value(cell_val) or col_sched_date
            elif "month date" in left_label_str or "effective" in left_label_str:
                col_month_date = parse_date_value(cell_val) or col_month_date
            elif "month" in left_label_str or "period" in left_label_str:
                try:
                    col_period = parse_period_string(cell_val)
                    col_month_label = str(cell_val).strip()
                except ApplicationError:
                    pass

            if not col_period:
                try:
                    col_period = parse_period_string(cell_val)
                    col_month_label = str(cell_val).strip()
                except ApplicationError:
                    pass

            d_val = parse_date_value(cell_val)
            if d_val:
                if d_val.day == 1 and not col_month_date:
                    col_month_date = d_val
                elif d_val.day != 1 and not col_sched_date:
                    col_sched_date = d_val

        final_period = col_period or sheet_wide_period
        final_month_label = col_month_label or sheet_wide_month_label or (header_text if "r" not in header_text.lower() else "Sales Schedule")
        if not final_period:
            try:
                final_period = parse_period_string(header_text)
            except ApplicationError:
                final_period = datetime.now().strftime("%Y-%m")

        final_sched_date = col_sched_date or sheet_wide_schedule_date
        final_month_date = col_month_date or sheet_wide_month_date

        quantity_columns.append(
            QuantityColumnMetadata(
                col_index=col_idx,
                month_label=final_month_label,
                planning_period=final_period,
                revision_label=rev_label,
                schedule_given_on=final_sched_date,
                month_date=final_month_date,
            )
        )

    if not quantity_columns:
        raise ApplicationError("QUANTITY_COLUMNS_MISSING", "No valid planned-quantity revision columns found.", 422)

    return header_row_idx, left_columns, quantity_columns


def is_closing_row(text: str) -> bool:
    """Check if cell text denotes a block closing subtotal row."""
    if not text:
        return False
    upper = str(text).strip().upper()
    if "GRAND TOTAL" in upper:
        return False
    return "TOTAL" in upper


def is_grand_total_row(text: str) -> bool:
    """Check if cell text denotes Grand Total row."""
    if not text:
        return False
    upper = str(text).strip().upper()
    return "GRAND TOTAL" in upper


def parse_decimal_safe(value: Any) -> Decimal:
    """Parse numeric cell value to Decimal. Blank or None becomes 0."""
    if value is None or str(value).strip() == "":
        return Decimal("0.0000")
    try:
        d = Decimal(str(value).strip().replace(",", ""))
        if not d.is_finite() or d < 0:
            return Decimal("0.0000")
        return d.quantize(Decimal("0.0001"))
    except (InvalidOperation, ValueError):
        return Decimal("0.0000")


# ---------------------------------------------------------------------------
# Two-Pass Block Extraction & Alias Matching
# ---------------------------------------------------------------------------

def parse_sales_schedule_data(
    rows: list[tuple[Any, ...]],
    session: Session,
    target_period: Optional[str] = None,
    revision_label: Optional[str] = None,
    *, seed_defaults: bool = True,
) -> dict[str, Any]:
    """Parse sales schedule workbook rows using KNL two-pass block parser.
    
    Dynamic: no hardcoded row numbers, counts, or customer positions.
    """
    if seed_defaults:
        seed_default_customers_and_aliases(session)

    header_row_idx, left_cols, qty_cols = detect_header_layout(rows)

    errors: list[str] = []
    warnings: list[str] = []
    needs_review_rows: list[dict[str, Any]] = []

    # 1. Month-wise validation: check target_period against file banner/quantity columns
    active_qty_cols = qty_cols
    if target_period:
        try:
            target_p = parse_period_string(target_period)
            matched_qty_cols = [q for q in qty_cols if q.planning_period == target_p]
            if not matched_qty_cols:
                # Blocking error if selected month differs from file banner month
                file_month = qty_cols[0].planning_period if qty_cols else "Unknown"
                file_month_label = qty_cols[0].month_label if qty_cols else ""
                errors.append(
                    f"Month mismatch: Selected planning month '{target_period}' does not match the file banner month '{file_month}' ({file_month_label}). Please select {file_month} or upload the schedule for {target_period}."
                )
            else:
                active_qty_cols = matched_qty_cols
        except ApplicationError:
            pass

    if revision_label:
        rev_clean = revision_label.strip().upper()
        rev_matched = [q for q in active_qty_cols if q.revision_label == rev_clean]
        if rev_matched:
            active_qty_cols = rev_matched

    if not active_qty_cols:
        active_qty_cols = qty_cols

    # Load active aliases into memory map: normalized_alias -> (Customer, CustomerAlias)
    alias_results = session.execute(
        select(CustomerAlias, Customer)
        .join(Customer, CustomerAlias.customer_id == Customer.id)
        .where(CustomerAlias.is_active.is_(True), Customer.is_active.is_(True))
    ).all()
    alias_map: dict[str, tuple[Customer, CustomerAlias]] = {
        alias.normalized_alias: (cust, alias)
        for alias, cust in alias_results
    }

    # Load active products for part number checks
    active_products = list(session.scalars(select(Product).where(Product.is_active.is_(True))))
    product_code_set = {p.code for p in active_products}
    product_item_set = {p.item_id for p in active_products if p.item_id}
    product_part_set = {p.part_number for p in active_products if p.part_number}

    data_rows = rows[header_row_idx + 1 :]

    # Pass 1: Scan rows to identify blocks and part rows dynamically
    blocks: list[dict[str, Any]] = []
    current_heading: Optional[str] = None
    current_part_rows: list[dict[str, Any]] = []
    grand_total_quantities: dict[str, Decimal] = {}

    for row_offset, row in enumerate(data_rows):
        source_row_num = header_row_idx + 2 + row_offset
        if not any(c is not None and str(c).strip() != "" for c in row):
            continue  # empty row

        def get_val(idx: Optional[int]) -> str:
            if idx is not None and idx < len(row) and row[idx] is not None:
                return str(row[idx]).strip()
            return ""

        s_no = get_val(left_cols.get("s_no"))
        part_no = get_val(left_cols.get("part_number"))
        item_id = get_val(left_cols.get("item_id"))
        description = get_val(left_cols.get("description"))

        row_summary = f"{part_no} {item_id} {description}".strip()

        # Check for Grand Total
        if is_grand_total_row(row_summary) or any(is_grand_total_row(str(c)) for c in row[:5]):
            for q_col in active_qty_cols:
                raw_qty = row[q_col.col_index] if q_col.col_index < len(row) else None
                grand_total_quantities[q_col.key] = parse_decimal_safe(raw_qty)
            break

        # Check for Closing Total Row
        row_has_closing = is_closing_row(row_summary) or any(is_closing_row(str(c)) for c in row[: max(left_cols.values()) + 1])
        if row_has_closing:
            raw_closing_label = description or part_no or (s_no if not re.match(r"^\d+$", s_no) else "") or "TOTAL"
            # If current_heading is None, this closing row identifies the block (e.g. "ASHOK LEYLAND - TOTAL")
            if not current_heading:
                current_heading = re.sub(r"\s*[-–—:]?\s*TOTAL\s*$", "", raw_closing_label, flags=re.IGNORECASE).strip()
            if not current_heading:
                current_heading = "First Block"

            closing_totals: dict[str, Decimal] = {}
            for q_col in active_qty_cols:
                raw_qty = row[q_col.col_index] if q_col.col_index < len(row) else None
                closing_totals[q_col.key] = parse_decimal_safe(raw_qty)

            blocks.append({
                "heading_text": current_heading,
                "closing_label": raw_closing_label,
                "part_rows": current_part_rows,
                "closing_totals": closing_totals,
            })
            current_heading = None
            current_part_rows = []
            continue

        # Check for Heading Row: text in description, part_no, or s_no, no item_id, no numeric S.No, no positive qty
        has_qty = any(
            parse_decimal_safe(row[q_col.col_index] if q_col.col_index < len(row) else None) > 0
            for q_col in active_qty_cols
        )
        candidate_heading = (
            description
            or (part_no if not item_id and not re.match(r"^\d+$", part_no) else "")
            or (s_no if not re.match(r"^\d+$", s_no) else "")
        ).strip()
        is_heading_candidate = (
            bool(candidate_heading)
            and not item_id
            and not has_qty
            and not re.match(r"^\d+$", s_no)
            and not row_has_closing
        )

        if is_heading_candidate:
            if current_part_rows:
                blocks.append({
                    "heading_text": current_heading or "Preceding Block",
                    "closing_label": current_heading or "Preceding Block",
                    "part_rows": current_part_rows,
                    "closing_totals": {},
                })
            current_heading = candidate_heading
            current_part_rows = []
            continue

        # Check for Part Row: Part No or Item ID filled, not containing TOTAL
        if (part_no or item_id) and not row_has_closing:
            quantities: dict[str, Decimal] = {}
            for q_col in active_qty_cols:
                raw_qty = row[q_col.col_index] if q_col.col_index < len(row) else None
                quantities[q_col.key] = parse_decimal_safe(raw_qty)

            is_mapped_part = (
                part_no in product_code_set
                or part_no in product_part_set
                or item_id in product_item_set
                or item_id in product_code_set
            )

            current_part_rows.append({
                "source_row": source_row_num,
                "s_no": s_no,
                "part_no": part_no,
                "item_id": item_id,
                "description": description,
                "quantities": quantities,
                "is_mapped_part": is_mapped_part,
            })
            continue

        # Unrecognized row: record in needs_review_rows
        if any(c is not None and str(c).strip() != "" for c in row):
            needs_review_rows.append({
                "excel_row": source_row_num,
                "text": row_summary[:120],
                "reason": "Unrecognised row format (appears as both heading and part or has irregular structure)",
            })
            if has_qty:
                warnings.append(f"Row {source_row_num} has unrecognized format with quantity; verify reconciliation.")

    # Close any dangling block
    if current_part_rows:
        blocks.append({
            "heading_text": current_heading or "Remaining Parts",
            "closing_label": current_heading or "Remaining Parts",
            "part_rows": current_part_rows,
            "closing_totals": {},
        })

    # Pass 2: Alias Resolution & Reconciliation
    processed_blocks: list[dict[str, Any]] = []
    unmapped_headings: list[str] = []
    all_flat_rows: list[dict[str, Any]] = []

    grand_sum_all: dict[str, Decimal] = {q.key: Decimal("0.0000") for q in active_qty_cols}
    grand_sum_customer: dict[str, Decimal] = {q.key: Decimal("0.0000") for q in active_qty_cols}

    for b_idx, block in enumerate(blocks):
        heading = block["heading_text"]
        norm = normalize_alias(heading)
        section_type = "JOB_WORK" if "job work" in norm else "STANDARD"

        matched_entry = alias_map.get(norm)
        customer_id: Optional[UUID] = None
        customer_name: Optional[str] = None
        customer_code: Optional[str] = None
        group_label: Optional[str] = None
        is_mapped = False

        if matched_entry:
            cust, alias_obj = matched_entry
            customer_id = cust.id
            customer_name = cust.name
            customer_code = cust.code
            group_label = alias_obj.group_label
            is_mapped = True
        else:
            customer_name = heading
            customer_code = ""
            customer_id = None
            group_label = "Job Work" if section_type == "JOB_WORK" else None
            is_mapped = False
            if heading not in unmapped_headings:
                unmapped_headings.append(heading)
            warnings.append(f"Unknown company heading detected: '{heading}'. Map this heading to an existing Customer Master record.")

        # Block quantity reconciliation against ITS OWN closing row
        block_sums: dict[str, Decimal] = {q.key: Decimal("0.0000") for q in active_qty_cols}
        for p_row in block["part_rows"]:
            for q_key, q_val in p_row["quantities"].items():
                block_sums[q_key] += q_val
                grand_sum_all[q_key] += q_val
                if section_type != "JOB_WORK":
                    grand_sum_customer[q_key] += q_val

        reconciliation_ok = True
        reconciliation_diffs: dict[str, str] = {}
        for q_col in active_qty_cols:
            q_key = q_col.key
            part_sum = block_sums[q_key]
            sheet_total = block["closing_totals"].get(q_key)
            if sheet_total is not None and sheet_total > 0:
                diff = abs(part_sum - sheet_total)
                if diff > Decimal("0.0001"):
                    reconciliation_ok = False
                    reconciliation_diffs[q_key] = str(diff)
                    warnings.append(
                        f"Block '{heading}' sum ({part_sum}) differs from sheet closing total ({sheet_total}) for {q_col.revision_label}. Diff: {diff}."
                    )

        # Build flat rows with Customer, Customer ID, Group
        for p_row in block["part_rows"]:
            for q_col in active_qty_cols:
                q_key = q_col.key
                qty = p_row["quantities"].get(q_key, Decimal("0.0000"))
                all_flat_rows.append({
                    "source_row": p_row["source_row"],
                    "customer_id": str(customer_id) if customer_id else None,
                    "customer": customer_name or heading,
                    "customer_code": customer_code or "",
                    "customer_group": group_label or ("Job Work" if section_type == "JOB_WORK" else None),
                    "product_code": p_row["part_no"] or p_row["item_id"],
                    "item_id": p_row["item_id"],
                    "part_no": p_row["part_no"],
                    "description": p_row["description"],
                    "planned_quantity": format(qty, ".4f"),
                    # This schedule layout has no production-unit column.
                    # The importer must use the matched Product Master unit.
                    "uom": "",
                    "target_period": q_col.planning_period,
                    "planning_version": q_col.revision_label,
                    "revision": q_col.revision_label,
                    "schedule_given_on": q_col.schedule_given_on.isoformat() if q_col.schedule_given_on else None,
                    "section_type": section_type,
                    "is_mapped_part": p_row["is_mapped_part"],
                })

        processed_blocks.append({
            "block_index": b_idx + 1,
            "heading_text": heading,
            "closing_label": block["closing_label"],
            "customer_name": customer_name,
            "customer_code": customer_code,
            "customer_id": str(customer_id) if customer_id else None,
            "group_label": group_label,
            "section_type": section_type,
            "part_count": len(block["part_rows"]),
            "quantities": {k: format(v, ".4f") for k, v in block_sums.items()},
            "sheet_totals": {k: format(v, ".4f") for k, v in block["closing_totals"].items()},
            "reconciliation_ok": reconciliation_ok,
            "reconciliation_diff": reconciliation_diffs,
            "is_mapped": is_mapped,
        })

    # Grand Total Reconciliation
    grand_reconciliation_ok = True
    grand_diffs: dict[str, str] = {}
    for q_col in active_qty_cols:
        q_key = q_col.key
        parts_total = grand_sum_all[q_key]
        cust_total = grand_sum_customer[q_key]
        sheet_grand = grand_total_quantities.get(q_key)
        if sheet_grand is not None and sheet_grand > 0:
            diff_all = abs(parts_total - sheet_grand)
            diff_cust = abs(cust_total - sheet_grand)
            best_diff = min(diff_all, diff_cust)
            if best_diff > Decimal("0.0001"):
                grand_reconciliation_ok = False
                grand_diffs[q_key] = str(best_diff)
                warnings.append(
                    f"Grand total check failed for {q_col.revision_label}: Sum of blocks ({parts_total}) does not match sheet Grand Total ({sheet_grand}). Diff: {best_diff}."
                )

    # Check for duplicate parts within blocks
    for b in processed_blocks:
        seen_parts = set()
        block_parts = [r for r in all_flat_rows if r["customer"] == (b["customer_name"] or b["heading_text"])]
        for r in block_parts:
            key = (r["product_code"], r["target_period"], r["revision"])
            if key in seen_parts:
                warnings.append(f"Duplicate part '{r['product_code']}' in block '{b['heading_text']}'.")
            seen_parts.add(key)

    can_confirm = len(errors) == 0 and len(unmapped_headings) == 0 and len(all_flat_rows) > 0

    return {
        "header_row_index": header_row_idx + 1,
        "available_revisions": [q.to_dict() for q in qty_cols],
        "active_revisions": [q.to_dict() for q in active_qty_cols],
        "blocks": processed_blocks,
        "unmapped_aliases": unmapped_headings,
        "grand_total": {
            "parts_sum": {k: format(v, ".4f") for k, v in grand_sum_all.items()},
            "sheet_grand_total": {k: format(v, ".4f") for k, v in grand_total_quantities.items()},
            "reconciliation_ok": grand_reconciliation_ok,
            "diffs": grand_diffs,
        },
        "total_part_rows": len(all_flat_rows),
        "sample_rows": all_flat_rows[:50],
        "needs_review_rows": needs_review_rows,
        "can_confirm": can_confirm,
        "warnings": warnings,
        "errors": errors,
        "_all_flat_rows": all_flat_rows,
    }


def read_workbook_rows(file_bytes: bytes, filename: str, sheet_name: Optional[str] = None) -> tuple[str, list[tuple[Any, ...]]]:
    """Safely open workbook and extract rows from chosen sheet."""
    if not file_bytes:
        raise ApplicationError("EMPTY_FILE", "Uploaded file is empty.", 422)

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                raise ApplicationError("SOURCE_TOO_LARGE", "Expanded workbook exceeds safe limit (200MB).", 422)
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except ApplicationError:
        raise
    except Exception as exc:
        raise ApplicationError("INVALID_EXCEL", f"Could not read Excel file: {str(exc)}", 422)

    try:
        target_sheet = sheet_name or wb.sheetnames[0]
        if target_sheet not in wb.sheetnames:
            raise ApplicationError("SHEET_NOT_FOUND", f"Sheet '{target_sheet}' not found in workbook.", 422)
        sheet = wb[target_sheet]
        rows = list(sheet.iter_rows(values_only=True))
        return target_sheet, rows
    finally:
        wb.close()

