import hashlib
import io
import time
import zipfile
from collections import defaultdict
import logging
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Optional
from uuid import UUID, uuid4

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import func, or_, select, update
from sqlalchemy.orm import Session

logger = logging.getLogger(__name__)

from app.core.errors import ApplicationError
from app.models.mappings import ProductPlant, ProductProcessConsumable
from app.models.masters import Customer, CustomerAlias, Product
from app.models.rules import ConsumptionNorm
from app.models.prd import ImportBatch, PlanningVersion, PRDOrderHeader, PRDOrderItem, PRDRecord
from app.models.production import Plant, Route, RouteStep
from app.services.basic_norms import match_product
from app.schemas.prd_workspace import (
    PRDExcelImportResponse,
    PRDExcelInspectResponse,
    PRDExcelSheetInspectInfo,
    PRDWorkspaceRecordResponse,
    PRDWorkspaceSaveRequest,
    PRDWorkspaceSaveResponse,
    SalesScheduleBlockInfo,
    SalesScheduleGrandTotal,
    SalesScheduleImportPreviewResponse,
    SalesScheduleRevisionInfo,
)
from app.services.prd import promote_batch_to_planning_version
from app.services.purchase_order import audit
from app.services.sales_schedule_parser import (
    normalize_alias,
    parse_sales_schedule_data,
    read_workbook_rows,
    seed_default_customers_and_aliases,
)

# In-memory storage for staging uploaded files between preview and confirmation
_STAGED_FILES: dict[str, tuple[bytes, str, float]] = {}


def _cleanup_stale_staged_files() -> None:
    now = time.time()
    stale_keys = [k for k, (_, _, t) in _STAGED_FILES.items() if now - t > 3600]
    for k in stale_keys:
        _STAGED_FILES.pop(k, None)


def to_prd_record_response(row: PRDRecord) -> PRDWorkspaceRecordResponse:
    return PRDWorkspaceRecordResponse(
        id=row.id,
        row_index=row.row_index,
        plant=row.plant,
        customer=row.customer,
        customer_id=row.customer_id,
        customer_code=row.customer_code,
        customer_group=row.customer_group,
        product_code=row.product_code,
        description=row.description,
        planned_quantity=format(row.planned_quantity, ".4f") if row.planned_quantity is not None else "0.0000",
        uom=row.uom or "Nos",
        target_period=row.target_period,
        planning_version=row.planning_version or "V1",
        revision=row.revision or row.planning_version or "R0",
        schedule_given_on=row.schedule_given_on,
        section_type=row.section_type or "STANDARD",
        status=row.status or "SAVED",
        remarks=row.remarks,
        created_at=row.created_at,
        updated_at=row.updated_at,
    )


def list_prd_records(
    session: Session,
    search: Optional[str] = None,
    plant: Optional[str] = None,
    target_period: Optional[str] = None,
    planning_version: Optional[str] = None,
    status: Optional[str] = None,
    customer: Optional[str] = None,
    limit: int = 500,
    offset: int = 0,
) -> list[PRDWorkspaceRecordResponse]:
    query = select(PRDRecord).where(PRDRecord.is_deleted.is_(False))

    if plant and plant.upper() != "ALL":
        query = query.where(PRDRecord.plant.ilike(f"%{plant.strip()}%"))
    if target_period and target_period.upper() != "ALL":
        query = query.where(PRDRecord.target_period == target_period.strip())
    if planning_version and planning_version.upper() != "ALL":
        query = query.where(
            or_(
                PRDRecord.planning_version == planning_version.strip(),
                PRDRecord.revision == planning_version.strip(),
            )
        )
    if status and status.upper() != "ALL":
        query = query.where(PRDRecord.status == status.upper())
    if customer and customer.upper() != "ALL":
        query = query.where(
            or_(
                PRDRecord.customer.ilike(f"%{customer.strip()}%"),
                PRDRecord.customer_code.ilike(f"%{customer.strip()}%"),
            )
        )

    if search:
        s = f"%{search.strip()}%"
        query = query.where(
            or_(
                PRDRecord.product_code.ilike(s),
                PRDRecord.description.ilike(s),
                PRDRecord.plant.ilike(s),
                PRDRecord.customer.ilike(s),
                PRDRecord.customer_code.ilike(s),
                PRDRecord.customer_group.ilike(s),
                PRDRecord.target_period.ilike(s),
                PRDRecord.planning_version.ilike(s),
                PRDRecord.remarks.ilike(s),
            )
        )

    rows = session.scalars(
        query.order_by(PRDRecord.row_index.asc().nulls_last(), PRDRecord.created_at.asc())
        .limit(limit)
        .offset(offset)
    ).all()

    return [to_prd_record_response(r) for r in rows]


def save_prd_records(
    session: Session,
    data: PRDWorkspaceSaveRequest,
    actor: UUID,
) -> PRDWorkspaceSaveResponse:
    saved_records = []

    # Process deletions
    for del_id in data.deleted_ids:
        row = session.get(PRDRecord, del_id)
        if row and not row.is_deleted:
            if row.status == "VALIDATED":
                raise ApplicationError("PRD_HISTORY_PROTECTED", "Import a new revision instead of deleting validated history.", 409)
            row.is_deleted = True
            row.updated_at = datetime.now(timezone.utc)
            audit(
                session, actor, row.id, "DELETE_ROW", data.reason,
                {"id": str(row.id), "product_code": row.product_code},
                old={"status": row.status, "is_deleted": False},
                entity="prd_records"
            )

    # Process additions / updates
    for item in data.records:
        try:
            qty = Decimal(str(item.planned_quantity).strip())
            if qty < 0:
                raise ValueError("Quantity must be non-negative")
        except (InvalidOperation, ValueError):
            raise ApplicationError(
                "INVALID_QUANTITY",
                f"Invalid planned quantity '{item.planned_quantity}' for part {item.product_code}.",
                422,
            )

        existing = session.get(PRDRecord, item.id) if item.id else None
        if existing and not existing.is_deleted:
            if existing.status == "VALIDATED":
                raise ApplicationError("PRD_HISTORY_PROTECTED", "Import a new revision instead of editing validated history.", 409)
            old_vals = {
                "product_code": existing.product_code,
                "plant": existing.plant,
                "planned_quantity": str(existing.planned_quantity),
                "status": existing.status,
            }
            existing.plant = item.plant.strip()
            existing.customer = item.customer.strip() if item.customer else None
            existing.customer_id = item.customer_id
            existing.customer_code = item.customer_code
            existing.customer_group = item.customer_group
            existing.product_code = item.product_code.strip()
            existing.description = item.description.strip()
            existing.planned_quantity = qty
            existing.uom = item.uom.strip() or "Nos"
            existing.target_period = item.target_period.strip()
            existing.planning_version = item.planning_version.strip() or "V1"
            existing.revision = item.revision or item.planning_version or "R0"
            existing.section_type = item.section_type or "STANDARD"
            existing.status = item.status or "SAVED"
            existing.remarks = item.remarks
            existing.row_index = item.row_index
            existing.updated_at = datetime.now(timezone.utc)
            saved_records.append(existing)
            audit(
                session, actor, existing.id, "UPDATE_ROW", data.reason,
                {"id": str(existing.id), "product_code": item.product_code, "planned_quantity": str(qty)},
                old=old_vals, entity="prd_records"
            )
        else:
            new_row = PRDRecord(
                id=item.id or uuid4(),
                row_index=item.row_index,
                plant=item.plant.strip(),
                customer=item.customer.strip() if item.customer else None,
                customer_id=item.customer_id,
                customer_code=item.customer_code,
                customer_group=item.customer_group,
                product_code=item.product_code.strip(),
                description=item.description.strip(),
                planned_quantity=qty,
                uom=item.uom.strip() or "Nos",
                target_period=item.target_period.strip(),
                planning_version=item.planning_version.strip() or "V1",
                revision=item.revision or item.planning_version or "R0",
                section_type=item.section_type or "STANDARD",
                status="SAVED",
                remarks=item.remarks,
                created_by=actor,
            )
            session.add(new_row)
            saved_records.append(new_row)
            audit(
                session, actor, new_row.id, "CREATE_ROW", data.reason,
                {"id": str(new_row.id), "product_code": item.product_code, "planned_quantity": str(qty)},
                entity="prd_records"
            )

    session.commit()
    return PRDWorkspaceSaveResponse(
        saved_count=len(saved_records),
        deleted_count=len(data.deleted_ids),
        records=[to_prd_record_response(r) for r in saved_records],
    )


def delete_prd_record(session: Session, identity: UUID, actor: UUID, reason: str = "Row deleted") -> None:
    row = session.get(PRDRecord, identity)
    if not row or row.is_deleted:
        raise ApplicationError("RECORD_NOT_FOUND", "PRD planning record not found.", 404)
    if row.status == "VALIDATED":
        raise ApplicationError("PRD_HISTORY_PROTECTED", "Validated revision history cannot be deleted.", 409)
    row.is_deleted = True
    row.updated_at = datetime.now(timezone.utc)
    audit(
        session, actor, row.id, "DELETE_ROW", reason,
        {"id": str(row.id), "product_code": row.product_code},
        old={"is_deleted": False}, entity="prd_records"
    )
    session.commit()


def bulk_delete_prd_records(
    session: Session,
    ids: list[UUID] | None = None,
    delete_all_matching: bool = False,
    search: str | None = None,
    plant: str | None = None,
    target_period: str | None = None,
    planning_version: str | None = None,
    status: str | None = None,
    actor: UUID | None = None,
    reason: str = "Bulk delete PRD records",
) -> int:
    query = select(PRDRecord).where(PRDRecord.is_deleted.is_(False))
    if delete_all_matching:
        if plant and plant.upper() != "ALL":
            query = query.where(PRDRecord.plant.ilike(f"%{plant.strip()}%"))
        if target_period and target_period.upper() != "ALL":
            query = query.where(PRDRecord.target_period == target_period.strip())
        if planning_version and planning_version.upper() != "ALL":
            query = query.where(PRDRecord.planning_version == planning_version.strip())
        if status and status.upper() != "ALL":
            query = query.where(PRDRecord.status == status.upper())
        if search and search.strip():
            s = f"%{search.strip()}%"
            query = query.where(
                or_(
                    PRDRecord.product_code.ilike(s),
                    PRDRecord.description.ilike(s),
                    PRDRecord.plant.ilike(s),
                    PRDRecord.customer.ilike(s),
                    PRDRecord.target_period.ilike(s),
                    PRDRecord.planning_version.ilike(s),
                    PRDRecord.remarks.ilike(s),
                )
            )
    elif ids:
        query = query.where(PRDRecord.id.in_(ids))
    else:
        return 0

    records_to_delete = session.scalars(query).all()
    if any(r.status == "VALIDATED" for r in records_to_delete):
        raise ApplicationError("PRD_HISTORY_PROTECTED", "Validated revision history cannot be deleted.", 409)
    count = len(records_to_delete)
    if count == 0:
        return 0

    target_ids = [r.id for r in records_to_delete]
    session.execute(
        update(PRDRecord)
        .where(PRDRecord.id.in_(target_ids))
        .values(is_deleted=True, updated_at=datetime.now(timezone.utc))
    )

    if actor:
        audit(
            session, actor, uuid4(), "BULK_DELETE_ROWS", reason,
            {"deleted_count": count, "deleted_ids": [str(i) for i in target_ids[:50]]},
            entity="prd_records"
        )

    session.commit()
    return count


def inspect_prd_excel(file_bytes: bytes, filename: str) -> PRDExcelInspectResponse:
    if not file_bytes:
        raise ApplicationError("EMPTY_FILE", "The uploaded file is empty.", 422)

    sheets_info = []
    if filename.lower().endswith(".csv"):
        import csv
        try:
            text = file_bytes.decode("utf-8", errors="replace")
            csv_rows = list(csv.reader(io.StringIO(text)))
        except Exception as exc:
            raise ApplicationError("INVALID_CSV", f"Could not read CSV file: {str(exc)}", 422)

        header_row = [str(c).strip() for c in csv_rows[0]] if csv_rows else []
        data_rows = csv_rows[1:] if len(csv_rows) > 1 else []
        sample = []
        for d_row in data_rows[:5]:
            row_dict = {}
            for idx, h in enumerate(header_row):
                val = d_row[idx] if idx < len(d_row) else ""
                row_dict[h] = str(val) if val is not None else ""
            sample.append(row_dict)

        sheets_info.append(PRDExcelSheetInspectInfo(
            name="Sheet1",
            row_count=len(data_rows),
            column_count=len(header_row),
            headers=header_row,
            sample_rows=sample,
        ))
    else:
        try:
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                    raise ApplicationError('SOURCE_TOO_LARGE', 'The expanded workbook exceeds the import limit.', 422)
            workbook = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except ApplicationError:
            raise
        except Exception:
            raise ApplicationError("INVALID_EXCEL", "Could not read Excel workbook.", 422) from None

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            if not hasattr(sheet, "iter_rows"):
                continue
            rows = list(sheet.iter_rows(max_row=20, max_col=min(sheet.max_column, 128), values_only=True))
            header_row = None
            data_rows = []
            for r in rows:
                if any(c is not None and str(c).strip() != "" for c in r):
                    if header_row is None:
                        header_row = [str(c).strip() if c is not None else f"Col_{i+1}" for i, c in enumerate(r)]
                    else:
                        data_rows.append(r)

            headers = header_row or []
            sample = []
            for d_row in data_rows[:5]:
                row_dict = {}
                for idx, h in enumerate(headers):
                    val = d_row[idx] if idx < len(d_row) else ""
                    row_dict[h] = str(val) if val is not None else ""
                sample.append(row_dict)

            sheets_info.append(PRDExcelSheetInspectInfo(
                name=sheet_name,
                row_count=max(0, sheet.max_row - 1),
                column_count=sheet.max_column,
                headers=headers,
                sample_rows=sample,
            ))
        workbook.close()

    if not sheets_info:
        raise ApplicationError("NO_DATA_SHEETS", "Workbook contains no readable sheets.", 422)

    return PRDExcelInspectResponse(filename=filename, sheets=sheets_info)


# ---------------------------------------------------------------------------
# Sales Schedule Multi-Step Preview & Confirmation
# ---------------------------------------------------------------------------

def _resolve_route_for_product(
    session: Session,
    product_id: UUID,
    plant_id: Optional[UUID] = None,
    product_plants_by_prod: Optional[dict[UUID, list[ProductPlant]]] = None,
    routes_by_id: Optional[dict[UUID, Route]] = None,
    ppc_proc_by_prod: Optional[dict[UUID, set[UUID]]] = None,
    route_proc_map: Optional[dict[UUID, set[UUID]]] = None,
) -> Optional[UUID]:
    """Deterministically resolve a valid production route associated with a product.

    1. First, check active ProductPlant mappings for this product across any plant.
       Prioritize mappings for the same plant, then primary mappings, then earliest created.
    2. Second, check if the product has assigned processes via ProductProcessConsumable,
       and find an active Route containing those processes via RouteStep.
    3. Return None if no valid route is deterministically associated.
    """
    if product_plants_by_prod is not None and routes_by_id is not None:
        pps = product_plants_by_prod.get(product_id, [])
        sorted_pps = sorted(
            pps,
            key=lambda pp: (
                0 if (plant_id and pp.plant_id == plant_id) else 1,
                0 if pp.is_primary else 1,
                pp.created_at,
            ),
        )
        for pp in sorted_pps:
            if pp.route_id and pp.route_id in routes_by_id:
                route = routes_by_id[pp.route_id]
                if route.is_active:
                    return route.id

        if ppc_proc_by_prod is not None and route_proc_map is not None:
            proc_ids = ppc_proc_by_prod.get(product_id, set())
            if proc_ids:
                matching_routes = [
                    r for r_id, r in routes_by_id.items()
                    if r.is_active and proc_ids.intersection(route_proc_map.get(r_id, set()))
                ]
                if matching_routes:
                    matching_routes.sort(key=lambda r: r.name or "")
                    return matching_routes[0].id
        return None

    # Fallback to direct DB query if in-memory dicts not provided
    stmt = (
        select(ProductPlant)
        .where(
            ProductPlant.product_id == product_id,
            ProductPlant.is_active.is_(True),
        )
    )
    if plant_id:
        stmt = stmt.order_by(
            (ProductPlant.plant_id == plant_id).desc(),
            ProductPlant.is_primary.desc(),
            ProductPlant.created_at.asc(),
        )
    else:
        stmt = stmt.order_by(
            ProductPlant.is_primary.desc(),
            ProductPlant.created_at.asc(),
        )

    for pp in session.scalars(stmt).all():
        if pp.route_id:
            route = session.get(Route, pp.route_id)
            if route and route.is_active:
                return route.id

    ppc_process_ids = list(session.scalars(
        select(ProductProcessConsumable.process_id)
        .where(
            ProductProcessConsumable.product_id == product_id,
            ProductProcessConsumable.is_active.is_(True),
        )
    ).all())
    if ppc_process_ids:
        candidate_route = session.scalars(
            select(Route)
            .join(RouteStep, Route.id == RouteStep.route_id)
            .where(
                Route.is_active.is_(True),
                RouteStep.process_id.in_(ppc_process_ids),
            )
            .order_by(Route.name.asc())
        ).first()
        if candidate_route:
            return candidate_route.id

    return None


def _lookup_or_create_product(
    session: Session,
    product_cache: dict[str, Product],
    p_code: str,
    part_no: Optional[str],
    item_id: Optional[str],
    description: Optional[str],
    uom: Optional[str],
) -> Product:
    """Find existing product or create new one safely without violating uniqueness constraints."""
    norm_part = part_no.strip() if part_no else None
    norm_item = item_id.strip() if item_id else None
    norm_code = p_code.strip() if p_code else None

    existing_prod: Optional[Product] = None
    if norm_part and norm_item:
        existing_prod = product_cache.get(f"pair:{norm_part.upper()}:{norm_item.upper()}")
    if not existing_prod and norm_code:
        existing_prod = product_cache.get(f"code:{norm_code.upper()}")
    if not existing_prod and norm_part:
        existing_prod = product_cache.get(f"part:{norm_part.upper()}")
    if not existing_prod and norm_item:
        existing_prod = product_cache.get(f"item:{norm_item.upper()}")

    if not existing_prod:
        # 1. Exact match on (part_number, item_id)
        if norm_part and norm_item:
            existing_prod = session.scalar(
                select(Product).where(
                    func.lower(func.trim(Product.part_number)) == norm_part.lower(),
                    func.lower(func.trim(Product.item_id)) == norm_item.lower(),
                )
            )

        # 2. Composite code match
        if not existing_prod and norm_part and norm_item:
            existing_prod = session.scalar(
                select(Product).where(
                    Product.code.in_([
                        f"{norm_part}_{norm_item}",
                        f"{norm_part}-{norm_item}",
                        f"{norm_item}_{norm_part}",
                    ])
                )
            )

        # 3. Match on code
        if not existing_prod and norm_code:
            existing_prod = session.scalar(
                select(Product).where(func.lower(func.trim(Product.code)) == norm_code.lower())
            )

        # 4. Match on part_number alone if unique
        if not existing_prod and norm_part:
            matches = list(session.scalars(
                select(Product).where(func.lower(func.trim(Product.part_number)) == norm_part.lower())
            ).all())
            if len(matches) == 1:
                existing_prod = matches[0]

        # 5. Match on item_id alone if unique
        if not existing_prod and norm_item:
            matches = list(session.scalars(
                select(Product).where(func.lower(func.trim(Product.item_id)) == norm_item.lower())
            ).all())
            if len(matches) == 1:
                existing_prod = matches[0]

    if existing_prod:
        if not existing_prod.is_active:
            existing_prod.is_active = True
        if norm_part and not existing_prod.part_number:
            target_item = existing_prod.item_id or norm_item
            if target_item:
                collision = session.scalar(
                    select(Product.id).where(
                        Product.id != existing_prod.id,
                        func.lower(func.trim(Product.part_number)) == norm_part.lower(),
                        func.lower(func.trim(Product.item_id)) == target_item.lower(),
                    )
                )
                if not collision:
                    existing_prod.part_number = norm_part
            else:
                existing_prod.part_number = norm_part

        if norm_item and not existing_prod.item_id:
            target_part = existing_prod.part_number or norm_part
            if target_part:
                collision = session.scalar(
                    select(Product.id).where(
                        Product.id != existing_prod.id,
                        func.lower(func.trim(Product.part_number)) == target_part.lower(),
                        func.lower(func.trim(Product.item_id)) == norm_item.lower(),
                    )
                )
                if not collision:
                    existing_prod.item_id = norm_item
            else:
                existing_prod.item_id = norm_item
    else:
        if norm_part and norm_item:
            dupe = session.scalar(
                select(Product).where(
                    func.lower(func.trim(Product.part_number)) == norm_part.lower(),
                    func.lower(func.trim(Product.item_id)) == norm_item.lower(),
                )
            )
            if dupe:
                existing_prod = dupe
                if not existing_prod.is_active:
                    existing_prod.is_active = True

        if not existing_prod:
            code_cand = norm_code or (f"{norm_part}_{norm_item}" if norm_part and norm_item else (norm_part or norm_item or "PROD"))
            code_exists = session.scalar(select(Product.id).where(Product.code == code_cand))
            if code_exists:
                code_cand = f"{code_cand}_{uuid4().hex[:6]}"

            existing_prod = Product(
                code=code_cand,
                name=description or code_cand,
                uom=uom or "PCS",
                part_number=norm_part,
                item_id=norm_item,
                is_active=True,
            )
            session.add(existing_prod)
            session.flush()

    if norm_part and norm_item:
        product_cache[f"pair:{norm_part.upper()}:{norm_item.upper()}"] = existing_prod
    if norm_code:
        product_cache[f"code:{norm_code.upper()}"] = existing_prod
    if existing_prod.code:
        product_cache[f"code:{existing_prod.code.upper()}"] = existing_prod
    if norm_part:
        product_cache[f"part:{norm_part.upper()}"] = existing_prod
    if norm_item:
        product_cache[f"item:{norm_item.upper()}"] = existing_prod
    product_cache[str(existing_prod.id)] = existing_prod

    return existing_prod


def get_previous_period_str(period: str) -> str:
    try:
        parts = period.split("-")
        y, m = int(parts[0]), int(parts[1])
        m -= 1
        if m == 0:
            m = 12
            y -= 1
        return f"{y:04d}-{m:02d}"
    except Exception:
        return period


def compute_changes_vs_previous_month(
    session: Session,
    current_period: str,
    current_flat_rows: list[dict[str, Any]],
    current_blocks: list[dict[str, Any]],
) -> dict[str, Any]:
    prev_period = get_previous_period_str(current_period)

    # 1. Fetch previous month records from PRDRecord
    prev_records = session.scalars(
        select(PRDRecord).where(
            PRDRecord.target_period == prev_period,
            PRDRecord.is_deleted.is_(False),
        )
    ).all()

    prev_map: dict[tuple[str, str, str], Decimal] = {}
    prev_blocks: set[str] = set()

    for r in prev_records:
        cust = (r.customer or "").strip().lower()
        part = (r.product_code or "").strip().lower()
        item = ""
        prev_map[(cust, part, item)] = r.planned_quantity
        if r.customer:
            prev_blocks.add(r.customer.strip())

    # If no PRDRecords found, query PRDOrderItem
    if not prev_records:
        prev_items = session.execute(
            select(PRDOrderItem, Customer.name, Product.part_number, Product.item_id)
            .outerjoin(Customer, PRDOrderItem.customer_id == Customer.id)
            .outerjoin(Product, PRDOrderItem.product_id == Product.id)
            .where(PRDOrderItem.target_period == prev_period)
        ).all()
        for item_obj, c_name, p_no, i_id in prev_items:
            cust = (c_name or "").strip().lower()
            part = (p_no or item_obj.product_code or "").strip().lower()
            it = (i_id or "").strip().lower()
            prev_map[(cust, part, it)] = item_obj.planned_quantity
            if c_name:
                prev_blocks.add(c_name.strip())

    # Build current month map (group by block/customer, part_no, item_id)
    current_map: dict[tuple[str, str, str], dict[str, Any]] = {}
    curr_blocks: set[str] = set()

    for r in current_flat_rows:
        cust = (r.get("customer") or "").strip()
        part = (r.get("part_no") or r.get("product_code") or "").strip()
        item = (r.get("item_id") or "").strip()
        key = (cust.lower(), part.lower(), item.lower())
        qty = Decimal(str(r.get("planned_quantity") or "0"))
        if key not in current_map:
            current_map[key] = {
                "customer": cust,
                "part_no": part,
                "item_id": item,
                "description": r.get("description", ""),
                "planned_quantity": str(qty),
                "qty_dec": qty,
            }
        if cust:
            curr_blocks.add(cust)

    for b in current_blocks:
        b_name = b.get("customer_name") or b.get("heading_text")
        if b_name:
            curr_blocks.add(b_name.strip())

    parts_added = []
    parts_removed = []
    parts_changed = []

    for key, c_data in current_map.items():
        old_qty = prev_map.get(key)
        if old_qty is None and (key[0], key[1], "") in prev_map:
            old_qty = prev_map[(key[0], key[1], "")]

        if old_qty is None:
            parts_added.append({
                "customer": c_data["customer"],
                "part_no": c_data["part_no"],
                "item_id": c_data["item_id"],
                "description": c_data["description"],
                "planned_quantity": c_data["planned_quantity"],
            })
        else:
            new_qty = c_data["qty_dec"]
            if abs(new_qty - old_qty) > Decimal("0.0001"):
                parts_changed.append({
                    "customer": c_data["customer"],
                    "part_no": c_data["part_no"],
                    "item_id": c_data["item_id"],
                    "description": c_data["description"],
                    "old_quantity": format(old_qty, ".4f"),
                    "new_quantity": format(new_qty, ".4f"),
                })

    for key, old_qty in prev_map.items():
        found = key in current_map or (key[0], key[1], "") in current_map
        if not found:
            parts_removed.append({
                "customer": key[0].title(),
                "part_no": key[1].upper(),
                "item_id": key[2].upper(),
                "planned_quantity": format(old_qty, ".4f"),
            })

    blocks_added = sorted(list(curr_blocks - prev_blocks))
    blocks_removed = sorted(list(prev_blocks - curr_blocks))

    return {
        "current_period": current_period,
        "previous_period": prev_period,
        "has_previous_data": bool(prev_map),
        "added_count": len(parts_added),
        "removed_count": len(parts_removed),
        "changed_count": len(parts_changed),
        "blocks_added_count": len(blocks_added),
        "blocks_removed_count": len(blocks_removed),
        "parts_added": parts_added[:200],
        "parts_removed": parts_removed[:200],
        "parts_quantity_changed": parts_changed[:200],
        "blocks_added": blocks_added,
        "blocks_removed": blocks_removed,
    }


def preview_sales_schedule_upload(
    session: Session,
    file_bytes: bytes,
    filename: str,
    sheet_name: Optional[str] = None,
    target_period: Optional[str] = None,
    revision_label: Optional[str] = None,
    require_route_mapping: bool = False,
) -> SalesScheduleImportPreviewResponse:
    t_start = time.perf_counter()
    _cleanup_stale_staged_files()

    # Get available sheets
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True)
        sheet_names = wb.sheetnames
        wb.close()
    except Exception as exc:
        raise ApplicationError("INVALID_EXCEL", f"Could not open workbook: {str(exc)}", 422)

    chosen_sheet, rows = read_workbook_rows(file_bytes, filename, sheet_name)

    t_parse_start = time.perf_counter()
    parse_result = parse_sales_schedule_data(
        rows=rows,
        session=session,
        target_period=target_period,
        revision_label=revision_label,
    )
    t_parse_end = time.perf_counter()

    # 1. Non-blocking Norms Match check against CURRENT ConsumptionNorms
    all_products = session.scalars(select(Product).where(Product.is_active.is_(True))).all()
    active_norm_product_ids = set(session.scalars(
        select(ConsumptionNorm.product_id).where(ConsumptionNorm.is_active.is_(True)).distinct()
    ).all())

    for r in parse_result["_all_flat_rows"]:
        p_part = (r.get("part_no") or "").strip() or None
        p_item = (r.get("item_id") or "").strip() or None
        p_code = (r.get("product_code") or "").strip() or None
        prod = None
        if p_part and p_item:
            prod, _ = match_product(all_products, part_number=p_part, item_id=p_item)
        if not prod and p_code:
            prod, _ = match_product(all_products, code=p_code)
        if not prod and p_part:
            prod, _ = match_product(all_products, part_number=p_part)
        if not prod and p_item:
            prod, _ = match_product(all_products, item_id=p_item)
        has_norm = prod is not None and prod.id in active_norm_product_ids
        r["norm_match"] = "Matched" if has_norm else "No norms found"
        r["norms_matched"] = has_norm

    # Re-slice sample_rows to include norm_match
    parse_result["sample_rows"] = parse_result["_all_flat_rows"][:50]

    # 2. Check for missing routes ONLY IF require_route_mapping is requested (default False for new workflow)
    if require_route_mapping:
        active_plant = session.scalar(select(Plant).where(Plant.is_active.is_(True)))
        if active_plant:
            all_pps = session.scalars(select(ProductPlant).where(ProductPlant.is_active.is_(True))).all()
            product_plants_by_prod: dict[UUID, list[ProductPlant]] = defaultdict(list)
            for pp in all_pps:
                product_plants_by_prod[pp.product_id].append(pp)

            all_routes = session.scalars(select(Route).where(Route.is_active.is_(True))).all()
            routes_by_id: dict[UUID, Route] = {r.id: r for r in all_routes}
            all_ppcs = session.scalars(select(ProductProcessConsumable).where(ProductProcessConsumable.is_active.is_(True))).all()
            ppc_proc_by_prod: dict[UUID, set[UUID]] = defaultdict(set)
            for ppc in all_ppcs:
                ppc_proc_by_prod[ppc.product_id].add(ppc.process_id)
            all_steps = session.scalars(select(RouteStep)).all()
            route_proc_map: dict[UUID, set[UUID]] = defaultdict(set)
            for step in all_steps:
                route_proc_map[step.route_id].add(step.process_id)

            prod_by_pair: dict[tuple[str, str], Product] = {}
            prod_by_code: dict[str, Product] = {}
            for p in all_products:
                if p.part_number and p.item_id:
                    prod_by_pair[(p.part_number.strip().lower(), p.item_id.strip().lower())] = p
                if p.code:
                    prod_by_code[p.code.strip().lower()] = p

            for r in parse_result["_all_flat_rows"]:
                p_part = (r.get("part_no") or "").strip() or None
                p_item = (r.get("item_id") or "").strip() or None
                p_code = (r.get("product_code") or "").strip()
                prod = None
                if p_part and p_item:
                    prod = prod_by_pair.get((p_part.lower(), p_item.lower()))
                if not prod and p_code:
                    prod = prod_by_code.get(p_code.lower())

                if prod:
                    r_id = _resolve_route_for_product(
                        session=session,
                        product_id=prod.id,
                        plant_id=active_plant.id,
                        product_plants_by_prod=product_plants_by_prod,
                        routes_by_id=routes_by_id,
                        ppc_proc_by_prod=ppc_proc_by_prod,
                        route_proc_map=route_proc_map,
                    )
                    if not r_id:
                        warn = f"Product '{prod.code}' (Part: '{p_part or '-'}', Item: '{p_item or '-'}') has no production route assigned for plant '{active_plant.name}'. Map a valid route before importing."
                        if warn not in parse_result["warnings"]:
                            parse_result["warnings"].append(warn)
                else:
                    warn = f"Part '{p_part or p_item or p_code}' does not exist in Product Master and will require route assignment."
                    if warn not in parse_result["warnings"]:
                        parse_result["warnings"].append(warn)

    # 3. Changes vs Previous Month
    rev0 = parse_result["active_revisions"][0] if parse_result.get("active_revisions") else None
    effective_period = (
        (rev0.get("planning_period") if isinstance(rev0, dict) else getattr(rev0, "planning_period", None))
        if rev0
        else (target_period or "2026-10")
    ) or (target_period or "2026-10")
    comparison = compute_changes_vs_previous_month(
        session=session,
        current_period=effective_period,
        current_flat_rows=parse_result["_all_flat_rows"],
        current_blocks=parse_result["blocks"],
    )

    file_token = str(uuid4())
    _STAGED_FILES[file_token] = (file_bytes, filename, time.time())
    t_total = time.perf_counter() - t_start
    logger.info(
        "preview_sales_schedule_upload finished in %.3fs (parse=%.3fs, rows=%d)",
        t_total,
        t_parse_end - t_parse_start,
        len(parse_result.get("_all_flat_rows", [])),
    )

    return SalesScheduleImportPreviewResponse(
        filename=filename,
        sheet_name=chosen_sheet,
        sheet_names=sheet_names,
        header_row_index=parse_result["header_row_index"],
        available_revisions=[SalesScheduleRevisionInfo(**r) for r in parse_result["available_revisions"]],
        active_revisions=[SalesScheduleRevisionInfo(**r) for r in parse_result["active_revisions"]],
        blocks=[SalesScheduleBlockInfo(**b) for b in parse_result["blocks"]],
        unmapped_aliases=parse_result["unmapped_aliases"],
        grand_total=SalesScheduleGrandTotal(**parse_result["grand_total"]),
        total_part_rows=parse_result["total_part_rows"],
        sample_rows=parse_result["sample_rows"],
        can_confirm=parse_result["can_confirm"],
        warnings=parse_result["warnings"],
        errors=parse_result["errors"],
        comparison_vs_previous_month=comparison,
        needs_review_rows=parse_result.get("needs_review_rows", []),
        file_token=file_token,
    )


def confirm_sales_schedule_import(
    session: Session,
    file_token: str,
    sheet_name: str,
    target_period: Optional[str],
    revision_label: Optional[str],
    plant_code: Optional[str],
    actor: UUID,
    alias_mappings: Optional[dict[str, UUID]] = None,
    mode: str = "APPEND",
    reason: str = "Confirmed KNL sales schedule import",
    require_route_mapping: bool = False,
) -> PRDExcelImportResponse:
    """Confirm and atomically save the verified sales schedule into database."""
    _cleanup_stale_staged_files()

    staged_tuple = _STAGED_FILES.get(file_token)
    if not staged_tuple:
        raise ApplicationError("EXPIRED_TOKEN", "Staged file session expired. Please re-upload.", 422)

    file_bytes, filename, _ = staged_tuple

    try:
        # 1. Process inline alias mappings if user supplied them in modal
        if alias_mappings:
            for alias_text, cust_id in alias_mappings.items():
                cust = session.get(Customer, cust_id)
                if cust and cust.is_active:
                    norm = normalize_alias(alias_text)
                    existing = session.scalar(select(CustomerAlias).where(CustomerAlias.normalized_alias == norm))
                    if existing:
                        existing.customer_id = cust.id
                        existing.alias_text = alias_text.strip()
                        existing.is_active = True
                    else:
                        session.add(
                            CustomerAlias(
                                customer_id=cust.id,
                                alias_text=alias_text.strip(),
                                normalized_alias=norm,
                                is_active=True,
                            )
                        )
            session.flush()

        # 2. Re-parse with newly committed aliases
        chosen_sheet, rows = read_workbook_rows(file_bytes, filename, sheet_name)
        parse_result = parse_sales_schedule_data(
            rows=rows,
            session=session,
            target_period=target_period,
            revision_label=revision_label,
        )

        if parse_result["unmapped_aliases"]:
            raise ApplicationError(
                "UNMAPPED_ALIASES",
                f"Cannot import with unmapped company headings: {', '.join(parse_result['unmapped_aliases'])}. Map them first.",
                422,
            )

        all_flat_rows = parse_result["_all_flat_rows"]
        if not all_flat_rows:
            raise ApplicationError("EMPTY_DATA", "No valid parts found to import.", 422)

        # 3. Determine target period & revision
        period = all_flat_rows[0]["target_period"]
        rev_label = all_flat_rows[0]["revision"]
        rev_number = int(rev_label[1:]) if rev_label.startswith("R") and rev_label[1:].isdigit() else 0

        # Ensure revision doesn't collide
        existing_version = session.scalar(
            select(PlanningVersion).where(
                PlanningVersion.planning_period == period,
                PlanningVersion.version_number == rev_number,
            )
        )
        if existing_version and mode == "APPEND":
            raise ApplicationError(
                "PRD_REVISION_EXISTS",
                f"Revision {rev_label} for {period} already exists. Increment revision (e.g. R1) to preserve audit history.",
                409,
            )

        # 4. Ensure plant exists
        effective_plant_name = plant_code.strip() if plant_code else "Plant 1"
        plant_obj = session.scalar(select(Plant).where(Plant.name == effective_plant_name))
        if not plant_obj:
            plant_obj = session.scalar(select(Plant).where(Plant.is_active.is_(True)))
            if not plant_obj:
                plant_obj = Plant(name=effective_plant_name, location=effective_plant_name, is_active=True)
                session.add(plant_obj)
                session.flush()

        # 5. Ensure products and plant mappings exist
        product_cache: dict[str, Product] = {}
        for p in session.scalars(select(Product)).all():
            if p.part_number and p.item_id:
                product_cache[f"pair:{p.part_number.strip().upper()}:{p.item_id.strip().upper()}"] = p
            if p.code:
                product_cache[f"code:{p.code.strip().upper()}"] = p
            if p.part_number:
                product_cache[f"part:{p.part_number.strip().upper()}"] = p
            if p.item_id:
                product_cache[f"item:{p.item_id.strip().upper()}"] = p
            product_cache[str(p.id)] = p

        pps_by_prod_plant: dict[tuple[UUID, UUID], ProductPlant] = {
            (pp.product_id, pp.plant_id): pp
            for pp in session.scalars(select(ProductPlant)).all()
        }
        product_plants_by_prod: dict[UUID, list[ProductPlant]] = defaultdict(list)
        for pp in pps_by_prod_plant.values():
            if pp.is_active:
                product_plants_by_prod[pp.product_id].append(pp)

        routes_by_id: dict[UUID, Route] = {
            r.id: r for r in session.scalars(select(Route).where(Route.is_active.is_(True))).all()
        }
        all_ppcs = session.scalars(select(ProductProcessConsumable).where(ProductProcessConsumable.is_active.is_(True))).all()
        ppc_proc_by_prod: dict[UUID, set[UUID]] = defaultdict(set)
        for ppc in all_ppcs:
            ppc_proc_by_prod[ppc.product_id].add(ppc.process_id)
        all_steps = session.scalars(select(RouteStep)).all()
        route_proc_map: dict[UUID, set[UUID]] = defaultdict(set)
        for step in all_steps:
            route_proc_map[step.route_id].add(step.process_id)

        for r in all_flat_rows:
            p_code = r.get("product_code") or ""
            part_no = r.get("part_no")
            item_id = r.get("item_id")
            description = r.get("description")
            uom = r.get("uom")

            existing_prod = _lookup_or_create_product(
                session=session,
                product_cache=product_cache,
                p_code=p_code,
                part_no=part_no,
                item_id=item_id,
                description=description,
                uom=uom,
            )

            # Route resolution and ProductPlant verification
            target_route_id = _resolve_route_for_product(
                session,
                existing_prod.id,
                plant_obj.id,
                product_plants_by_prod=product_plants_by_prod,
                routes_by_id=routes_by_id,
                ppc_proc_by_prod=ppc_proc_by_prod,
                route_proc_map=route_proc_map,
            )
            if require_route_mapping and not target_route_id:
                raise ApplicationError(
                    "ROUTE_REQUIRED",
                    f"Product '{existing_prod.code}' (Part: '{part_no or '-'}', Item: '{item_id or '-'}') has no production route assigned for plant '{plant_obj.name}'. Map a valid route in Production Mappings before importing.",
                    422,
                )

            pp = pps_by_prod_plant.get((existing_prod.id, plant_obj.id))
            if pp:
                if target_route_id:
                    pp.route_id = target_route_id
                if not pp.is_active:
                    pp.is_active = True
            elif target_route_id:
                pp = ProductPlant(
                    product_id=existing_prod.id,
                    plant_id=plant_obj.id,
                    route_id=target_route_id,
                    is_active=True,
                    is_primary=True,
                )
                session.add(pp)
                pps_by_prod_plant[(existing_prod.id, plant_obj.id)] = pp
                product_plants_by_prod[existing_prod.id].append(pp)

        session.flush()

        # 6. Save PRDRecord rows
        workspace_records: list[PRDRecord] = []
        staged_items: list[dict[str, Any]] = []

        m_parts = period.split("-")
        p_year = int(m_parts[0]) if len(m_parts) > 0 and m_parts[0].isdigit() else None
        p_month = int(m_parts[1]) if len(m_parts) > 1 and m_parts[1].isdigit() else None

        for idx, r in enumerate(all_flat_rows, start=1):
            p_code = r.get("product_code") or ""
            part_no = r.get("part_no")
            item_id = r.get("item_id")
            cached_prod = (
                (product_cache.get(f"pair:{part_no.upper()}:{item_id.upper()}") if part_no and item_id else None)
                or product_cache.get(f"code:{p_code.upper()}")
                or (product_cache.get(f"part:{part_no.upper()}") if part_no else None)
                or (product_cache.get(f"item:{item_id.upper()}") if item_id else None)
            )
            prod_id_str = str(cached_prod.id) if cached_prod else str(uuid4())
            effective_code = cached_prod.code if cached_prod else p_code
            qty_dec = Decimal(r["planned_quantity"])

            rec = PRDRecord(
                id=uuid4(),
                row_index=idx,
                plant=plant_obj.name,
                customer=r["customer"],
                customer_id=UUID(r["customer_id"]) if r.get("customer_id") else None,
                customer_code=r.get("customer_code"),
                customer_group=r.get("customer_group"),
                source_customer_label=r["customer"],
                source_row=r["source_row"],
                product_code=effective_code,
                description=r["description"] or (cached_prod.name if cached_prod else effective_code),
                planned_quantity=qty_dec,
                uom=r["uom"] or (cached_prod.uom if cached_prod else "Nos"),
                target_period=period,
                planning_version=rev_label,
                revision=rev_label,
                schedule_given_on=datetime.fromisoformat(r["schedule_given_on"]) if r.get("schedule_given_on") else None,
                plan_month=p_month,
                plan_year=p_year,
                section_type=r.get("section_type") or "STANDARD",
                status="VALIDATED",
                created_by=actor,
            )
            session.add(rec)
            workspace_records.append(rec)

            staged_items.append({
                "product_code": effective_code,
                "product_id": prod_id_str,
                "plant_code": plant_obj.name,
                "planned_quantity": str(qty_dec),
                "uom": r["uom"] or (cached_prod.uom if cached_prod else "Nos"),
                "target_period": period,
                "source_row_number": r["source_row"],
                "customer_id": r.get("customer_id"),
            })

        # 7. Create ImportBatch and promote to PlanningVersion
        file_sha256 = hashlib.sha256(file_bytes).hexdigest()
        batch = ImportBatch(
            filename=filename,
            file_size_bytes=len(file_bytes),
            uploaded_by=actor,
            status="VALIDATED",
            row_count=len(staged_items),
            valid_row_count=len(staged_items),
            error_row_count=0,
            file_hash=file_sha256,
            validation_summary={
                "blocks_count": len(parse_result["blocks"]),
                "grand_total": parse_result["grand_total"],
                "planning_period": period,
                "revision": rev_label,
                "workflow": "BASIC_CENTRAL_V1",
                "source_format": "SALES_SCHEDULE_V1",
                "production_unit_source": "PRODUCT_MASTER",
            },
            staged_data=staged_items,
        )
        session.add(batch)
        session.flush()

        audit(
            session, actor, batch.id, "IMPORT_VERIFIED_PRD", reason,
            {
                "sheet": sheet_name,
                "filename": filename,
                "sha256": file_sha256,
                "planning_period": period,
                "revision": rev_label,
                "row_count": len(staged_items),
            },
            entity="import_batches",
        )

        promote_batch_to_planning_version(session, batch.id, actor, period, rev_label)
        session.commit()

        # Clean up staged token
        _STAGED_FILES.pop(file_token, None)

        return PRDExcelImportResponse(
            sheet_name=sheet_name,
            imported_count=len(workspace_records),
            records=[to_prd_record_response(r) for r in workspace_records],
        )
    except Exception:
        session.rollback()
        raise


def import_prd_excel(
    session: Session,
    file_bytes: bytes,
    filename: str,
    sheet_name: str,
    actor: UUID,
    mode: str = "APPEND",
    reason: str = "Excel PRD import",
    target_period: Optional[str] = None,
    revision_label: Optional[str] = None,
) -> PRDExcelImportResponse:
    from app.services.product_source import source_rows
    from app.services.prd_source import parse, revision
    from app.models.masters import Product
    from app.models.mappings import ProductPlant
    from app.models.production import Plant
    from app.models.prd import ImportBatch, PlanningVersion
    from app.services.prd import promote_batch_to_planning_version

    if mode != "APPEND":
        raise ApplicationError("PRD_HISTORY_PROTECTED", "Import a new revision instead of replacing revision history.", 409)

    rows = parse(source_rows(file_bytes, filename, sheet_name), target_period, revision_label)
    products = list(session.scalars(select(Product).where(Product.is_active.is_(True))))
    period, label = rows[0]["period"], rows[0]["revision"]

    if session.scalar(
        select(PlanningVersion.id).where(
            PlanningVersion.planning_period == period,
            PlanningVersion.version_number == revision(label),
        )
    ):
        raise ApplicationError("PRD_REVISION_EXISTS", "This revision already exists. Keep its history and use a new revision for changes.", 409)

    staged, workspace = [], []
    seen = set()
    for row in rows:
        matches = [
            p for p in products
            if any(source and source in {p.code, p.item_id, p.part_number} for source in (row["item_id"], row["part_number"], row["code"]))
        ]
        if len(matches) != 1:
            raise ApplicationError("PRODUCT_REVIEW_REQUIRED", f"Row {row['source_row_number']} needs one unambiguous active product master. Review it in Masters → Products.", 422)
        product = matches[0]
        if row["uom"] and row["uom"].casefold() != product.uom.casefold():
            raise ApplicationError("PRD_UNIT_MISMATCH", f"Product {product.code}: source production unit differs from its reviewed master.", 422)

        mappings = list(session.scalars(select(ProductPlant).where(ProductPlant.product_id == product.id, ProductPlant.is_active.is_(True))))
        if row["plant"]:
            mappings = [m for m in mappings if session.get(Plant, m.plant_id).name == row["plant"]]
        elif len(mappings) > 1:
            mappings = [m for m in mappings if m.is_primary]

        if len(mappings) != 1 or not session.get(Plant, mappings[0].plant_id).is_active:
            raise ApplicationError("PLANT_MAPPING_REQUIRED", f"Product {product.code} needs one verified active plant mapping or an explicit source plant.", 422)

        plant = session.get(Plant, mappings[0].plant_id)
        if (product.id, plant.id) in seen:
            raise ApplicationError("DUPLICATE_PRD_ITEM", f"Product {product.code} occurs more than once for {plant.name}; review the source rows.", 422)
        seen.add((product.id, plant.id))

        staged.append({
            "product_code": product.code,
            "product_id": str(product.id),
            "plant_code": plant.name,
            "planned_quantity": str(row["quantity"]),
            "uom": product.uom,
            "target_period": period,
            "source_row_number": row["source_row_number"],
            "customer_id": None,
            "source_item_id": row["item_id"],
            "source_part_number": row["part_number"],
        })
        workspace.append(
            PRDRecord(
                id=uuid4(),
                row_index=row["source_row_number"],
                plant=plant.name,
                product_code=product.code,
                description=product.name,
                planned_quantity=row["quantity"],
                uom=product.uom,
                target_period=period,
                planning_version=label,
                status="VALIDATED",
                created_by=actor,
            )
        )

    file_sha256 = hashlib.sha256(file_bytes).hexdigest()
    batch = ImportBatch(
        filename=filename,
        file_size_bytes=len(file_bytes),
        uploaded_by=actor,
        status="VALIDATED",
        row_count=len(staged),
        valid_row_count=len(staged),
        error_row_count=0,
        staged_data=staged,
    )
    session.add(batch)
    session.add_all(workspace)
    session.flush()

    audit(
        session,
        actor,
        batch.id,
        "IMPORT_VERIFIED_PRD",
        reason,
        {
            "sheet": sheet_name,
            "filename": filename,
            "sha256": file_sha256,
            "planning_period": period,
            "revision": label,
            "row_count": len(staged),
        },
        entity="import_batches",
    )
    promote_batch_to_planning_version(session, batch.id, actor, period, label)
    return PRDExcelImportResponse(
        sheet_name=sheet_name,
        imported_count=len(workspace),
        records=[to_prd_record_response(r) for r in workspace],
    )


# ---------------------------------------------------------------------------
# Excel Export with Full Customer Columns
# ---------------------------------------------------------------------------

def export_prd_excel(
    session: Session,
    search: Optional[str] = None,
    plant: Optional[str] = None,
    target_period: Optional[str] = None,
    planning_version: Optional[str] = None,
    status: Optional[str] = None,
    customer: Optional[str] = None,
    record_ids: Optional[list[UUID]] = None,
) -> bytes:
    query = select(PRDRecord).where(PRDRecord.is_deleted.is_(False))
    if record_ids:
        query = query.where(PRDRecord.id.in_(record_ids))
    else:
        if plant and plant.upper() != "ALL":
            query = query.where(PRDRecord.plant.ilike(f"%{plant.strip()}%"))
        if target_period and target_period.upper() != "ALL":
            query = query.where(PRDRecord.target_period == target_period.strip())
        if planning_version and planning_version.upper() != "ALL":
            query = query.where(
                or_(
                    PRDRecord.planning_version == planning_version.strip(),
                    PRDRecord.revision == planning_version.strip(),
                )
            )
        if status and status.upper() != "ALL":
            query = query.where(PRDRecord.status == status.upper())
        if customer and customer.upper() != "ALL":
            query = query.where(
                or_(
                    PRDRecord.customer.ilike(f"%{customer.strip()}%"),
                    PRDRecord.customer_code.ilike(f"%{customer.strip()}%"),
                )
            )
        if search:
            s = f"%{search.strip()}%"
            query = query.where(
                or_(
                    PRDRecord.product_code.ilike(s),
                    PRDRecord.description.ilike(s),
                    PRDRecord.plant.ilike(s),
                    PRDRecord.customer.ilike(s),
                    PRDRecord.customer_code.ilike(s),
                    PRDRecord.customer_group.ilike(s),
                    PRDRecord.target_period.ilike(s),
                    PRDRecord.planning_version.ilike(s),
                    PRDRecord.remarks.ilike(s),
                )
            )

    rows = session.scalars(
        query.order_by(PRDRecord.row_index.asc().nulls_last(), PRDRecord.created_at.asc())
    ).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PRD Planning"

    header_fill = PatternFill(start_color="1B3A5C", end_color="1B3A5C", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    headers = [
        "#",
        "Customer",
        "Customer ID",
        "Group",
        "Plant",
        "Product / Part No.",
        "Description",
        "Production Plan Qty",
        "Unit",
        "Month / Date",
        "Revision",
        "Section",
        "Status",
        "Remarks",
    ]
    ws.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center" if col_num in (1, 8, 9, 10, 11, 12, 13) else "left",
            vertical="center",
        )

    ws.row_dimensions[1].height = 24

    for r_idx, row in enumerate(rows, start=1):
        qty_val = float(row.planned_quantity) if row.planned_quantity is not None else 0.0
        row_data = [
            row.row_index or r_idx,
            row.customer or "",
            row.customer_code or "",
            row.customer_group or "",
            row.plant,
            row.product_code,
            row.description,
            qty_val,
            row.uom,
            row.target_period,
            row.revision or row.planning_version,
            row.section_type or "STANDARD",
            row.status,
            row.remarks or "",
        ]
        ws.append(row_data)
        current_row = r_idx + 1
        ws.row_dimensions[current_row].height = 20
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if col_num in (1, 9, 10, 11, 12, 13):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_num == 8:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0.0000"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
