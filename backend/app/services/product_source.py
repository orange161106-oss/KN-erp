"""Preview source identities without creating products or guessing units/mappings."""
import io
import hashlib
import re
import zipfile
import openpyxl
from app.core.errors import ApplicationError


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def source_rows(contents: bytes, filename: str, sheet_name: str):
    if not filename.lower().endswith('.xlsx') or not contents or len(contents) > 10 * 1024 * 1024:
        raise ApplicationError('INVALID_SOURCE_FILE', 'Use a nonempty .xlsx file of at most 10 MB.', 422)
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                raise ApplicationError('SOURCE_TOO_LARGE', 'The expanded workbook exceeds the import limit.', 422)
        workbook = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
    except ApplicationError:
        raise
    except Exception:
        raise ApplicationError('INVALID_SOURCE_FILE', 'The workbook could not be read.', 422) from None
    try:
        if sheet_name not in workbook.sheetnames:
            raise ApplicationError('SHEET_NOT_FOUND', 'Select the production-order sheet.', 422)
        sheet = workbook[sheet_name]
        if not hasattr(sheet, 'iter_rows') or sheet.max_column > 128 or sheet.max_row > 20000:
            raise ApplicationError('SOURCE_TOO_LARGE', 'Use a worksheet of at most 20,000 rows and 128 columns.', 422)
        return list(sheet.iter_rows(values_only=True))
    finally:
        workbook.close()


def identity_headers(rows):
    aliases = {'partno': 'part_number', 'partnumber': 'part_number', 'itemid': 'item_id',
               'partdescription': 'name', 'description': 'name', 'productcode': 'code',
               'productpartno': 'part_number', 'unit': 'uom', 'uom': 'uom'}
    for index, row in enumerate(rows[:10]):
        columns = {aliases[normalize_header(value)]: i for i, value in enumerate(row) if normalize_header(value) in aliases}
        if 'name' in columns and ('item_id' in columns or 'part_number' in columns or 'code' in columns):
            return index, columns
    raise ApplicationError('SOURCE_HEADERS_REQUIRED', 'Product identifier and description headers are required.', 422)


def preview(contents: bytes, filename: str, sheet_name: str):
    rows = source_rows(contents, filename, sheet_name)
    header, columns = identity_headers(rows)
    products = {}
    for number, row in enumerate(rows[header + 1:], start=header + 2):
        def value(field):
            cell = row[columns[field]] if field in columns and columns[field] < len(row) else None
            return str(cell).strip() if cell is not None else ''
        item_id, part_number, name = value('item_id'), value('part_number'), value('name')
        if not item_id and not part_number and not value('code'):
            continue  # section/subtotal labels do not have a product identity
        if not name:
            raise ApplicationError('PRODUCT_DESCRIPTION_REQUIRED', f'Source row {number} needs a product description.', 422)
        identity = (item_id, part_number, value('code'))
        candidate = {'item_id': item_id, 'part_number': part_number, 'code': value('code'),
                     'name': name, 'uom': value('uom'), 'source_rows': [number],
                     'description_options': [name], 'unit_options': [value('uom')] if value('uom') else []}
        if identity in products:
            existing = products[identity]
            if name not in existing['description_options']:
                existing['description_options'].append(name)
            if candidate['uom'] and candidate['uom'] not in existing['unit_options']:
                existing['unit_options'].append(candidate['uom'])
            if len(existing['description_options']) > 1:
                existing['name'] = ''  # Requires an explicit reviewer choice.
            existing['uom'] = existing['unit_options'][0] if len(existing['unit_options']) == 1 else ''
            existing['source_rows'].append(number)
        else:
            products[identity] = candidate
    return {'filename': filename, 'sheet': sheet_name, 'sha256': hashlib.sha256(contents).hexdigest(),
            'products': list(products.values()), 'message': 'Preview only. Review identifiers and units before creating masters.'}
