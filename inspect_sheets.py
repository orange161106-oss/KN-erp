import openpyxl
import json

wb = openpyxl.load_workbook("KNL_Consumable_Formula_Map.xlsx", data_only=False)
data_wb = openpyxl.load_workbook("KNL_Consumable_Formula_Map.xlsx", data_only=True)

print("Sheets:", wb.sheetnames)

for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    data_ws = data_wb[sheet_name]
    print(f"\n==================== SHEET: {sheet_name} (Rows: {ws.max_row}, Cols: {ws.max_column}) ====================")
    headers = [str(ws.cell(1, col).value or '').strip() for col in range(1, ws.max_column + 1)]
    print("Headers:", headers)
    
    rows = []
    for r in range(2, ws.max_row + 1):
        row_dict = {}
        has_content = False
        for c in range(1, ws.max_column + 1):
            h = headers[c - 1] if c - 1 < len(headers) and headers[c - 1] else f"Col_{c}"
            formula_val = ws.cell(r, c).value
            calc_val = data_ws.cell(r, c).value
            if formula_val is not None:
                has_content = True
            row_dict[h] = {
                "raw": str(formula_val) if formula_val is not None else None,
                "val": str(calc_val) if calc_val is not None else None
            }
        if has_content:
            rows.append(row_dict)
            
    print(f"Total rows with content: {len(rows)}")
    with open(f"scratch_dump_{sheet_name.replace(' ', '_')}.json", "w", encoding="utf-8") as f:
        json.dump(rows, f, indent=2)

print("\nDump completed successfully.")

